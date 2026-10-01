#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script Import & Đồng Bộ Dữ Liệu Vô Tuyến (RF Data & Cell Sectors) cho Tổ Viễn Thông 3 (TVT3)
Nguồn dữ liệu:
  - 3G.xlsx (Số lượng & trạng thái cell 3G)
  - 4G.xlsx (Số lượng & trạng thái cell 4G)
  - ERA_RF_ALL_2026.xlsx (Thiết kế RF Azimuth, Tilt cơ/điện, Độ cao của 4G và 5G Lớp 1 / Lớp 2)
Đích đến:
  - Bảng Supabase: 'datacells' (Chi tiết từng Cell & thông số RF)
  - Cập nhật trường 'technical_info.rf_summary' trong bảng 'datasites'
"""

import os
import sys
import re
import json
from collections import defaultdict
from datetime import datetime
import openpyxl
from dotenv import load_dotenv
from supabase import create_client

# Load môi trường Supabase
current_dir = os.path.dirname(os.path.abspath(__file__))
project_dir = os.path.dirname(current_dir)
load_dotenv(os.path.join(project_dir, 'backend', '.env'))
load_dotenv(os.path.join(project_dir, 'tvt3_v2', '.env'))

SUPABASE_URL = os.getenv("VITE_SUPABASE_URL", "https://lnmoczxjweuifacqujcu.supabase.co")
SUPABASE_KEY = os.getenv("VITE_SUPABASE_ANON_KEY")

if not SUPABASE_KEY:
    print("❌ Lỗi: Không tìm thấy SUPABASE_KEY trong file .env!")
    sys.exit(1)

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

DATA_DIR = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/data cell"
FILE_3G = os.path.join(DATA_DIR, "3G.xlsx")
FILE_4G = os.path.join(DATA_DIR, "4G.xlsx")
FILE_RF = os.path.join(DATA_DIR, "ERA_RF_ALL_2026.xlsx")

def clean_num(val):
    if val is None: return None
    s = str(val).strip()
    if not s or s.lower() in ['none', 'null', 'nan', '-', '#n/a', '']: return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else round(f, 4)
    except (ValueError, TypeError):
        return None

def extract_sector(cell_name):
    if not cell_name: return 'A'
    s = str(cell_name).strip().upper()
    # Match last character if it is a letter A-Z
    m = re.search(r'([A-Z])$', s)
    if m:
        return m.group(1)
    # Check if there is sector 1, 2, 3
    m2 = re.search(r'([1-9])$', s)
    if m2:
        idx = int(m2.group(1))
        return chr(ord('A') + idx - 1) if 1 <= idx <= 26 else 'A'
    return 'A'

def main():
    print("================================================================================")
    print("🚀 BẮT ĐẦU IMPORT DỮ LIỆU VÔ TUYẾN & CÁNH SÓNG CELL CHO ĐỊA BÀN TVT3")
    print(f"⏰ Thời gian: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("================================================================================")

    # 1. Nạp danh mục trạm TVT3 từ datasites
    print("🔍 Bước 1: Nạp danh mục trạm TVT3 từ bảng 'datasites'...")
    res_sites = supabase.table('datasites').select('site_id, site_id_old, location_info, technical_info').execute()
    datasites_map = {}
    site_lookup = {} # maps site_id_new / site_id_old -> canonical site_id

    for r in (res_sites.data or []):
        sid = (r.get('site_id') or '').strip().upper()
        sold = (r.get('site_id_old') or '').strip().upper()
        if sid:
            datasites_map[sid] = r
            site_lookup[sid] = sid
        if sold:
            site_lookup[sold] = sid

    print(f"✅ Đã nạp {len(datasites_map)} trạm TVT3 ({len(site_lookup)} alias tra cứu).")

    # Lưu trữ thông tin RF theo từng trạm & sector để chia sẻ cho 3G/4G
    # site_rf_sectors[site_id][sector] = { azimuth, height, tilt_total, tilt_mech, tilt_elec, band }
    site_rf_sectors = defaultdict(dict)
    all_cells_to_upsert = {} # cell_id -> cell dict

    # 2. Đọc file ERA_RF_ALL_2026.xlsx (Sheet 5G & 4G)
    print("\n📖 Bước 2: Đọc file ERA_RF_ALL_2026.xlsx...")
    wb_rf = openpyxl.load_workbook(FILE_RF, read_only=True)

    # 2a. Sheet 4G
    ws_4g_rf = wb_rf['4G']
    count_4g_rf = 0
    for r in ws_4g_rf.iter_rows(min_row=2, values_only=True):
        if not r or len(r) < 7: continue
        s_old = str(r[0] or '').strip().upper()
        s_new = str(r[1] or '').strip().upper()
        target_site = site_lookup.get(s_new) or site_lookup.get(s_old)
        if not target_site:
            continue

        cell_name = str(r[2] or '').strip()
        if not cell_name: continue
        band = str(r[3] or '1800').strip()
        lat = clean_num(r[4])
        lon = clean_num(r[5])
        azimuth = clean_num(r[6]) if clean_num(r[6]) is not None else clean_num(r[7])
        height = clean_num(r[8])
        tilt_total = clean_num(r[9])
        tilt_mech = clean_num(r[10])
        tilt_elec = clean_num(r[11])
        sector = extract_sector(cell_name)

        if azimuth is not None:
            site_rf_sectors[target_site][sector] = {
                'azimuth': azimuth,
                'height': height,
                'tilt_total': tilt_total,
                'tilt_mech': tilt_mech,
                'tilt_elec': tilt_elec,
                'band': band
            }

        cell_record = {
            'cell_id': cell_name,
            'site_id': target_site,
            'site_id_old': s_old or None,
            'cell_name_new': cell_name,
            'cell_name_old': None,
            'ran': '4G',
            'vendor': 'ERICSSON',
            'latitude': lat,
            'longitude': lon,
            'azimuth': azimuth,
            'height': height,
            'tilt_total': tilt_total,
            'tilt_mech': tilt_mech,
            'tilt_elec': tilt_elec,
            'band': band,
            'beamwidth': 65,
            'layer_5g': 0,
            'sector': sector,
            'status': 'ACTIVE',
            'updated_at': datetime.utcnow().isoformat()
        }
        all_cells_to_upsert[cell_name] = cell_record
        count_4g_rf += 1

    print(f"✅ Đã trích xuất {count_4g_rf} cell 4G từ sheet 4G khớp trạm TVT3.")

    # 2b. Sheet 5G
    ws_5g_rf = wb_rf['5G']
    count_5g_rf = 0
    for r in ws_5g_rf.iter_rows(min_row=2, values_only=True):
        if not r or len(r) < 7: continue
        s_old = str(r[0] or '').strip().upper()
        s_new = str(r[1] or '').strip().upper()
        target_site = site_lookup.get(s_new) or site_lookup.get(s_old)
        if not target_site:
            continue

        cell_name = str(r[2] or '').strip()
        if not cell_name: continue
        band_raw = clean_num(r[3])
        band_str = str(band_raw) if band_raw else '2600'
        lat = clean_num(r[4])
        lon = clean_num(r[5])
        azimuth = clean_num(r[6]) if clean_num(r[6]) is not None else clean_num(r[7])
        height = clean_num(r[8])
        tilt_total = clean_num(r[9])
        tilt_mech = clean_num(r[10])
        tilt_elec = clean_num(r[11])
        sector = extract_sector(cell_name)

        # Phân loại 5G Lớp 1 (2600 MHz) vs Lớp 2 (3800 MHz)
        layer_5g = 2 if '3800' in band_str else 1

        # Lưu thông số vào sector nếu chưa có
        if sector not in site_rf_sectors[target_site] and azimuth is not None:
            site_rf_sectors[target_site][sector] = {
                'azimuth': azimuth,
                'height': height,
                'tilt_total': tilt_total,
                'tilt_mech': tilt_mech,
                'tilt_elec': tilt_elec,
                'band': f"NR{band_str}"
            }

        cell_record = {
            'cell_id': cell_name,
            'site_id': target_site,
            'site_id_old': s_old or None,
            'cell_name_new': cell_name,
            'cell_name_old': None,
            'ran': '5G',
            'vendor': 'ERICSSON',
            'latitude': lat,
            'longitude': lon,
            'azimuth': azimuth,
            'height': height,
            'tilt_total': tilt_total,
            'tilt_mech': tilt_mech,
            'tilt_elec': tilt_elec,
            'band': f"NR{band_str}",
            'beamwidth': 65,
            'layer_5g': layer_5g,
            'sector': sector,
            'status': 'ACTIVE',
            'updated_at': datetime.utcnow().isoformat()
        }
        all_cells_to_upsert[cell_name] = cell_record
        count_5g_rf += 1

    print(f"✅ Đã trích xuất {count_5g_rf} cell 5G từ sheet 5G khớp trạm TVT3.")

    # 3. Đọc file 4G.xlsx (Active 4G Cells & Traffic/KPIs)
    print("\n📖 Bước 3: Đọc file 4G.xlsx (Active 4G cells & KPIs)...")
    wb_4g = openpyxl.load_workbook(FILE_4G, read_only=True)
    count_4g_active = 0
    for r in wb_4g.active.iter_rows(min_row=4, values_only=True):
        if not r or len(r) < 5: continue
        s_old = str(r[1] or '').strip().upper()
        s_new = str(r[2] or '').strip().upper()
        target_site = site_lookup.get(s_new) or site_lookup.get(s_old)
        if not target_site:
            continue

        c_old = str(r[3] or '').strip() or None
        c_new = str(r[4] or '').strip()
        if not c_new: continue

        vendor = str(r[10] or 'ERICSSON').strip()
        traffic = clean_num(r[11])
        succ_rate = clean_num(r[12])

        sector = extract_sector(c_new)
        rf_sec = site_rf_sectors[target_site].get(sector) or {}

        if c_new in all_cells_to_upsert:
            rec = all_cells_to_upsert[c_new]
            rec['cell_name_old'] = c_old
            rec['vendor'] = vendor
            if traffic is not None:
                rec.setdefault('metadata', {})['volte_traffic'] = traffic
                rec['metadata']['volte_sr'] = succ_rate
        else:
            # Cell có trong 4G.xlsx nhưng chưa có trong file RF -> Thừa hưởng RF của sector
            ds_info = datasites_map.get(target_site) or {}
            loc = ds_info.get('location_info') or {}
            all_cells_to_upsert[c_new] = {
                'cell_id': c_new,
                'site_id': target_site,
                'site_id_old': s_old,
                'cell_name_new': c_new,
                'cell_name_old': c_old,
                'ran': '4G',
                'vendor': vendor,
                'latitude': clean_num(loc.get('vi_do')),
                'longitude': clean_num(loc.get('kinh_do')),
                'azimuth': rf_sec.get('azimuth'),
                'height': rf_sec.get('height'),
                'tilt_total': rf_sec.get('tilt_total'),
                'tilt_mech': rf_sec.get('tilt_mech'),
                'tilt_elec': rf_sec.get('tilt_elec'),
                'band': rf_sec.get('band') or 'L1800',
                'beamwidth': 65,
                'layer_5g': 0,
                'sector': sector,
                'status': 'ACTIVE',
                'metadata': {'volte_traffic': traffic, 'volte_sr': succ_rate} if traffic is not None else {},
                'updated_at': datetime.utcnow().isoformat()
            }
        count_4g_active += 1

    print(f"✅ Đã xử lý {count_4g_active} cell 4G active từ file 4G.xlsx.")

    # 4. Đọc file 3G.xlsx (Active 3G Cells)
    print("\n📖 Bước 4: Đọc file 3G.xlsx (Active 3G cells)...")
    wb_3g = openpyxl.load_workbook(FILE_3G, read_only=True)
    count_3g_active = 0
    for r in wb_3g.active.iter_rows(min_row=4, values_only=True):
        if not r or len(r) < 7: continue
        s_old = str(r[3] or '').strip().upper()
        s_new = str(r[4] or '').strip().upper()
        target_site = site_lookup.get(s_new) or site_lookup.get(s_old)
        if not target_site:
            continue

        c_old = str(r[5] or '').strip() or None
        c_new = str(r[6] or '').strip()
        if not c_new: continue

        vendor = str(r[12] or 'NOKIA').strip()
        rrc_succ = clean_num(r[13])
        rrc_att = clean_num(r[14])

        sector = extract_sector(c_new)
        # Trong SRAN, 3G dùng chung hệ thống anten SRAN với 4G nên thừa hưởng góc Azimuth/Height cùng sector
        rf_sec = site_rf_sectors[target_site].get(sector) or {}

        ds_info = datasites_map.get(target_site) or {}
        loc = ds_info.get('location_info') or {}

        all_cells_to_upsert[c_new] = {
            'cell_id': c_new,
            'site_id': target_site,
            'site_id_old': s_old,
            'cell_name_new': c_new,
            'cell_name_old': c_old,
            'ran': '3G',
            'vendor': vendor,
            'latitude': clean_num(loc.get('vi_do')),
            'longitude': clean_num(loc.get('kinh_do')),
            'azimuth': rf_sec.get('azimuth'),
            'height': rf_sec.get('height'),
            'tilt_total': rf_sec.get('tilt_total'),
            'tilt_mech': rf_sec.get('tilt_mech'),
            'tilt_elec': rf_sec.get('tilt_elec'),
            'band': '2100',
            'beamwidth': 65,
            'layer_5g': 0,
            'sector': sector,
            'status': 'ACTIVE',
            'metadata': {'rrc_succ': rrc_succ, 'rrc_att': rrc_att} if rrc_succ is not None else {},
            'updated_at': datetime.utcnow().isoformat()
        }
        count_3g_active += 1

    print(f"✅ Đã xử lý {count_3g_active} cell 3G active từ file 3G.xlsx.")

    total_cells = len(all_cells_to_upsert)
    print(f"\n📊 TỔNG HỢP: Chuẩn bị nạp {total_cells} cell thuộc các trạm TVT3 vào database!")

    # 5. Upsert vào bảng 'datacells' theo batch
    print("\n💾 Bước 5: Tiến hành Upsert vào bảng 'datacells' (Batch 100)...")
    cell_list = list(all_cells_to_upsert.values())
    batch_size = 100
    for i in range(0, len(cell_list), batch_size):
        batch = cell_list[i:i + batch_size]
        try:
            supabase.table('datacells').upsert(batch).execute()
            print(f"  → Đã lưu {min(i + batch_size, len(cell_list))}/{len(cell_list)} cell...")
        except Exception as e:
            print(f"  ⚠️ Lỗi tại batch {i}-{i+batch_size}: {e}")

    # 6. Tính toán rf_summary và cập nhật vào bảng 'datasites'
    print("\n⚙️ Bước 6: Cập nhật tóm tắt RF 'technical_info.rf_summary' vào bảng 'datasites'...")
    cells_by_site = defaultdict(list)
    for c in cell_list:
        cells_by_site[c['site_id']].append(c)

    updated_sites_count = 0
    for site_id, c_list in cells_by_site.items():
        ds_row = datasites_map.get(site_id)
        if not ds_row: continue

        cells_3g = [c for c in c_list if c['ran'] == '3G']
        cells_4g = [c for c in c_list if c['ran'] == '4G']
        cells_5g = [c for c in c_list if c['ran'] == '5G']
        cells_5g_l1 = [c for c in cells_5g if c['layer_5g'] == 1]
        cells_5g_l2 = [c for c in cells_5g if c['layer_5g'] == 2]

        is_dual_5g = len(cells_5g_l1) > 0 and len(cells_5g_l2) > 0

        # Gom nhóm sectors theo góc hướng
        sec_map = {}
        for c in sorted(c_list, key=lambda x: (x.get('sector') or 'A', x.get('azimuth') or 0)):
            sec = c.get('sector') or 'A'
            az = c.get('azimuth')
            if sec not in sec_map:
                sec_map[sec] = {
                    'sector': sec,
                    'azimuth': az,
                    'height': c.get('height'),
                    'has_3g': False,
                    'has_4g': False,
                    'has_5g_l1': False,
                    'has_5g_l2': False
                }
            if az is not None and sec_map[sec]['azimuth'] is None:
                sec_map[sec]['azimuth'] = az
                sec_map[sec]['height'] = c.get('height')

            if c['ran'] == '3G': sec_map[sec]['has_3g'] = True
            elif c['ran'] == '4G': sec_map[sec]['has_4g'] = True
            elif c['ran'] == '5G':
                if c['layer_5g'] == 2: sec_map[sec]['has_5g_l2'] = True
                else: sec_map[sec]['has_5g_l1'] = True

        rf_summary = {
            'total_cells': len(c_list),
            'cells_3g': len(cells_3g),
            'cells_4g': len(cells_4g),
            'cells_5g': len(cells_5g),
            'cells_5g_l1': len(cells_5g_l1),
            'cells_5g_l2': len(cells_5g_l2),
            'has_5g': len(cells_5g) > 0,
            'is_dual_5g': is_dual_5g,
            'sectors': list(sec_map.values()),
            'last_sync': datetime.utcnow().isoformat()
        }

        # Cập nhật vào technical_info
        current_tech = ds_row.get('technical_info') or {}
        if not isinstance(current_tech, dict): current_tech = {}
        current_tech['rf_summary'] = rf_summary

        try:
            supabase.table('datasites').update({'technical_info': current_tech}).eq('site_id', site_id).execute()
            updated_sites_count += 1
        except Exception as e:
            print(f"  ⚠️ Lỗi cập nhật datasite {site_id}: {e}")

    print(f"✅ Đã cập nhật technical_info.rf_summary cho {updated_sites_count} trạm TVT3!")
    print("\n🎉 HOÀN THÀNH TOÀN BỘ TIẾN TRÌNH IMPORT DỮ LIỆU VÔ TUYẾN!")
    print("================================================================================")

if __name__ == '__main__':
    main()
