import os, csv, math, re, json
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from supabase import create_client

load_dotenv('tvt3_v2/.env')
sb = create_client(os.getenv('VITE_SUPABASE_URL'), os.getenv('VITE_SUPABASE_ANON_KEY'))

# Load OSRM bike cache
cache_file = 'scratch/osrm_bike_cache.json'
cache = {}
if os.path.exists(cache_file):
    with open(cache_file) as f:
        cache = json.load(f)

def dms_to_dd(dms_str):
    s = dms_str.strip()
    if not s or 'bỏ' in s.lower() or 'trùng' in s.lower(): return None, None
    m_dec = re.match(r'^([\d\.]+)\s*,\s*([\d\.]+)$', s)
    if m_dec: return float(m_dec.group(1)), float(m_dec.group(2))
    pat = r'(\d+)[°\s]+(\d+)[\'\s]+([\d\.]+)[\"\s]*([NSns])\s*[,;\s]*(\d+)[°\s]+(\d+)[\'\s]+([\d\.]+)[\"\s]*([EWew])'
    m = re.search(pat, s)
    if m:
        lat = float(m.group(1)) + float(m.group(2))/60.0 + float(m.group(3))/3600.0
        if m.group(4).upper() == 'S': lat = -lat
        lon = float(m.group(5)) + float(m.group(6))/60.0 + float(m.group(7))/3600.0
        if m.group(8).upper() == 'W': lon = -lon
        return lat, lon
    return None, None

def haversine(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi, dlam = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dphi/2.0)**2 + math.cos(phi1)*math.cos(phi2)*math.sin(dlam/2.0)**2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))

# 1. Fetch datasites and filter ACTIVE, non-MORAN, non-PTM
print("Fetching datasites from Supabase...")
datasites = sb.table('datasites').select('site_id, site_id_old, status, technical_info, management_info, location_info').execute().data

site_dict = {}
active_bts = []

for r in datasites:
    s_new = (r.get('site_id') or '').strip().upper()
    s_old = (r.get('site_id_old') or '').strip().upper()
    status = (r.get('status') or '').strip().upper()
    
    if 'UL_' in s_old or '4G-' in s_old or 'MBF_' in s_old or 'MOCN' in s_old: continue
    if status != 'ACTIVE': continue
        
    loc = r.get('location_info') or {}
    lat, lon = loc.get('vi_do'), loc.get('kinh_do')
    if not lat or not lon: continue
    try:
        lat, lon = float(lat), float(lon)
    except:
        continue
    
    mi = r.get('management_info') or {}
    ti = r.get('technical_info') or {}
    rf = ti.get('rf_summary') or {}
    
    has_5g = rf.get('has_5g') or False
    # CRITICAL RULE from User: All 5G deployed sites from SRAN + 5G project have CSG! (Except DNDQ15, DNDQ33)
    is_5g_deployed = has_5g and s_old not in ['DNDQ15', 'DNDQ33']
    
    vung = str(mi.get('vung_phu') or '').strip().upper()
    tram_main = str(mi.get('tram_main') or ti.get('tram_main_tx') or '').strip().upper()
    tb_csg = str(ti.get('thiet_bi_csg') or '').strip().upper()
    route = str(ti.get('last_mile_primary') or '').strip().upper()
    dv_vh = str(ti.get('don_vi_van_hanh_cap') or ti.get('don_vi_van_hanh') or '').strip().upper()
    
    is_cran = 'CRAN' in vung
    
    # CSG detection:
    # 1. Has 5G deployed (from SRAN + 5G project) -> DEFINITELY CSG (e.g. DNXL20, DNXL09, DNXL06, DNXL49, DNXL22, DNXL24)
    # 2. Known CSG Hubs in TVT3
    # 3. Explicit CSG in tb_csg or route
    known_csgs = [
        'DNXL20', 'DNXL09', 'DNXL06', 'DNXL49', 'DNXL22', 'DNXL24', 'DNXL44', 'DNXL23', 'DNXL07', 
        'DNXL75', 'DNXL38', 'DNXL47', 'DNXL13', 'DNXL39', 'DNXL40', 'DNXL56', 'DNXL64', 'DNXL28', 
        'DNXL01', 'DNXL33', 'DNXL34', 'DNXL37', 'DNXL02', 'DNXL03', 'DNXL04', 'DNXL12', 'DNXL14', 
        'DNXL16', 'DNXL17', 'DNXL30', 'DNXL32', 'DNXL35', 'DNXL48', 'DNXL58', 'DNXL65'
    ]
    is_csg = (not is_cran) and (is_5g_deployed or ('CSG' in tb_csg) or ('LOCAL CSG' in route) or (s_old in known_csgs))
    has_lsw = ('LSW' in tb_csg) or ('LSW' in route) or (s_old in ['DNXL36', 'DNXL31', 'DNXL41', 'DNXL05'])
    is_macro = not is_cran and not is_csg
    
    item = {
        'site_id': s_new, 'site_id_old': s_old, 'lat': lat, 'lon': lon,
        'vung': vung, 'tram_main': tram_main, 'is_csg': is_csg,
        'has_lsw': has_lsw, 'is_cran': is_cran, 'is_macro': is_macro,
        'has_5g': is_5g_deployed, 'tb_csg': tb_csg,
        'qlt': mi.get('qlt', '')
    }
    site_dict[s_new] = item
    site_dict[s_old] = item
    active_bts.append(item)

# Resolve Main for CRAN sites
for s in active_bts:
    if s['is_cran']:
        tm = s['tram_main']
        m = re.search(r'([A-Za-z0-9_]+)', tm)
        code = m.group(1).upper() if m else tm
        resolved = site_dict.get(code)
        if resolved:
            s['main_site'] = resolved
            s['main_code'] = resolved['site_id_old'] or resolved['site_id']
            s['main_is_csg'] = resolved['is_csg']
            s['main_has_lsw'] = resolved['has_lsw']
        else:
            s['main_site'] = None
            s['main_code'] = tm
            s['main_is_csg'] = False
            s['main_has_lsw'] = False
    else:
        s['main_site'] = s
        s['main_code'] = s['site_id_old'] or s['site_id']
        s['main_is_csg'] = s['is_csg']
        s['main_has_lsw'] = s['has_lsw']

# 2. Road Network Routing Rules & Cable Capacity Sizing
road_route_defs = {
    '1': ('DNXL21', 'Route 1: QL1A Đông (Xuân Hòa - Bình Thuận)', '8FO'),
    '24': ('DNXL21', 'Route 1: QL1A Đông (Xuân Hòa - Bình Thuận)', '8FO'),
    '55': ('DNXL21', 'Route 1: QL1A Đông (Xuân Hòa - Bình Thuận)', '8FO'),
    
    '26': ('DNXL49', 'Route 2: Tuyến Đường Cọ Dầu Ấp 3', '8FO'),
    '27': ('DNXL49', 'Route 2: Tuyến Đường Cọ Dầu Ấp 3', '8FO'),
    '42': ('DNXL49', 'Route 2: Tuyến Đường Cọ Dầu Ấp 3', '8FO'),
    
    # Route 3: Điểm 7-53-52 kéo về DNXL20 (Đã có CSG sau 5G!)
    '7': ('DNXL20', 'Route 3: Tuyến Đường Suối Bà Rùa (Điểm 7-53-52)', '8FO'),
    '53': ('DNXL20', 'Route 3: Tuyến Đường Suối Bà Rùa (Điểm 7-53-52)', '8FO'),
    '52': ('DNXL20', 'Route 3: Tuyến Đường Suối Bà Rùa (Điểm 7-53-52)', '8FO'),
    
    '54': ('DNXL05', 'Route 4: Trục QL1A & Tuyến Suối Đá (Chùa Phổ Minh)', '8FO'),
    '5': ('DNXL05', 'Route 4: Trục QL1A & Tuyến Suối Đá (Chùa Phổ Minh)', '8FO'),
    '33': ('DNXL05', 'Route 4: Trục QL1A & Tuyến Suối Đá (Chùa Phổ Minh)', '8FO'),
    
    '28': ('DNXL05', 'Route 5: Cụm QL1A - Trạm Dừng Chân & Nhà Thờ', '8FO'),
    '29': ('DNXL05', 'Route 5: Cụm QL1A - Trạm Dừng Chân & Nhà Thờ', '8FO'),
    '30': ('DNXL05', 'Route 5: Cụm QL1A - Trạm Dừng Chân & Nhà Thờ', '8FO'),
    
    '4': ('DNXL69', 'Route 6: Cụm CRAN DNXL69 & Vòng Cung Tín Nghĩa', '8FO'),
    '39': ('DNXL69', 'Route 6: Cụm CRAN DNXL69 & Vòng Cung Tín Nghĩa', '8FO'),
    '2': ('DNXL69', 'Route 6: Cụm CRAN DNXL69 & Vòng Cung Tín Nghĩa', '8FO'),
    
    '31': ('DNXL50', 'Route 7: Tuyến Đường Sóc Ba Buông & Mỏ Cát', '8FO'),
    '32': ('DNXL50', 'Route 7: Tuyến Đường Sóc Ba Buông & Mỏ Cát', '8FO'),
    '49': ('DNXL50', 'Route 7: Tuyến Đường Sóc Ba Buông & Mỏ Cát', '8FO'),
    
    '46': ('DNXL44', 'Route 8: Tuyến Đồi Môn Tây - Trại Heo & Tà Lú', '12FO'),
    '25': ('DNXL44', 'Route 8: Tuyến Đồi Môn Tây - Trại Heo & Tà Lú', '12FO'),
    '45': ('DNXL44', 'Route 8: Tuyến Đồi Môn Tây - Trại Heo & Tà Lú', '12FO'),
    '43': ('DNXL44', 'Route 8: Tuyến Đồi Môn Tây - Trại Heo & Tà Lú', '12FO'),
    '13': ('DNXL44', 'Route 8: Tuyến Đồi Môn Tây - Trại Heo & Tà Lú', '12FO'),
    
    '6': ('DNXL24', 'Route 9: Cụm THPT Xuân Hưng - Suối Cầu - Be 79', '12FO'),
    '8': ('DNXL24', 'Route 9: Cụm THPT Xuân Hưng - Suối Cầu - Be 79', '12FO'),
    '34': ('DNXL24', 'Route 9: Cụm THPT Xuân Hưng - Suối Cầu - Be 79', '12FO'),
    '35': ('DNXL24', 'Route 9: Cụm THPT Xuân Hưng - Suối Cầu - Be 79', '12FO'),
    '51': ('DNXL24', 'Route 9: Cụm THPT Xuân Hưng - Suối Cầu - Be 79', '12FO'),
    
    '11': ('DNXL78', 'Route 10: Cụm Làng Chăm - Thánh Đường - Riverside', '8FO'),
    '12': ('DNXL78', 'Route 10: Cụm Làng Chăm - Thánh Đường - Riverside', '8FO'),
    '41': ('DNXL78', 'Route 10: Cụm Làng Chăm - Thánh Đường - Riverside', '8FO'),
    '50': ('DNXL78', 'Route 10: Cụm Làng Chăm - Thánh Đường - Riverside', '8FO'),
    
    # Route 11: 4 cam QL1A (14, 15, 40, 17) gom trục 8FO về CSG-DNXL09
    '14': ('DNXL09', 'Route 11: Trục QL1A Xuân Tâm (Điểm 14-15-40-17)', '8FO'),
    '15': ('DNXL09', 'Route 11: Trục QL1A Xuân Tâm (Điểm 14-15-40-17)', '8FO'),
    '40': ('DNXL09', 'Route 11: Trục QL1A Xuân Tâm (Điểm 14-15-40-17)', '8FO'),
    '17': ('DNXL09', 'Route 11: Trục QL1A Xuân Tâm (Điểm 14-15-40-17)', '8FO'),
    
    # Route 11B: Điểm 23 tách riêng tuyến nhánh 2FO về CSG-DNXL09
    '23': ('DNXL09', 'Route 11B: Tuyến Nhánh BQL Rừng Phòng Hộ', '2FO'),
    
    '10': ('DNXL63', 'Route 12: Cụm Hầm Chui Cao Tốc & Trại Heo Anh Khôi', '4FO'),
    '48': ('DNXL63', 'Route 12: Cụm Hầm Chui Cao Tốc & Trại Heo Anh Khôi', '4FO'),
    
    '18': ('DNXL06', 'Route 13: Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', '12FO'),
    '47': ('DNXL06', 'Route 13: Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', '12FO'),
    '19': ('DNXL06', 'Route 13: Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', '12FO'),
    '36': ('DNXL06', 'Route 13: Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', '12FO'),
    '20': ('DNXL67', 'Route 13: Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', '12FO'),
    
    '21': ('DNXL75', 'Route 14: Vòng Xoay KCN & Tuyến Phân Trường 5', '8FO'),
    '44': ('DNXL36', 'Route 14: Vòng Xoay KCN & Tuyến Phân Trường 5', '8FO'),
    '37': ('DNXL36', 'Route 14: Vòng Xoay KCN & Tuyến Phân Trường 5', '8FO'),
    
    '38': ('DNXL26', 'Route 15: Cụm Hồ Gia Ui - Đập & Xã Gia Ui', '8FO'),
    '16': ('DNXL18', 'Route 15: Cụm Hồ Gia Ui - Đập & Xã Gia Ui', '8FO'),
    '22': ('DNXL18', 'Route 15: Cụm Hồ Gia Ui - Đập & Xã Gia Ui', '8FO')
}

with open('scratch/camera_survey.csv', 'r', encoding='utf-8') as f:
    survey_rows = list(csv.reader(f))

coord_overrides = {'49': (10.854361, 107.533917)}

processed_cams = []

for r in survey_rows[3:]:
    if len(r) > 4 and r[0].strip() and r[0].strip() != 'xxx':
        stt = r[0].strip()
        name = r[1].strip()
        point_id = r[2].strip()
        coord_raw = r[4].strip()
        map_link = r[5].strip() if len(r) > 5 else ''
        note_loc = r[21].strip() if len(r) > 21 else (r[22].strip() if len(r) > 22 else '')
        
        if stt in coord_overrides:
            lat, lon = coord_overrides[stt]
        else:
            lat, lon = dms_to_dd(coord_raw)
            
        if not lat or not lon:
            status_text = 'Điểm bỏ' if 'bỏ' in coord_raw.lower() else ('Điểm trùng' if 'trùng' in coord_raw.lower() else 'Chưa có tọa độ')
            processed_cams.append({
                'stt': stt, 'point_id': point_id, 'name': name, 'valid': False,
                'status_text': status_text, 'lat': None, 'lon': None,
                'col_h': '', 'col_i': '', 'col_j': '', 'col_k': '', 'col_l': '',
                'note': status_text, 'raw_row': r
            })
            continue

        route_tuple = road_route_defs.get(stt)
        if route_tuple:
            target_bts_id, route_name, cable_type = route_tuple
            target_bts = site_dict.get(target_bts_id)
        else:
            closest_bts = min(active_bts, key=lambda s: haversine(lat, lon, s['lat'], s['lon']))
            target_bts = closest_bts
            target_bts_id = target_bts['site_id_old']
            route_name = f"Nhánh cục bộ {target_bts_id}"
            cable_type = "2FO"

        cache_key = f"{lon:.5f},{lat:.5f}->{target_bts['lon']:.5f},{target_bts['lat']:.5f}"
        route_res = cache.get(cache_key)
        
        if route_res:
            road_dist_m = route_res['dist']
            road_coords = route_res['coords']
        else:
            road_dist_m = haversine(lat, lon, target_bts['lat'], target_bts['lon']) * 1.3
            road_coords = [[lon, lat], [target_bts['lon'], target_bts['lat']]]

        est_cable_m = int(math.ceil((road_dist_m * 1.12) / 50.0) * 50)
        if est_cable_m < 150: est_cable_m = 150

        t_old = target_bts['site_id_old']
        is_csg = target_bts['is_csg']
        is_cran = target_bts['is_cran']
        
        if is_cran:
            bts_type = 'CRAN OUT DOOR'
            main_code = target_bts['main_code']
            main_is_csg = target_bts['main_is_csg']
            cran_main_bts = target_bts.get('main_site')
            is_cran_tap = True
            
            if main_is_csg:
                col_h = f"CSG-{main_code}"
                col_i = "Port CSG"
                note = f"Kéo vào CRAN {t_old} ({int(road_dist_m)}m đường chạy xe), mượn core quang hiện hữu về CSG-{main_code}"
            else:
                main_has_lsw = target_bts.get('main_has_lsw', False)
                if main_has_lsw:
                    col_h = f"{main_code} (Đã có LSW)"
                    col_i = "Port LSW"
                    note = f"Kéo vào CRAN {t_old} ({int(road_dist_m)}m đường chạy xe), mượn core quang hiện hữu về LSW {main_code}"
                else:
                    col_h = f"{main_code} (Trang bị thêm LSW)"
                    col_i = "Port LSW"
                    note = f"Kéo vào CRAN {t_old} ({int(road_dist_m)}m đường chạy xe), mượn core quang hiện hữu về MAIN {main_code} (cần trang bị thêm 01 Switch LSW)"
        elif is_csg:
            bts_type = 'CSG Hub'
            col_h = f"CSG-{t_old}"
            col_i = "Port CSG"
            note = f"Kéo cáp theo đường chạy xe ({int(road_dist_m)}m) trực tiếp về cổng CSG-{t_old}"
            is_cran_tap = False
            cran_main_bts = None
        else: # MACRO
            bts_type = 'MACRO'
            is_cran_tap = False
            cran_main_bts = None
            is_viba = 'VR4' in str(target_bts.get('tb_csg', '')).upper() or 'VIBA' in str(target_bts.get('tb_csg', '')).upper()
            viba_note = " [Trạm hiện hữu dùng Viba iPSL_VR4 -> Tuyến cáp camera kết hợp quang hóa luôn cho trạm Lastmile]" if is_viba else ""
            if target_bts['has_lsw']:
                col_h = f"{t_old} (Đã có LSW)"
                col_i = "Port LSW"
                note = f"Kéo cáp theo đường chạy xe ({int(road_dist_m)}m) về Switch LSW trạm {t_old}{viba_note}"
            else:
                col_h = f"{t_old} (Trang bị thêm LSW)"
                col_i = "Port LSW"
                note = f"Kéo cáp theo đường chạy xe ({int(road_dist_m)}m) về trạm Macro {t_old} (cần trang bị thêm 01 Switch LSW){viba_note}"

        # Special notes
        if stt in ['7', '53', '52']:
            note = f"Kéo cáp theo đường chạy xe ({int(road_dist_m)}m) về CSG-DNXL20 | Tuyến Suối Bà Rùa gom chung cáp 8FO rẽ măng xông FOSC (3 cam dùng 6 core, dư 2 core)"
        elif stt in ['14', '15', '40', '17']:
            note = f"Kéo cáp theo đường chạy xe ({int(road_dist_m)}m) về CSG-DNXL09 | Tuyến trục QL1A gom chung cáp 8FO rẽ măng xông FOSC (4 cam dùng 8 core)"
        elif stt == '23':
            note = f"Kéo riêng 1 sợi cáp 2FO theo đường độc đạo lên rừng ({int(road_dist_m)}m chạy xe) về CSG-DNXL09 (dùng 2 core cho 1 cam)"

        processed_cams.append({
            'stt': stt, 'point_id': point_id, 'name': name, 'valid': True,
            'status_text': 'Hoạt Động',
            'lat': lat, 'lon': lon,
            'bts_old': t_old,
            'bts_type': bts_type,
            'bts_lat': target_bts['lat'],
            'bts_lon': target_bts['lon'],
            'is_cran_tap': is_cran_tap,
            'cran_main_bts': cran_main_bts,
            'road_dist_m': int(road_dist_m),
            'dist_m': est_cable_m,
            'col_h': col_h,
            'col_i': col_i,
            'col_j': est_cable_m,
            'col_k': route_name,
            'col_l': cable_type,
            'note': note,
            'map_link': map_link,
            'note_loc': note_loc,
            'road_coords': road_coords
        })

print(f"Processed {len(processed_cams)} camera records with post-SRAN 5G CSG classification.")

# 3. Export CSV for Copy/Paste
out_csv_path = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/Cam Xa Xuan Hoa/Cot_H_den_L_De_Copy_Paste.csv"
with open(out_csv_path, 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow([
        'STT', 'Mã Điểm', 'Tên Vị Trí Camera', 'Tọa Độ', 
        'Cột H (CSG / Trạm LM)', 'Cột I (Port CSG/LSW)', 'Cột J (Khoảng Cách Cáp Thực Tế m)', 
        'Cột K (Route Tuyến Đường)', 'Cột L (Loại Cáp Tối Ưu)', 'Ghi Chú Kỹ Thuật Đấu Nối'
    ])
    for c in processed_cams:
        coord_str = f"{c['lat']:.6f}, {c['lon']:.6f}" if c['valid'] else c['status_text']
        writer.writerow([
            c['stt'], c['point_id'], c['name'], coord_str,
            c['col_h'], c['col_i'], c['col_j'], c['col_k'], c['col_l'], c['note']
        ])
print(f"Updated CSV: {out_csv_path}")

# 4. Export Excel with Rich Formatting
out_xlsx_path = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/Cam Xa Xuan Hoa/Phan_Tich_Keo_Cap_Camera_Xuan_Loc_20261016.xlsx"
wb = openpyxl.Workbook()

# Sheet 1: Quy hoạch chi tiết 64 Cam
ws1 = wb.active
ws1.title = "Quy_Hoach_Camera_Xuan_Loc"
ws1.views.sheetView[0].showGridLines = True

f_title = Font(name='Calibri', size=14, bold=True, color='1F497D')
f_sub = Font(name='Calibri', size=10, italic=True, color='595959')
f_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
f_data = Font(name='Calibri', size=9)
f_data_bold = Font(name='Calibri', size=9, bold=True)
f_red = Font(name='Calibri', size=9, color='C00000', bold=True)
f_green = Font(name='Calibri', size=9, color='008000', bold=True)
f_blue = Font(name='Calibri', size=9, color='002060', bold=True)

fill_h_group1 = PatternFill('solid', fgColor='1F497D')
fill_h_group2 = PatternFill('solid', fgColor='366092')
fill_h_group3 = PatternFill('solid', fgColor='C00000')

fill_zebra = PatternFill('solid', fgColor='F2F5F8')
fill_route3 = PatternFill('solid', fgColor='E2EFDA')
fill_route11 = PatternFill('solid', fgColor='FCE4D6') # Peach highlight for Route 11
fill_invalid = PatternFill('solid', fgColor='F2F2F2')

thin_side = Side(style='thin', color='D9D9D9')
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')

ws1.merge_cells('A1:L1')
ws1['A1'] = "QUY HOẠCH HẠ TẦNG TRUYỀN DẪN QUANG CAMERA XUÂN LỘC (CẬP NHẬT TRẠM CSG HẬU DỰ ÁN SRAN + 5G)"
ws1['A1'].font = f_title
ws1['A1'].alignment = align_left

ws1.merge_cells('A2:L2')
ws1['A2'] = "Các trạm lên 5G (DNXL20, DNXL09, DNXL06, DNXL49, DNXL22, DNXL24...) đều là CSG Hub. 1 cam = 2FO; 2 cam = 4FO; 3-4 cam = 8FO; 5-6 cam = 12FO."
ws1['A2'].font = f_sub
ws1['A2'].alignment = align_left

headers_s1 = [
    'STT', 'Vị Trí / Mạng', 'Channel', 'Tọa Độ Khảo Sát', 'Map', 'Cột Camera',
    'Cột H\nCSG / Trạm LM', 'Cột I\nPort CSG/LSW', 'Cột J\nKC Cáp Thực Tế (m)', 
    'Cột K\nRoute Gom Tuyến', 'Cột L\nLoại Cáp Tối Ưu', 'Ghi Chú Kỹ Thuật & Đấu Nối'
]

ws1.row_dimensions[4].height = 32
for c_idx, h in enumerate(headers_s1, 1):
    cell = ws1.cell(row=4, column=c_idx, value=h)
    cell.font = f_header
    cell.alignment = align_center
    cell.border = border_cell
    if c_idx in [7, 8, 9, 10, 11, 12]:
        cell.fill = fill_h_group3
    elif c_idx <= 3:
        cell.fill = fill_h_group1
    else:
        cell.fill = fill_h_group2

row_cur = 5
for idx, c in enumerate(processed_cams):
    ws1.row_dimensions[row_cur].height = 24
    stt_val = c['stt']
    pt_val = c['point_id']
    name_val = c['name']
    coord_val = f"{c['lat']:.5f}, {c['lon']:.5f}" if c['valid'] else c['status_text']
    map_val = c.get('map_link', '')
    cot_val = pt_val
    col_h_val = c['col_h']
    col_i_val = c['col_i']
    col_j_val = c['col_j']
    col_k_val = c['col_k']
    col_l_val = c['col_l']
    note_val = c['note']
    
    vals = [stt_val, name_val, pt_val, coord_val, map_val, cot_val,
            col_h_val, col_i_val, col_j_val, col_k_val, col_l_val, note_val]
    
    for c_idx, val in enumerate(vals, 1):
        cell = ws1.cell(row=row_cur, column=c_idx, value=val)
        cell.border = border_cell
        cell.font = f_data
        
        if not c['valid']:
            cell.fill = fill_invalid
            cell.alignment = align_center
            if c_idx == 4 or c_idx == 12: cell.font = f_red
        else:
            if stt_val in ['7', '53', '52']:
                cell.fill = fill_route3
            elif stt_val in ['14', '15', '40', '17', '23']:
                cell.fill = fill_route11
            elif idx % 2 == 1:
                cell.fill = fill_zebra
                
            if c_idx in [7, 8, 10, 11]:
                cell.font = f_data_bold
                if 'CSG-' in str(val): cell.font = f_blue
                if 'Trang bị thêm LSW' in str(val): cell.font = f_red
                if stt_val in ['7', '53', '52'] and c_idx in [10, 11]: cell.font = f_green
            if c_idx in [1, 3, 5, 6, 8, 11]:
                cell.alignment = align_center
            elif c_idx in [9]:
                cell.alignment = align_right
            else:
                cell.alignment = align_left
    row_cur += 1

col_widths = {1: 8, 2: 36, 3: 12, 4: 24, 5: 12, 6: 12, 7: 28, 8: 14, 9: 18, 10: 45, 11: 18, 12: 55}
for col_i, w in col_widths.items():
    ws1.column_dimensions[get_column_letter(col_i)].width = w

# Sheet 2: Tổng hợp Tuyến Trục
ws2 = wb.create_sheet(title="Tong_Hop_15_Route_Giao_Thong")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:H1')
ws2['A1'] = "TỔNG HỢP CÁC TUYẾN TRỤC GOM QUANG THEO QUY TẮC TỐI ƯU CÁP & TRẠM CSG HẬU 5G"
ws2['A1'].font = f_title

headers_s2 = ['Route ID', 'Tên Tuyến Đường / Hành Lang', 'Số Cam', 'Core Dùng', 'Cáp Tối Ưu', 'Core Dư', 'Đích Đấu Nối (Cột H)', 'Phương Án Đấu Nối Măng Xông']
ws2.row_dimensions[3].height = 28
for c_idx, h in enumerate(headers_s2, 1):
    cell = ws2.cell(row=3, column=c_idx, value=h)
    cell.font = f_header
    cell.fill = fill_h_group1
    cell.alignment = align_center
    cell.border = border_cell

routes_summary_data = [
    ('Route 1', 'QL1A Đông (Xuân Hòa - Bình Thuận giáp ranh Cầu Sông Phan)', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL21 (Trang bị thêm LSW)', 'Trục 8FO dọc QL1A và ngã 3 cầu Thanh Niên Xung Phong về trạm DNXL21.'),
    ('Route 2', 'Tuyến Đường Cọ Dầu Ấp 3 & Cồn Tùng Lâm', 3, '6FO (3x2)', '8FO', '2FO', 'CSG-DNXL49', 'Trục 8FO dọc đường Cọ Dầu ấp 3 về trạm CSG-DNXL49 (đã lên 5G có CSG).'),
    ('Route 3', 'Tuyến Đường Suối Bà Rùa (Điểm 7 - 53 - 52)', 3, '6FO (3x2)', '8FO', '2FO', 'CSG-DNXL20', 'Chung 1 tuyến cáp trục 8FO dài 1.450m dọc đường Suối Bà Rùa về CSG-DNXL20 (đã lên 5G có CSG). Đặt FOSC tại Điểm 7 và 53, kết thúc tại Điểm 52.'),
    ('Route 4', 'Trục QL1A & Tuyến Suối Đá (Chùa Phổ Minh)', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL05 (Đã có Switch Cisco A901)', 'Trục 8FO dọc QL1A đối diện cty Trường Vinh và đường Suối Đá về trạm DNXL05.'),
    ('Route 5', 'Cụm QL1A - Trạm Dừng Chân & Nhà Thờ', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL05 (Đã có Switch Cisco A901)', 'Cụm camera ngay sát bên hông trạm DNXL05, kéo cáp trục 8FO gom về trạm.'),
    ('Route 6', 'Cụm CRAN DNXL69 & Vòng Cung Cây Xăng Tín Nghĩa', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL05 (qua CRAN DNXL69)', 'Kéo cáp trục 8FO về tủ CRAN DNXL69, mượn sợi quang có sẵn về MAIN DNXL05. Không cần quang hóa mới.'),
    ('Route 7', 'Tuyến Đường Sóc Ba Buông & Mỏ Cát', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL50 (Trang bị thêm LSW)', 'Trục 8FO dọc đường Sóc Ba Buông vượt cao tốc Phương Trang về trạm DNXL50.'),
    ('Route 8', 'Tuyến Đồi Môn Tây - Trại Heo Kim Hoàng Hưng & Tà Lú', 5, '10FO (5x2)', '12FO', '2FO', 'CSG-DNXL44', 'Trục chính 12FO chạy dọc toàn bộ trục đường Đồi Môn về phòng máy trung tâm CSG-DNXL44.'),
    ('Route 9', 'Cụm THPT Xuân Hưng - Suối Cầu - Be 79', 5, '10FO (5x2)', '12FO', '2FO', 'CSG-DNXL24', 'Trục 12FO gom 5 điểm sau trường THPT Xuân Hưng và Suối Cầu về trạm CSG-DNXL24 (đã lên 5G có CSG).'),
    ('Route 10', 'Cụm Làng Chăm - Thánh Đường - Dốc Cầu Riverside', 4, '8FO (4x2)', '8FO', '0FO', 'CSG-DNXL44 (qua CRAN DNXL78)', 'Trục 8FO trong KDC ấp 4 về tủ CRAN DNXL78, mượn cáp quang hiện hữu về CSG-DNXL44.'),
    ('Route 11', 'Trục QL1A Xuân Tâm (Điểm 14-15-40-17)', 4, '8FO (4x2)', '8FO', '0FO', 'CSG-DNXL09', 'Tuyến trục 8FO gom 4 cam dọc QL1A về trạm CSG-DNXL09 (đã lên 5G có CSG). FOSC tại Điểm 14, 15, 40.'),
    ('Route 11B', 'Tuyến Nhánh BQL Rừng Phòng Hộ (Điểm 23)', 1, '2FO (1x2)', '2FO', '0FO', 'CSG-DNXL09', 'Tuyến nhánh độc đạo dài 2.4km lên rừng kéo riêng 1 sợi cáp 2FO về phòng máy CSG-DNXL09.'),
    ('Route 12', 'Cụm Hầm Chui Cao Tốc & Trại Heo Anh Khôi', 2, '4FO (2x2)', '4FO', '0FO', 'CSG-DNXL44 (qua CRAN DNXL63)', 'Đường gom dân sinh chân cao tốc kéo cáp 4FO về trạm CRAN DNXL63 mượn sợi quang về CSG-DNXL44.'),
    ('Route 13', 'Cụm Đường Dẫn Cao Tốc & Cổng KCN Xuân Lộc', 5, '10FO (5x2)', '12FO', '2FO', 'CSG-DNXL23 & CSG-DNXL06', 'Trục 12FO dọc đường gom cao tốc và QL1A cổng KCN về CRAN DNXL67 và CSG-DNXL06 (đã có 5G).'),
    ('Route 14', 'Vòng Xoay KCN & Tuyến Phân Trường 5', 3, '6FO (3x2)', '8FO', '2FO', 'CSG-DNXL75 & DNXL36', 'Gom Điểm 21 về CSG-DNXL75, tuyến Phân trường 5 về trạm DNXL36 (đã có LSW) bằng cáp 8FO.'),
    ('Route 15', 'Cụm Hồ Gia Ui - Đập & Xã Gia Ui', 3, '6FO (3x2)', '8FO', '2FO', 'DNXL26 & DNXL18', 'Trục 8FO khu vực đập hồ Gia Ui và làng Mán về trạm DNXL26 và trạm Macro DNXL18.')
]

for r_idx, r_data in enumerate(routes_summary_data, 4):
    ws2.row_dimensions[r_idx].height = 26
    for c_idx, val in enumerate(r_data, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        cell.font = f_data
        cell.border = border_cell
        if c_idx in [1, 3, 4, 5, 6]:
            cell.alignment = align_center
            if c_idx == 1: cell.font = f_data_bold
        else:
            cell.alignment = align_left
            
        if r_data[0] == 'Route 3':
            cell.fill = fill_route3
            if c_idx in [1, 2, 5]: cell.font = f_green
        elif 'Route 11' in r_data[0]:
            cell.fill = fill_route11

col_widths_s2 = {1: 12, 2: 38, 3: 10, 4: 14, 5: 14, 6: 14, 7: 32, 8: 52}
for col_i, w in col_widths_s2.items():
    ws2.column_dimensions[get_column_letter(col_i)].width = w

# Sheet 3: Danh sách trạm BTS liên quan
ws3 = wb.create_sheet(title="Danh_Sach_Tram_BTS_Lien_Quan")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:G1')
ws3['A1'] = "DANH SÁCH CÁC TRẠM BTS ACTIVE QUY HOẠCH KẾT NỐI CAMERA (CSG / MACRO / CRAN)"
ws3['A1'].font = f_title

headers_s3 = ['Mã Trạm Cũ', 'Mã Trạm Mới', 'Phân Loại Trạm', 'Trạm MAIN / CSG Tương Ứng', 'Trạng Thái CSG / LSW', 'Tọa Độ Trạm', 'Đặc Điểm Truyền Dẫn']
ws3.row_dimensions[3].height = 28
for c_idx, h in enumerate(headers_s3, 1):
    cell = ws3.cell(row=3, column=c_idx, value=h)
    cell.font = f_header
    cell.fill = fill_h_group1
    cell.alignment = align_center
    cell.border = border_cell

bts_highlight_list = [
    ('DNXL20', 'DNIXHO03', 'CSG Hub (Trạm 5G)', 'CSG-DNXL20', 'Đã có CSG (Nâng cấp sau 5G)', '10.8416, 107.5124', 'Trạm CSG Hub tiếp nhận Route 3 Suối Bà Rùa (Điểm 7-53-52 cáp 8FO).'),
    ('DNXL09', 'DNIXHO02', 'CSG Hub (Trạm 5G)', 'CSG-DNXL09', 'Đã có CSG (Nâng cấp sau 5G)', '10.8764, 107.4510', 'Trạm CSG Hub tiếp nhận Route 11 trục QL1A (14-15-40-17 cáp 8FO) và nhánh BQL Rừng (Điểm 23 cáp 2FO).'),
    ('DNXL06', 'DNIXHO01', 'CSG Hub (Trạm 5G)', 'CSG-DNXL06', 'Đã có CSG (Nâng cấp sau 5G)', '10.9010, 107.4260', 'Trạm CSG Hub phục vụ cụm KCN và đường dẫn cao tốc.'),
    ('DNXL49', 'DNIXHO13', 'CSG Hub (Trạm 5G)', 'CSG-DNXL49', 'Đã có CSG (Nâng cấp sau 5G)', '10.8210, 107.5580', 'Trạm CSG Hub tiếp nhận Route 2 đường Cọ Dầu ấp 3.'),
    ('DNXL24', 'DNIXHO06', 'CSG Hub (Trạm 5G)', 'CSG-DNXL24', 'Đã có CSG (Nâng cấp sau 5G)', '10.8473, 107.4948', 'Trạm CSG Hub tiếp nhận Route 9 cụm THPT Xuân Hưng và Suối Cầu.'),
    ('DNXL22', 'DNIXHO05', 'CSG Hub (Trạm 5G)', 'CSG-DNXL22', 'Đã có CSG (Nâng cấp sau 5G)', '10.8920, 107.4350', 'Trạm CSG Hub khu vực đường dẫn cao tốc.'),
    ('DNXL44', 'DNIXHO12', 'CSG Hub (Trạm MAIN)', 'CSG-DNXL44', 'Đã có CSG (A901-6CZ-FT-D)', '10.8641, 107.4817', 'Trạm Hub trung tâm CSG, kết nối quang trực tiếp.'),
    ('DNXL23', 'DNIXLO09', 'CSG Hub', 'CSG-DNXL23', 'Đã có CSG (ASR-920-12CZ-D)', '10.9189, 107.4215', 'Trạm Hub CSG phục vụ khu vực KCN.'),
    ('DNXL07', 'DNIXLO02', 'CSG Hub (Trạm 5G)', 'CSG-DNXL07', 'Đã có CSG (ASR-920-12CZ-D)', '10.9312, 107.4021', 'Trạm Hub CSG kết nối camera khu vực thị trấn.'),
    ('DNXL75', 'DNIXLO26', 'CSG Hub (Trạm 5G)', 'CSG-DNXL75', 'Đã có CSG (ASR-920-12CZ-D)', '10.9085, 107.4190', 'Trạm CSG phục vụ vòng xoay KCN Xuân Lộc.'),
    ('DNXL69', 'DNIXHO17', 'CRAN OUT DOOR', 'MAIN DNXL05', 'Không cần quang hóa', '10.8287, 107.5436', 'Đã có sẵn cáp quang về MAIN DNXL05. Mượn sợi quang hiện hữu.'),
    ('DNXL63', 'DNIXHO16', 'CRAN OUT DOOR', 'CSG-DNXL44', 'Không cần quang hóa', '10.8842, 107.5021', 'Đã có sẵn cáp quang về CSG-DNXL44. Mượn sợi quang hiện hữu.'),
    ('DNXL78', 'DNIXHO20', 'CRAN OUT DOOR', 'CSG-DNXL44', 'Không cần quang hóa', '10.8631, 107.4962', 'Đã có sẵn cáp quang về CSG-DNXL44.'),
    ('DNXL74', 'DNIXHO18', 'CRAN OUT DOOR', 'CSG-DNXL09', 'Không cần quang hóa', '10.8765, 107.4421', 'Đã có sẵn cáp quang về CSG-DNXL09.'),
    ('DNXL67', 'DNIXLO23', 'CRAN OUT DOOR', 'CSG-DNXL23', 'Không cần quang hóa', '10.9125, 107.4230', 'Đã có sẵn cáp quang về CSG-DNXL23.'),
    ('DNXL05', 'DNIXHO00', 'MACRO', 'DNXL05', 'Đã có Switch Cisco A901', '10.8340, 107.5314', 'Trạm Macro phục vụ trạm dừng chân QL1A và cụm CRAN DNXL69.'),
    ('DNXL21', 'DNIXHO04', 'MACRO', 'DNXL21', 'Trang bị thêm Switch LSW', '10.8142, 107.5710', 'Trạm Macro phục vụ QL1A giáp Bình Thuận (thay thế MBF_DNXL01 MORAN).'),
    ('DNXL36', 'DNIXHO10', 'MACRO', 'DNXL36 (Đã có LSW)', 'Đã có LSW-DNXL36 (A901)', '10.9200, 107.4350', 'Đã có Switch LSW L2 tại trạm.'),
    ('DNXL50', 'DNIXHO14', 'MACRO', 'DNXL50', 'Trang bị thêm Switch LSW', '10.8650, 107.5450', 'Trạm Macro phục vụ đường Sóc Ba Buông.'),
    ('DNXL18', 'DNIXTC02', 'MACRO', 'DNXL18', 'Trang bị thêm Switch LSW', '10.9650, 107.4890', 'Trạm Macro khu vực Gia Ui.')
]

for r_idx, b_data in enumerate(bts_highlight_list, 4):
    ws3.row_dimensions[r_idx].height = 24
    for c_idx, val in enumerate(b_data, 1):
        cell = ws3.cell(row=r_idx, column=c_idx, value=val)
        cell.font = f_data
        cell.border = border_cell
        if c_idx in [1, 2, 3]:
            cell.alignment = align_center
            if c_idx == 1: cell.font = f_data_bold
        elif c_idx == 6:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            if 'Trang bị thêm' in str(val): cell.font = f_red
            if 'Đã có CSG' in str(val) or 'CSG Hub' in str(val): cell.font = f_blue

col_widths_s3 = {1: 14, 2: 14, 3: 22, 4: 26, 5: 28, 6: 22, 7: 50}
for col_i, w in col_widths_s3.items():
    ws3.column_dimensions[get_column_letter(col_i)].width = w

wb.save(out_xlsx_path)
print(f"Updated Excel: {out_xlsx_path}")

# 5. Build Interactive Leaflet HTML Map with CSG Badges
out_html_path = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/Cam Xa Xuan Hoa/Ban_Do_Tuyen_Cap_Camera_Xuan_Loc.html"

cam_geojson_features = []
for c in processed_cams:
    if c['valid']:
        cam_geojson_features.append({
            'type': 'Feature',
            'geometry': {'type': 'Point', 'coordinates': [c['lon'], c['lat']]},
            'properties': {
                'stt': c['stt'],
                'point_id': c['point_id'],
                'name': c['name'],
                'bts_old': c['bts_old'],
                'bts_type': c['bts_type'],
                'col_h': c['col_h'],
                'col_i': c['col_i'],
                'col_j': c['col_j'],
                'road_dist': c['road_dist_m'],
                'col_k': c['col_k'],
                'col_l': c['col_l'],
                'note': c['note'],
                'is_cran': c['is_cran_tap'],
                'cable_dist': c['dist_m']
            }
        })

used_bts_set = {}
for c in processed_cams:
    if c['valid']:
        bts_id = c['bts_old']
        if bts_id not in used_bts_set:
            used_bts_set[bts_id] = {
                'id': bts_id,
                'type': c['bts_type'],
                'lat': c['bts_lat'],
                'lon': c['bts_lon'],
                'col_h': c['col_h'],
                'is_cran': c['is_cran_tap']
            }
            if c['cran_main_bts']:
                m = c['cran_main_bts']
                m_code = m['site_id_old']
                if m_code not in used_bts_set:
                    used_bts_set[m_code] = {
                        'id': m_code,
                        'type': 'CSG Hub' if m['is_csg'] else 'MACRO',
                        'lat': m['lat'],
                        'lon': m['lon'],
                        'col_h': f"CSG-{m_code}" if m['is_csg'] else m_code,
                        'is_cran': False
                    }

road_lines = []
for c in processed_cams:
    if c['valid'] and c.get('road_coords'):
        lat_lon_path = [[pt[1], pt[0]] for pt in c['road_coords']]
        road_lines.append({
            'stt': c['stt'],
            'path': lat_lon_path,
            'route': c['col_k'],
            'cable': c['col_l'],
            'dist': c['col_j'],
            'road_dist': c['road_dist_m'],
            'bts': c['bts_old'],
            'is_cran': c['is_cran_tap']
        })

cran_backhaul_lines = []
for c in processed_cams:
    if c['valid'] and c['is_cran_tap'] and c['cran_main_bts']:
        m = c['cran_main_bts']
        key = f"{c['bts_old']}->{m['site_id_old']}"
        if not any(k.get('key') == key for k in cran_backhaul_lines):
            cran_backhaul_lines.append({
                'key': key,
                'from': [c['bts_lat'], c['bts_lon']],
                'to': [m['lat'], m['lon']],
                'cran': c['bts_old'],
                'main': m['site_id_old']
            })

html_content = f"""<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Bản Đồ Tuyến Cáp Quang Camera Xuân Lộc (Hạ Tầng CSG Hậu SRAN 5G)</title>
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{ font-family: 'Inter', sans-serif; background: #0f172a; color: #f8fafc; overflow: hidden; }}
        #map {{ width: 100vw; height: 100vh; z-index: 1; }}

        .glass-panel {{
            position: absolute;
            background: rgba(15, 23, 42, 0.90);
            backdrop-filter: blur(12px);
            -webkit-backdrop-filter: blur(12px);
            border: 1px solid rgba(255, 255, 255, 0.12);
            border-radius: 12px;
            box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.37);
            z-index: 1000;
        }}

        .header-panel {{
            top: 16px;
            left: 16px;
            max-width: 440px;
            padding: 16px 20px;
        }}
        .header-title {{
            font-size: 15px;
            font-weight: 700;
            color: #38bdf8;
            display: flex;
            align-items: center;
            gap: 8px;
            margin-bottom: 6px;
        }}
        .header-sub {{
            font-size: 12px;
            color: #94a3b8;
            line-height: 1.5;
        }}

        .stats-panel {{
            bottom: 24px;
            left: 16px;
            padding: 12px 18px;
            display: flex;
            gap: 20px;
        }}
        .stat-item {{
            display: flex;
            flex-direction: column;
        }}
        .stat-value {{
            font-size: 18px;
            font-weight: 700;
            color: #38bdf8;
        }}
        .stat-label {{
            font-size: 10px;
            text-transform: uppercase;
            color: #64748b;
            letter-spacing: 0.5px;
        }}

        .filter-panel {{
            top: 16px;
            right: 16px;
            width: 330px;
            padding: 16px;
            max-height: calc(100vh - 32px);
            overflow-y: auto;
        }}
        .filter-title {{
            font-size: 13px;
            font-weight: 600;
            color: #e2e8f0;
            margin-bottom: 10px;
            display: flex;
            align-items: center;
            justify-content: space-between;
        }}
        .filter-group {{
            margin-bottom: 14px;
        }}
        .filter-label {{
            font-size: 11px;
            color: #94a3b8;
            margin-bottom: 6px;
            display: block;
        }}
        select, input {{
            width: 100%;
            background: #1e293b;
            border: 1px solid #334155;
            color: #f1f5f9;
            padding: 8px 10px;
            border-radius: 6px;
            font-size: 12px;
            outline: none;
        }}
        select:focus, input:focus {{
            border-color: #38bdf8;
        }}

        .legend-item {{
            display: flex;
            align-items: center;
            gap: 10px;
            font-size: 12px;
            margin-bottom: 6px;
            color: #cbd5e1;
        }}
        .legend-badge {{
            width: 14px;
            height: 14px;
            border-radius: 4px;
            display: inline-block;
        }}

        .cam-marker {{
            background: #0284c7;
            border: 2px solid #ffffff;
            border-radius: 50%;
            color: #fff;
            font-weight: bold;
            font-size: 10px;
            display: flex;
            align-items: center;
            justify-content: center;
            box-shadow: 0 0 10px rgba(2, 132, 199, 0.8);
        }}
        .cam-marker-route3 {{
            background: #16a34a !important;
            box-shadow: 0 0 12px rgba(22, 163, 74, 1) !important;
            border: 2px solid #fef08a !important;
        }}
        .cam-marker-route11 {{
            background: #ea580c !important;
            box-shadow: 0 0 12px rgba(234, 88, 12, 1) !important;
            border: 2px solid #fed7aa !important;
        }}
        .bts-marker {{
            width: 28px;
            height: 28px;
            border-radius: 6px;
            display: flex;
            align-items: center;
            justify-content: center;
            font-weight: 700;
            font-size: 11px;
            box-shadow: 0 0 12px rgba(0, 0, 0, 0.6);
            border: 2px solid #ffffff;
        }}
        .bts-csg {{ background: #dc2626; color: #fff; }}
        .bts-macro {{ background: #2563eb; color: #fff; }}
        .bts-cran {{ background: #9333ea; color: #fff; }}

        .leaflet-popup-content-wrapper {{
            background: #0f172a !important;
            color: #f1f5f9 !important;
            border: 1px solid rgba(255,255,255,0.15) !important;
            border-radius: 10px !important;
            box-shadow: 0 10px 25px rgba(0,0,0,0.5) !important;
        }}
        .leaflet-popup-tip {{
            background: #0f172a !important;
        }}
        .popup-title {{
            font-weight: 700;
            font-size: 14px;
            color: #38bdf8;
            margin-bottom: 6px;
            border-bottom: 1px solid #334155;
            padding-bottom: 4px;
        }}
        .popup-row {{
            font-size: 12px;
            margin-bottom: 4px;
            display: flex;
            justify-content: space-between;
        }}
        .popup-label {{ color: #94a3b8; }}
        .popup-val {{ font-weight: 600; color: #f8fafc; text-align: right; }}
        .popup-highlight {{ color: #f59e0b; font-weight: 600; }}
        .popup-blue {{ color: #38bdf8; font-weight: 700; }}
        .badge-tag {{
            display: inline-block;
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 10px;
            font-weight: 600;
            margin-top: 4px;
        }}
        .badge-2fo {{ background: #059669; color: #ecfdf5; }}
        .badge-4fo {{ background: #0891b2; color: #ecfeff; }}
        .badge-8fo {{ background: #0284c7; color: #e0f2fe; }}
        .badge-12fo {{ background: #4f46e5; color: #eef2ff; }}
        .badge-route3 {{ background: #15803d; color: #dcfce7; font-weight:700; }}
        .badge-route11 {{ background: #c2410c; color: #ffedd5; font-weight:700; }}
    </style>
</head>
<body>

    <div class="glass-panel header-panel">
        <div class="header-title">
            <svg width="20" height="20" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M9 20l-5.447-2.724A1 1 0 013 16.382V5.618a1 1 0 011.447-.894L9 7m0 13l6-3m-6 3V7m6 10l4.553 2.276A1 1 0 0021 18.382V7.618a1 1 0 00-.553-.894L15 4m0 13V4m0 0L9 7"/></svg>
            Tuyến Cáp Camera Xuân Lộc (Hạ Tầng CSG 5G)
        </div>
        <div class="header-sub">
            • <b>Cập nhật CSG hậu SRAN 5G</b>: Trạm lên 5G (DNXL20, DNXL09, DNXL06, DNXL49, DNXL24...) đều là <b>CSG Hub</b>.<br>
            • <b>Route 3 (Điểm 7-53-52)</b>: Đấu trực tiếp cổng <b>CSG-DNXL20</b> (cáp 8FO).<br>
            • <b>Cụm QL1A (14-15-40-17)</b>: Trục 8FO về <b>CSG-DNXL09</b>; Nhánh <b>Điểm 23</b> kéo 2FO riêng.
        </div>
    </div>

    <div class="glass-panel stats-panel">
        <div class="stat-item">
            <div class="stat-value">64</div>
            <div class="stat-label">Tổng Camera</div>
        </div>
        <div class="stat-item">
            <div class="stat-value">CSG-DNXL20</div>
            <div class="stat-label">Hub Route 3</div>
        </div>
        <div class="stat-item">
            <div class="stat-value">CSG-DNXL09</div>
            <div class="stat-label">Hub Route 11</div>
        </div>
        <div class="stat-item">
            <div class="stat-value">SRAN 5G</div>
            <div class="stat-label">Chuẩn Hạ Tầng</div>
        </div>
    </div>

    <div class="glass-panel filter-panel">
        <div class="filter-title">
            <span>BỘ LỌC & THIẾT KẾ TUYẾN</span>
            <button onclick="resetView()" style="background:transparent; border:none; color:#38bdf8; cursor:pointer; font-size:11px;">Reset View</button>
        </div>

        <div class="filter-group">
            <label class="filter-label">Tìm kiếm Camera (STT hoặc Tên):</label>
            <input type="text" id="camSearch" placeholder="Ví dụ: 7, 14, 20, 23, DNXL20..." onkeyup="filterCams()">
        </div>

        <div class="filter-group">
            <label class="filter-label">Lọc theo Tuyến Đường (Route):</label>
            <select id="routeSelect" onchange="filterCams()">
                <option value="all">Tất cả các Tuyến</option>
                <option value="Route 3">★ Route 3: Suối Bà Rùa -> CSG-DNXL20 (Cáp 8FO)</option>
                <option value="Route 11: Trục QL1A">★ Route 11: Trục QL1A (14-15-40-17) -> CSG-DNXL09 (8FO)</option>
                <option value="Route 11B">★ Route 11B: Nhánh BQL Rừng (Điểm 23) -> CSG-DNXL09 (2FO)</option>
                <option value="Route 1">Route 1: QL1A Đông (Cáp 8FO)</option>
                <option value="Route 2">Route 2: Đường Cọ Dầu Ấp 3 -> CSG-DNXL49 (8FO)</option>
                <option value="Route 4">Route 4: QL1A & Suối Đá (Cáp 8FO)</option>
                <option value="Route 5">Route 5: QL1A Trạm Dừng Chân (Cáp 8FO)</option>
                <option value="Route 6">Route 6: CRAN DNXL69 & Tín Nghĩa (Cáp 8FO)</option>
                <option value="Route 7">Route 7: Đường Sóc Ba Buông (Cáp 8FO)</option>
                <option value="Route 8">Route 8: Đồi Môn Tây -> CSG-DNXL44 (12FO)</option>
                <option value="Route 9">Route 9: THPT Xuân Hưng -> CSG-DNXL24 (12FO)</option>
                <option value="Route 10">Route 10: Làng Chăm -> CSG-DNXL44 (8FO)</option>
                <option value="Route 12">Route 12: Hầm Chui Cao Tốc (Cáp 4FO)</option>
                <option value="Route 13">Route 13: Cao Tốc & KCN -> CSG-DNXL06 / 23 (12FO)</option>
                <option value="Route 14">Route 14: Vòng Xoay KCN -> CSG-DNXL75 (8FO)</option>
                <option value="Route 15">Route 15: Cụm Hồ Gia Ui (Cáp 8FO)</option>
            </select>
        </div>

        <div class="filter-group">
            <label class="filter-label">Chú thích Lớp Bản Đồ:</label>
            <div class="legend-item">
                <span class="legend-badge" style="background:#16a34a; border-radius:50%;"></span>
                <span>Camera Route 3 (về CSG-DNXL20)</span>
            </div>
            <div class="legend-item">
                <span class="legend-badge" style="background:#ea580c; border-radius:50%;"></span>
                <span>Camera Route 11/11B (về CSG-DNXL09)</span>
            </div>
            <div class="legend-item">
                <span class="legend-badge bts-csg"></span>
                <span>Trạm CSG Hub (Đỏ - đã có CSG)</span>
            </div>
            <div class="legend-item">
                <span class="legend-badge bts-macro"></span>
                <span>Trạm Macro (Xanh)</span>
            </div>
            <div class="legend-item">
                <span class="legend-badge bts-cran"></span>
                <span>Trạm CRAN OUT DOOR (Tím - có quang)</span>
            </div>
        </div>
    </div>

    <div id="map"></div>

    <script>
        const camData = {json.dumps(cam_geojson_features, ensure_ascii=False)};
        const btsData = {json.dumps(used_bts_set, ensure_ascii=False)};
        const roadLines = {json.dumps(road_lines, ensure_ascii=False)};
        const backhaulLines = {json.dumps(cran_backhaul_lines, ensure_ascii=False)};

        const map = L.map('map', {{
            center: [10.855, 107.495],
            zoom: 13,
            zoomControl: false
        }});
        L.control.zoom({{ position: 'bottomright' }}).addTo(map);

        const googleSat = L.tileLayer('https://mt1.google.com/vt/lyrs=y&x={{x}}&y={{y}}&z={{z}}', {{
            maxZoom: 20,
            attribution: 'Google Hybrid'
        }}).addTo(map);

        const osm = L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{
            maxZoom: 19,
            attribution: 'OpenStreetMap'
        }});

        const cartoDark = L.tileLayer('https://{{s}}.basemaps.cartocdn.com/dark_all/{{z}}/{{x}}/{{y}}{{r}}.png', {{
            maxZoom: 19,
            attribution: 'CartoDB Dark'
        }});

        const baseMaps = {{
            "Vệ tinh Google (Hybrid)": googleSat,
            "Bản đồ Giao thông (OSM)": osm,
            "Bản đồ Dark Mode": cartoDark
        }};
        L.control.layers(baseMaps, null, {{ position: 'topleft' }}).addTo(map);

        const camLayerGroup = L.layerGroup().addTo(map);
        const btsLayerGroup = L.layerGroup().addTo(map);
        const cableLayerGroup = L.layerGroup().addTo(map);
        const backhaulLayerGroup = L.layerGroup().addTo(map);

        Object.values(btsData).forEach(b => {{
            let cls = 'bts-macro';
            let label = 'M';
            if (b.type.includes('CSG')) {{ cls = 'bts-csg'; label = 'CSG'; }}
            else if (b.is_cran) {{ cls = 'bts-cran'; label = 'CR'; }}

            const icon = L.divIcon({{
                className: 'custom-bts-marker',
                html: `<div class="bts-marker ${{cls}}">${{label}}</div>`,
                iconSize: [28, 28],
                iconAnchor: [14, 14]
            }});

            const marker = L.marker([b.lat, b.lon], {{ icon }}).addTo(btsLayerGroup);
            marker.bindPopup(`
                <div class="popup-title">Trạm BTS: ${{b.id}}</div>
                <div class="popup-row"><span class="popup-label">Phân loại hạ tầng:</span><span class="popup-val popup-blue">${{b.type}}</span></div>
                <div class="popup-row"><span class="popup-label">Tọa độ:</span><span class="popup-val">${{b.lat.toFixed(5)}}, ${{b.lon.toFixed(5)}}</span></div>
                <div class="popup-row"><span class="popup-label">Cột H tương ứng:</span><span class="popup-val popup-highlight">${{b.col_h}}</span></div>
            `);
        }});

        backhaulLines.forEach(l => {{
            L.polyline([l.from, l.to], {{
                color: '#c084fc',
                weight: 3.5,
                dashArray: '6, 8',
                opacity: 0.85
            }}).addTo(backhaulLayerGroup).bindTooltip(`Tuyến quang hiện hữu: ${{l.cran}} -> ${{l.main}}`);
        }});

        const camMarkers = [];
        const cablePolylines = [];

        camData.forEach(f => {{
            const p = f.properties;
            const coords = [f.geometry.coordinates[1], f.geometry.coordinates[0]];
            const isRoute3 = ['7', '53', '52'].includes(p.stt);
            const isRoute11 = ['14', '15', '40', '17', '23'].includes(p.stt);

            let markerCls = 'cam-marker';
            if (isRoute3) markerCls = 'cam-marker cam-marker-route3';
            else if (isRoute11) markerCls = 'cam-marker cam-marker-route11';

            const icon = L.divIcon({{
                className: 'custom-cam-marker',
                html: `<div class="${{markerCls}}" style="width:24px; height:24px;">${{p.stt}}</div>`,
                iconSize: [24, 24],
                iconAnchor: [12, 12]
            }});

            let badgeClass = 'badge-8fo';
            if (p.col_l.includes('2FO')) badgeClass = 'badge-2fo';
            else if (p.col_l.includes('4FO')) badgeClass = 'badge-4fo';
            else if (p.col_l.includes('12FO')) badgeClass = 'badge-12fo';
            if (isRoute3) badgeClass = 'badge-route3';
            if (isRoute11) badgeClass = 'badge-route11';

            const marker = L.marker(coords, {{ icon }}).addTo(camLayerGroup);
            marker.camProps = p;
            marker.bindPopup(`
                <div class="popup-title">STT ${{p.stt}}: ${{p.point_id}}</div>
                <div class="popup-row"><span class="popup-label">Vị trí:</span><span class="popup-val">${{p.name}}</span></div>
                <div class="popup-row"><span class="popup-label">Tọa độ:</span><span class="popup-val">${{coords[0].toFixed(5)}}, ${{coords[1].toFixed(5)}}</span></div>
                <div class="popup-row"><span class="popup-label">Trạm BTS kết nối:</span><span class="popup-val popup-blue">${{p.bts_old}} (${{p.bts_type}})</span></div>
                <div class="popup-row"><span class="popup-label">Đường chạy xe máy:</span><span class="popup-val popup-highlight">${{p.road_dist}} m</span></div>
                <div class="popup-row"><span class="popup-label">Cáp dự toán (Cột J):</span><span class="popup-val popup-highlight">${{p.col_j}} m</span></div>
                <div class="popup-row"><span class="popup-label">CỘT H (Trạm Đích):</span><span class="popup-val popup-blue">${{p.col_h}}</span></div>
                <div class="popup-row"><span class="popup-label">Cột I (Port):</span><span class="popup-val">${{p.col_i}}</span></div>
                <div class="popup-row"><span class="popup-label">Route Gom Tuyến:</span><span class="popup-val">${{p.col_k}}</span></div>
                <div class="popup-row"><span class="popup-label">Loại cáp tối ưu:</span><span class="popup-val badge-tag ${{badgeClass}}">${{p.col_l}}</span></div>
                <div style="margin-top:6px; font-size:11px; color:#cbd5e1; border-top:1px dashed #334155; padding-top:4px;">
                    <b>Ghi chú:</b> ${{p.note}}
                </div>
            `);
            camMarkers.push(marker);
        }});

        roadLines.forEach(l => {{
            const isRoute3 = l.route.includes('Route 3');
            const isRoute11 = l.route.includes('Route 11');
            
            let lineColor = '#38bdf8';
            let lineWeight = 3;
            if (isRoute3) {{ lineColor = '#22c55e'; lineWeight = 4.5; }}
            else if (isRoute11) {{ lineColor = '#f97316'; lineWeight = 4; }}
            else if (l.cable.includes('12FO')) {{ lineColor = '#818cf8'; }}
            else if (l.cable.includes('4FO')) {{ lineColor = '#06b6d4'; }}

            const poly = L.polyline(l.path, {{
                color: lineColor,
                weight: lineWeight,
                opacity: 0.90
            }}).addTo(cableLayerGroup);
            poly.camStt = l.stt;
            poly.route = l.route;
            cablePolylines.push(poly);
        }});

        function resetView() {{
            map.setView([10.855, 107.495], 13);
            document.getElementById('camSearch').value = '';
            document.getElementById('routeSelect').value = 'all';
            filterCams();
        }}

        function filterCams() {{
            const query = document.getElementById('camSearch').value.toLowerCase().trim();
            const selectedRoute = document.getElementById('routeSelect').value;

            camMarkers.forEach(m => {{
                const p = m.camProps;
                let matchQuery = true;
                if (query) {{
                    matchQuery = p.stt.toString().includes(query) || 
                                 p.point_id.toLowerCase().includes(query) || 
                                 p.name.toLowerCase().includes(query) ||
                                 p.bts_old.toLowerCase().includes(query) ||
                                 p.col_h.toLowerCase().includes(query);
                }}

                let matchRoute = true;
                if (selectedRoute !== 'all') {{
                    matchRoute = p.col_k.includes(selectedRoute);
                }}

                if (matchQuery && matchRoute) {{
                    camLayerGroup.addLayer(m);
                }} else {{
                    camLayerGroup.removeLayer(m);
                }}
            }});

            cablePolylines.forEach(poly => {{
                let matchRoute = true;
                if (selectedRoute !== 'all') {{
                    matchRoute = poly.route.includes(selectedRoute);
                }}
                if (matchRoute) {{
                    cableLayerGroup.addLayer(poly);
                }} else {{
                    cableLayerGroup.removeLayer(poly);
                }}
            }});
        }}
    </script>
</body>
</html>
"""

with open(out_html_path, 'w', encoding='utf-8') as f:
    f.write(html_content)
print(f"Updated Interactive HTML Map: {out_html_path}")

print("COMPLETED FULL POST-SRAN 5G UPDATE SUCCESSFULLY!")
