#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script implementing Option 1 for Dong Nai 67 Tram:
- Removes 1 generator run of station DNXL20 on 2026-09-04 (Row 32 in 02A):
  Amount: 504,028 VND (16.45 Liters, 4.62 hours).
- Updates 02A_TTNB_DongNai_67Tram:
  Total runs: 101 (9 Gas + 92 Oil).
  Oil total: 28,656,372 VND | Gas total: 1,176,939 VND | Grand Total: 29,833,311 VND.
  F8 = 29,833,311 VND.
- Updates HD_DongNai_67Tram:
  Keeps exactly 14 invoices (13 Oil + 1 Gas):
  All 12 oil invoices from 14/09 to 29/09 + exactly 1 oil invoice on 30/09 (705060: 3,900,433 VND).
  Invoice 704928 (1,500,000 VND on 30/09) is moved to HD_Du_Thua_Khong_Su_Dung.
  -> Every single day in Dong Nai is now <= 3,900,433 VND (ALL <= 5,000,000 VND 100%!).
- Re-runs Waterfall allocation for Map_HD_Theo_Tram_Nhom1:
  Column B = Column D = 29,833,311 VND (EXACT 100% MATCH).
  Surplus preserved: Oil +4,185 VND, Gas +34,561 VND.
- Updates HD_Du_Thua_Khong_Su_Dung:
  38 invoices total (7 August + 20 DN surplus + 11 TC surplus).
- Preserves 100% Toan Cau sheets.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import defaultdict
import shutil
import os

HSTT_PATH = '/Users/cang_it/Desktop/Ho_So_Thanh_Toan_Chuan_Mau_09_2026_HSTT.xlsx'
DOWNLOADS_PATH = '/Users/cang_it/Downloads/Ho_So_Thanh_Toan_Chuan_Mau_09_2026_HSTT.xlsx'

# Formatting styles
FONT_NAME = 'Times New Roman'
FONT_MAIN = Font(name=FONT_NAME, size=11)
FONT_BOLD = Font(name=FONT_NAME, size=11, bold=True)
FONT_ITALIC = Font(name=FONT_NAME, size=10, italic=True)
FONT_SECTION = Font(name=FONT_NAME, size=11, bold=True, color='1F4E78')
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color='1F4E78')

FILL_HEADER = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
FILL_SECTION = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
FILL_TOTAL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
FILL_NOTE = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

THIN_SIDE = Side(border_style='thin', color='A6A6A6')
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
DOUBLE_BOTTOM = Side(border_style='double', color='000000')
BORDER_TOTAL = Border(top=THIN_SIDE, bottom=DOUBLE_BOTTOM, left=THIN_SIDE, right=THIN_SIDE)

ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

def clear_sheet_from_row(ws, from_row):
    merged_ranges = list(ws.merged_cells.ranges)
    for mr in merged_ranges:
        if mr.min_row >= from_row or mr.max_row >= from_row:
            ws.unmerge_cells(str(mr))
    if ws.max_row >= from_row:
        ws.delete_rows(from_row, ws.max_row - from_row + 1)

def main():
    print(">>> 1. Loading workbook...")
    wb = openpyxl.load_workbook(HSTT_PATH)

    # -------------------------------------------------------------
    # STEP 1: MODIFY 02A_TTNB_DongNai_67Tram
    # Remove Row 32 (DNXL20 on 2026-09-04: 504,028 VND)
    # -------------------------------------------------------------
    print(">>> 2. Removing target run at Row 32 in 02A_TTNB_DongNai_67Tram...")
    ws_02a1 = wb['02A_TTNB_DongNai_67Tram']
    
    # Confirm target row
    target_row = 32
    tid_check = ws_02a1.cell(target_row, 3).value
    amt_check = ws_02a1.cell(target_row, 13).value
    date_check = str(ws_02a1.cell(target_row, 7).value)[:10]
    print(f"Target check: Row {target_row} | TID={tid_check} | Date={date_check} | Amt={amt_check}")
    assert tid_check == 'DNXL20' and date_check == '2026-09-04' and amt_check == 504028, "Mismatch on target row!"

    ws_02a1.delete_rows(target_row, 1)

    # Re-number Oil runs: rows 26 to 117 (92 runs)
    oil_end_row = 117
    for idx, r in enumerate(range(26, oil_end_row + 1), 1):
        ws_02a1.cell(r, 1, idx)

    # Row 118: Tổng cộng máy chạy dầu
    row_tot_oil = 118
    ws_02a1.cell(row_tot_oil, 10, f"=SUM(J26:J{oil_end_row})")
    ws_02a1.cell(row_tot_oil, 11, f"=SUM(K26:K{oil_end_row})")
    ws_02a1.cell(row_tot_oil, 13, f"=SUM(M26:M{oil_end_row})")

    # Row 119: TỔNG CỘNG (XĂNG + DẦU)
    row_grand = 119
    ws_02a1.cell(row_grand, 10, f"=J22+J{row_tot_oil}")
    ws_02a1.cell(row_grand, 11, f"=K22+K{row_tot_oil}")
    ws_02a1.cell(row_grand, 13, f"=M22+M{row_tot_oil}")

    # Row 8: F8 = 29833311
    NEW_G1_TOTAL = 29833311
    ws_02a1['F8'] = NEW_G1_TOTAL
    ws_02a1['I8'] = 0
    ws_02a1['K8'] = "=F8+I8"
    ws_02a1['F9'] = "=F8"
    ws_02a1['I9'] = "=I8"
    ws_02a1['K9'] = "=K8"

    print(f"Updated 02A G1: 92 oil runs, grand total formula in M{row_grand}, F8={NEW_G1_TOTAL}")

    # -------------------------------------------------------------
    # STEP 2: RE-READ DEMANDS FROM UPDATED 02A SHEET
    # -------------------------------------------------------------
    print(">>> 3. Reading updated demands for Group 1...")
    site_lookup = {}
    if os.path.exists('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json'):
        import json
        with open('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json') as f:
            for s in json.load(f):
                site_lookup[s.get('site_id')] = s.get('site_name')

    g1_x_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(13, 22):
        tid = str(ws_02a1.cell(r, 3).value).strip()
        g1_x_demands[tid]['tien'] += float(ws_02a1.cell(r, 13).value or 0)
        g1_x_demands[tid]['lit'] += float(ws_02a1.cell(r, 11).value or 0)
        g1_x_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a1.cell(r, 2).value or '').strip()

    g1_d_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(26, oil_end_row + 1):
        tid = str(ws_02a1.cell(r, 3).value).strip()
        g1_d_demands[tid]['tien'] += float(ws_02a1.cell(r, 13).value or 0)
        g1_d_demands[tid]['lit'] += float(ws_02a1.cell(r, 11).value or 0)
        g1_d_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a1.cell(r, 2).value or '').strip()

    tot_d_demand = sum(g1_d_demands[t]['tien'] for t in g1_d_demands)
    tot_x_demand = sum(g1_x_demands[t]['tien'] for t in g1_x_demands)
    print(f"Updated G1 demands: Oil={tot_d_demand:,.0f} đ, Gas={tot_x_demand:,.0f} đ, Total={tot_d_demand+tot_x_demand:,.0f} đ")
    assert tot_d_demand == 28656372, f"Expected 28,656,372 but got {tot_d_demand}"
    assert tot_d_demand + tot_x_demand == NEW_G1_TOTAL, "Mismatch on new G1 total!"

    # -------------------------------------------------------------
    # STEP 3: SELECT INVOICES FOR DONG NAI (OPTION 1)
    # -------------------------------------------------------------
    # In HD_DongNai_67Tram currently:
    # 14 oil invoices (665105 to 705060) + 1 gas invoice (656300).
    # We REMOVE 704928 (1,500,000 VND on 30/09) from active list and send it to surplus!
    # The active oil invoices will be 13 invoices:
    # 665105, 666711, 667931, 671634, 676989, 682090, 687125, 690332, 692839, 694889, 701215, 703454, 705060.
    # Total oil money = 24,760,124 + 3,900,433 = 28,660,557 VND >= 28,656,372 VND!
    # Active gas invoice = 656300 (1,211,500 VND).
    
    # Read current active invoices from HD_DongNai_67Tram
    ws_hd1 = wb['HD_DongNai_67Tram']
    current_hd1 = []
    inv_704928 = None
    for r in range(8, ws_hd1.max_row):
        val3 = ws_hd1.cell(r, 3).value
        if val3 and str(val3).isdigit() and len(str(val3)) >= 6:
            inv_info = {
                'stt': ws_hd1.cell(r, 1).value,
                'ngay': str(ws_hd1.cell(r, 2).value)[:10],
                'so_hd': str(val3).strip(),
                'don_vi_ban': ws_hd1.cell(r, 4).value,
                'loai_nl': str(ws_hd1.cell(r, 5).value or '').strip(),
                'dien_giai': ws_hd1.cell(r, 6).value,
                'lit': float(ws_hd1.cell(r, 7).value or 0),
                'mst_ban': str(ws_hd1.cell(r, 8).value or '').strip(),
                'mau_so': ws_hd1.cell(r, 9).value,
                'ky_hieu': ws_hd1.cell(r, 10).value,
                'tien_chua_vat': float(ws_hd1.cell(r, 11).value or 0),
                'vat': float(ws_hd1.cell(r, 12).value or 0),
                'tong_tien': float(ws_hd1.cell(r, 13).value or 0),
                'link': ws_hd1.cell(r, 14).value,
                'fkey': ws_hd1.cell(r, 15).value,
            }
            if inv_info['so_hd'] == '704928':
                inv_704928 = inv_info
            else:
                current_hd1.append(inv_info)

    assert inv_704928 is not None, "Invoice 704928 not found in HD_DongNai_67Tram!"
    print(f"Removed 704928 from active list: Date={inv_704928['ngay']}, Amt={inv_704928['tong_tien']:,.0f} đ")

    # Sort remaining active invoices: Oil first (by date, so_hd), then Gas
    active_oil_dn = sorted([x for x in current_hd1 if 'Dầu' in x['loai_nl']], key=lambda x: (x['ngay'], x['so_hd']))
    active_gas_dn = sorted([x for x in current_hd1 if 'Xăng' in x['loai_nl']], key=lambda x: (x['ngay'], x['so_hd']))
    print(f"Active DN Invoices: Oil={len(active_oil_dn)} HĐ (Sum={sum(x['tong_tien'] for x in active_oil_dn):,.0f} đ), Gas={len(active_gas_dn)} HĐ")

    # -------------------------------------------------------------
    # STEP 4: RUN WATERFALL ALLOCATION FOR GROUP 1
    # -------------------------------------------------------------
    def run_waterfall(demands_dict, invoices_list):
        sorted_tids = sorted(demands_dict.keys())
        assignments = []
        inv_idx = 0
        inv_rem = invoices_list[0]['tong_tien'] if invoices_list else 0

        for tid in sorted_tids:
            needed = demands_dict[tid]['tien']
            is_first = True
            while needed > 0 and inv_idx < len(invoices_list):
                cur_inv = invoices_list[inv_idx]
                if inv_rem >= needed:
                    assignments.append({
                        'tid': tid,
                        'st_name': demands_dict[tid]['ten'],
                        'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                        'so_hd': cur_inv['so_hd'],
                        'tien_gan': needed,
                        'inv': cur_inv
                    })
                    inv_rem -= needed
                    needed = 0
                    is_first = False
                else:
                    assignments.append({
                        'tid': tid,
                        'st_name': demands_dict[tid]['ten'],
                        'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                        'so_hd': cur_inv['so_hd'],
                        'tien_gan': inv_rem,
                        'inv': cur_inv
                    })
                    needed -= inv_rem
                    is_first = False
                    inv_idx += 1
                    if inv_idx < len(invoices_list):
                        inv_rem = invoices_list[inv_idx]['tong_tien']

        tot_tram = sum(demands_dict[t]['tien'] for t in demands_dict)
        tot_gan = sum(a['tien_gan'] for a in assignments)
        return assignments, tot_tram, tot_gan, inv_rem

    a1_d, t1_d, g1_d, r1_d = run_waterfall(g1_d_demands, active_oil_dn)
    a1_x, t1_x, g1_x, r1_x = run_waterfall(g1_x_demands, active_gas_dn)

    print(f"G1 Waterfall Oil: Tram={t1_d:,.0f} đ | Gan={g1_d:,.0f} đ | Rem={r1_d:,.0f} đ")
    print(f"G1 Waterfall Gas: Tram={t1_x:,.0f} đ | Gan={g1_x:,.0f} đ | Rem={r1_x:,.0f} đ")
    print(f"G1 Total: Tram={t1_d+t1_x:,.0f} đ | Gan={g1_d+g1_x:,.0f} đ | Target={NEW_G1_TOTAL:,.0f} đ -> EXACT MATCH 100%!")
    assert t1_d == g1_d and t1_x == g1_x, "Waterfall mismatch!"

    # -------------------------------------------------------------
    # STEP 5: RE-WRITE Map_HD_Theo_Tram_Nhom1
    # -------------------------------------------------------------
    print(">>> 5. Rewriting Map_HD_Theo_Tram_Nhom1...")
    ws_map1 = wb['Map_HD_Theo_Tram_Nhom1']
    clear_sheet_from_row(ws_map1, 1)

    ws_map1.cell(1, 1, "BẢNG KÊ PHÂN BỔ HÓA ĐƠN XĂNG DẦU THEO TỪNG TRẠM CHẠY MÁY (MOBIFONE ĐỒNG NAI - 67 TRẠM ĐẶC THÙ)").font = FONT_TITLE
    ws_map1.cell(2, 1, "Tháng 09/2026 • Phân tách độc lập bảng kê Dầu DO và Xăng RON 95 • Ưu tiên đáp ứng đủ 100% số tiền bảng kê, bảo lưu số lít/tiền dư").font = FONT_ITALIC
    ws_map1.row_dimensions[1].height = 28
    ws_map1.row_dimensions[2].height = 20

    headers = [
        'ID trạm\n(Nhãn Hàng)', 'Thành tiền chạy máy\ntheo trạm (đồng)', 'Số hóa đơn',
        'Số tiền gán từ HĐ\n(đồng)', 'Đơn vị bán hàng', 'Mã số thuế\n(Bán)',
        'Ký hiệu HĐ', 'Link tra cứu hóa đơn', 'Mã tra cứu / Fkey', 'Ngày HĐ', 'Tên trạm / Địa bàn'
    ]
    ws_map1.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers, 1):
        cell = ws_map1.cell(4, col_idx, h)
        cell.font = FONT_BOLD
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL

    # Section I: Dầu DO
    curr_row = 5
    sec1 = ws_map1.cell(curr_row, 1, "I. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU DẦU DO (DO 0.05S) — GÁN HÓA ĐƠN THEO TRẠM")
    sec1.font = FONT_SECTION; sec1.fill = FILL_SECTION
    ws_map1.row_dimensions[curr_row].height = 24
    ws_map1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=11)
    for c in range(1, 12):
        ws_map1.cell(curr_row, c).border = BORDER_ALL; ws_map1.cell(curr_row, c).fill = FILL_SECTION

    dau_start = curr_row + 1
    curr_row += 1
    for a in a1_d:
        ws_map1.row_dimensions[curr_row].height = 20
        c1 = ws_map1.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map1.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map1.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map1.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map1.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT
        for c in range(5, 12):
            ws_map1.cell(curr_row, c).font = FONT_MAIN
            ws_map1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    dau_end = curr_row - 1
    # Total Dầu DO
    tot_dau_row = curr_row
    ws_map1.row_dimensions[tot_dau_row].height = 22
    ws_map1.cell(tot_dau_row, 1, "TỔNG CỘNG DẦU DO").font = FONT_BOLD
    ws_map1.cell(tot_dau_row, 1).alignment = ALIGN_LEFT

    t2 = ws_map1.cell(tot_dau_row, 2, f"=SUM(B{dau_start}:B{dau_end})")
    t2.font = FONT_BOLD; t2.alignment = ALIGN_RIGHT; t2.number_format = '#,##0'

    t4 = ws_map1.cell(tot_dau_row, 4, f"=SUM(D{dau_start}:D{dau_end})")
    t4.font = FONT_BOLD; t4.alignment = ALIGN_RIGHT; t4.number_format = '#,##0'

    for c in range(1, 12):
        ws_map1.cell(tot_dau_row, c).fill = FILL_TOTAL
        ws_map1.cell(tot_dau_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Dầu DO
    note_dau_row = curr_row
    ws_map1.row_dimensions[note_dau_row].height = 22
    tot_dau_lit_buy = sum(inv['lit'] for inv in active_oil_dn)
    tot_dau_tien_buy = sum(inv['tong_tien'] for inv in active_oil_dn)
    tot_dau_lit_used = sum(g1_d_demands[t]['lit'] for t in g1_d_demands)
    note_dau_text = f"📌 Ghi chú bảo lưu Dầu DO: Tổng HĐ mua {tot_dau_lit_buy:,.1f} L ({tot_dau_tien_buy:,.0f} đ) — Tiêu hao chạy máy {tot_dau_lit_used:,.1f} L ({t1_d:,.0f} đ) ➔ Tiền HĐ còn dư bảo lưu kho: +{r1_d:,.0f} đ"
    n_cell = ws_map1.cell(note_dau_row, 1, note_dau_text)
    n_cell.font = FONT_ITALIC; n_cell.fill = FILL_NOTE
    ws_map1.merge_cells(start_row=note_dau_row, start_column=1, end_row=note_dau_row, end_column=11)
    for c in range(1, 12):
        ws_map1.cell(note_dau_row, c).border = BORDER_ALL; ws_map1.cell(note_dau_row, c).fill = FILL_NOTE

    curr_row += 2
    # Section II: Xăng RON 95
    sec2 = ws_map1.cell(curr_row, 1, "II. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU XĂNG RON 95 (RON 95-III) — GÁN HÓA ĐƠN THEO TRẠM")
    sec2.font = FONT_SECTION; sec2.fill = FILL_SECTION
    ws_map1.row_dimensions[curr_row].height = 24
    ws_map1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=11)
    for c in range(1, 12):
        ws_map1.cell(curr_row, c).border = BORDER_ALL; ws_map1.cell(curr_row, c).fill = FILL_SECTION

    xang_start = curr_row + 1
    curr_row += 1
    for a in a1_x:
        ws_map1.row_dimensions[curr_row].height = 20
        c1 = ws_map1.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map1.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map1.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map1.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map1.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT
        for c in range(5, 12):
            ws_map1.cell(curr_row, c).font = FONT_MAIN
            ws_map1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    xang_end = curr_row - 1
    # Total Xăng
    tot_xang_row = curr_row
    ws_map1.row_dimensions[tot_xang_row].height = 22
    ws_map1.cell(tot_xang_row, 1, "TỔNG CỘNG XĂNG RON 95").font = FONT_BOLD
    ws_map1.cell(tot_xang_row, 1).alignment = ALIGN_LEFT

    t2_x = ws_map1.cell(tot_xang_row, 2, f"=SUM(B{xang_start}:B{xang_end})")
    t2_x.font = FONT_BOLD; t2_x.alignment = ALIGN_RIGHT; t2_x.number_format = '#,##0'

    t4_x = ws_map1.cell(tot_xang_row, 4, f"=SUM(D{xang_start}:D{xang_end})")
    t4_x.font = FONT_BOLD; t4_x.alignment = ALIGN_RIGHT; t4_x.number_format = '#,##0'

    for c in range(1, 12):
        ws_map1.cell(tot_xang_row, c).fill = FILL_TOTAL
        ws_map1.cell(tot_xang_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Xăng
    note_xang_row = curr_row
    ws_map1.row_dimensions[note_xang_row].height = 22
    tot_xang_lit_buy = sum(inv['lit'] for inv in active_gas_dn)
    tot_xang_tien_buy = sum(inv['tong_tien'] for inv in active_gas_dn)
    tot_xang_lit_used = sum(g1_x_demands[t]['lit'] for t in g1_x_demands)
    surplus_lit_xang = tot_xang_lit_buy - tot_xang_lit_used
    note_xang_text = f"📌 Ghi chú bảo lưu Xăng RON 95: Tổng HĐ mua {tot_xang_lit_buy:,.1f} L ({tot_xang_tien_buy:,.0f} đ) — Tiêu hao chạy máy {tot_xang_lit_used:,.1f} L ({t1_x:,.0f} đ) ➔ Số lít dư bảo lưu kho: +{surplus_lit_xang:,.1f} L (Tiền HĐ còn dư bảo lưu kho: +{r1_x:,.0f} đ)"
    n_cell_x = ws_map1.cell(note_xang_row, 1, note_xang_text)
    n_cell_x.font = FONT_ITALIC; n_cell_x.fill = FILL_NOTE
    ws_map1.merge_cells(start_row=note_xang_row, start_column=1, end_row=note_xang_row, end_column=11)
    for c in range(1, 12):
        ws_map1.cell(note_xang_row, c).border = BORDER_ALL; ws_map1.cell(note_xang_row, c).fill = FILL_NOTE

    curr_row += 2
    # Grand Total
    grand_total_row = curr_row
    ws_map1.row_dimensions[grand_total_row].height = 24
    ws_map1.cell(grand_total_row, 1, "TỔNG CỘNG TOÀN BỘ (DẦU DO + XĂNG RON 95)").font = FONT_BOLD
    ws_map1.cell(grand_total_row, 1).alignment = ALIGN_LEFT

    gt2 = ws_map1.cell(grand_total_row, 2, f"=B{tot_dau_row}+B{tot_xang_row}")
    gt2.font = FONT_BOLD; gt2.alignment = ALIGN_RIGHT; gt2.number_format = '#,##0'

    gt4 = ws_map1.cell(grand_total_row, 4, f"=D{tot_dau_row}+D{tot_xang_row}")
    gt4.font = FONT_BOLD; gt4.alignment = ALIGN_RIGHT; gt4.number_format = '#,##0'

    for c in range(1, 12):
        ws_map1.cell(grand_total_row, c).fill = FILL_TOTAL
        ws_map1.cell(grand_total_row, c).border = BORDER_TOTAL

    widths = {1: 14, 2: 24, 3: 15, 4: 24, 5: 35, 6: 16, 7: 14, 8: 30, 9: 18, 10: 14, 11: 25}
    for col_idx, w in widths.items():
        ws_map1.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # -------------------------------------------------------------
    # STEP 6: RE-WRITE HD_DongNai_67Tram (14 INVOICES)
    # -------------------------------------------------------------
    print(">>> 6. Rewriting HD_DongNai_67Tram (14 invoices)...")
    clear_sheet_from_row(ws_hd1, 8)
    final_hd1_list = active_oil_dn + active_gas_dn
    curr_row = 8
    for idx, inv in enumerate(final_hd1_list, 1):
        ws_hd1.row_dimensions[curr_row].height = 20
        ws_hd1.cell(curr_row, 1, f"L{idx}").alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 2, inv['ngay']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 3, inv['so_hd']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 4, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_hd1.cell(curr_row, 5, inv['loai_nl']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 6, inv['dien_giai']).alignment = ALIGN_LEFT

        c7 = ws_hd1.cell(curr_row, 7, inv['lit'])
        c7.alignment = ALIGN_RIGHT; c7.number_format = '#,##0.00'

        ws_hd1.cell(curr_row, 8, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 9, inv['mau_so']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 10, inv['ky_hieu']).alignment = ALIGN_CENTER

        c11 = ws_hd1.cell(curr_row, 11, inv['tien_chua_vat'])
        c11.alignment = ALIGN_RIGHT; c11.number_format = '#,##0'

        c12 = ws_hd1.cell(curr_row, 12, inv['vat'])
        c12.alignment = ALIGN_RIGHT; c12.number_format = '#,##0'

        c13 = ws_hd1.cell(curr_row, 13, inv['tong_tien'])
        c13.alignment = ALIGN_RIGHT; c13.number_format = '#,##0'

        ws_hd1.cell(curr_row, 14, inv['link']).alignment = ALIGN_LEFT
        ws_hd1.cell(curr_row, 15, inv['fkey']).alignment = ALIGN_CENTER

        for c in range(1, 16):
            ws_hd1.cell(curr_row, c).font = FONT_MAIN
            ws_hd1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    end_inv_row = curr_row - 1
    tot_row = curr_row
    ws_hd1.row_dimensions[tot_row].height = 24
    ws_hd1.cell(tot_row, 1, "TỔNG CỘNG TOÀN BỘ HÓA ĐƠN").font = FONT_BOLD
    ws_hd1.cell(tot_row, 1).alignment = ALIGN_LEFT
    ws_hd1.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)

    c7_tot = ws_hd1.cell(tot_row, 7, f"=SUM(G8:G{end_inv_row})")
    c7_tot.font = FONT_BOLD; c7_tot.alignment = ALIGN_RIGHT; c7_tot.number_format = '#,##0.00'

    c11_tot = ws_hd1.cell(tot_row, 11, f"=SUM(K8:K{end_inv_row})")
    c11_tot.font = FONT_BOLD; c11_tot.alignment = ALIGN_RIGHT; c11_tot.number_format = '#,##0'

    c12_tot = ws_hd1.cell(tot_row, 12, f"=SUM(L8:L{end_inv_row})")
    c12_tot.font = FONT_BOLD; c12_tot.alignment = ALIGN_RIGHT; c12_tot.number_format = '#,##0'

    c13_tot = ws_hd1.cell(tot_row, 13, f"=SUM(M8:M{end_inv_row})")
    c13_tot.font = FONT_BOLD; c13_tot.alignment = ALIGN_RIGHT; c13_tot.number_format = '#,##0'

    for c in range(1, 16):
        ws_hd1.cell(tot_row, c).fill = FILL_TOTAL
        ws_hd1.cell(tot_row, c).border = BORDER_TOTAL

    # -------------------------------------------------------------
    # STEP 7: RE-WRITE HD_Du_Thua_Khong_Su_Dung (38 INVOICES)
    # -------------------------------------------------------------
    print(">>> 7. Adding 704928 to HD_Du_Thua_Khong_Su_Dung...")
    ws_surplus = wb['HD_Du_Thua_Khong_Su_Dung']
    
    # Read existing 37 surplus invoices
    surplus_list = []
    for r in range(5, ws_surplus.max_row):
        val3 = ws_surplus.cell(r, 3).value
        if val3 and str(val3).isdigit() and len(str(val3)) >= 6:
            surplus_list.append({
                'ngay': str(ws_surplus.cell(r, 2).value)[:10],
                'so_hd': str(val3).strip(),
                'buyer': ws_surplus.cell(r, 4).value,
                'loai_nl': ws_surplus.cell(r, 5).value,
                'lit': float(ws_surplus.cell(r, 6).value or 0),
                'don_gia': float(ws_surplus.cell(r, 7).value or 0),
                'tien_chua_vat': float(ws_surplus.cell(r, 8).value or 0),
                'vat': float(ws_surplus.cell(r, 9).value or 0),
                'tong_tien': float(ws_surplus.cell(r, 10).value or 0),
                'don_vi_ban': ws_surplus.cell(r, 11).value,
                'mst_ban': ws_surplus.cell(r, 12).value,
                'ky_hieu': ws_surplus.cell(r, 13).value,
                'link': ws_surplus.cell(r, 14).value,
                'fkey': ws_surplus.cell(r, 15).value,
                'ghi_chu': ws_surplus.cell(r, 16).value,
            })

    # Add 704928
    surplus_list.append({
        'ngay': inv_704928['ngay'],
        'so_hd': inv_704928['so_hd'],
        'buyer': 'MobiFone Đồng Nai (0100686209-129)',
        'loai_nl': inv_704928['loai_nl'],
        'lit': inv_704928['lit'],
        'don_gia': round(inv_704928['tong_tien'] / inv_704928['lit']) if inv_704928['lit'] > 0 else 0,
        'tien_chua_vat': inv_704928['tien_chua_vat'],
        'vat': inv_704928['vat'],
        'tong_tien': inv_704928['tong_tien'],
        'don_vi_ban': inv_704928['don_vi_ban'],
        'mst_ban': inv_704928['mst_ban'],
        'ky_hieu': inv_704928['ky_hieu'],
        'link': inv_704928['link'],
        'fkey': inv_704928['fkey'],
        'ghi_chu': 'Hóa đơn Dầu DO tháng 09/2026 dư thừa (đã đủ 100% tiền trạm chạy dầu Đồng Nai, giúp ngày 30/09 <= 5 triệu)'
    })

    # Sort surplus: August first, then September by date, so_hd
    surplus_list.sort(key=lambda x: (0 if '-08-' in x['ngay'] else 1, x['ngay'], x['so_hd']))
    print(f"Total surplus invoices: {len(surplus_list)}")

    clear_sheet_from_row(ws_surplus, 1)

    ws_surplus.cell(1, 1, "BẢNG KÊ HÓA ĐƠN XĂNG DẦU DƯ THỪA / BẢO LƯU KHO KHÔNG SỬ DỤNG").font = FONT_TITLE
    ws_surplus.cell(2, 1, f"Tháng 09/2026 • Tổng cộng {len(surplus_list)} hóa đơn được bảo lưu trong kho dữ liệu (gồm 7 HĐ tháng 8 gối đầu + {len(surplus_list)-7} HĐ tháng 9 dư thừa)").font = FONT_ITALIC
    ws_surplus.row_dimensions[1].height = 28
    ws_surplus.row_dimensions[2].height = 20

    headers_surplus = [
        "STT", "Ngày Lập HĐ", "Số Hóa Đơn", "Bên Mua (Pháp Nhân / MST)", "Loại Nhiên Liệu",
        "Số Lượng (Lít)", "Đơn Giá (đ/L)", "Thành Tiền Chưa Thuế (đ)", "Thuế GTGT 8% (đ)",
        "Tổng Tiền Thanh Toán (đ)", "Đơn Vị Xuất Hóa Đơn", "Mã Số Thuế (Bán)",
        "Ký Hiệu HĐ", "Link Tra Cứu Hóa Đơn", "Mã Tra Cứu / Fkey", "Lý Do Dư Thừa / Ghi Chú Phân Loại"
    ]
    ws_surplus.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers_surplus, 1):
        cell = ws_surplus.cell(4, col_idx, h)
        cell.font = FONT_BOLD; cell.fill = FILL_HEADER; cell.alignment = ALIGN_CENTER; cell.border = BORDER_ALL

    curr_row = 5
    for idx, inv in enumerate(surplus_list, 1):
        ws_surplus.row_dimensions[curr_row].height = 20
        ws_surplus.cell(curr_row, 1, idx).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 2, inv['ngay']).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 3, inv['so_hd']).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 4, inv['buyer']).alignment = ALIGN_LEFT
        ws_surplus.cell(curr_row, 5, inv['loai_nl']).alignment = ALIGN_CENTER

        c6 = ws_surplus.cell(curr_row, 6, inv['lit'])
        c6.alignment = ALIGN_RIGHT; c6.number_format = '#,##0.00'

        c7 = ws_surplus.cell(curr_row, 7, inv['don_gia'])
        c7.alignment = ALIGN_RIGHT; c7.number_format = '#,##0'

        c8 = ws_surplus.cell(curr_row, 8, inv['tien_chua_vat'])
        c8.alignment = ALIGN_RIGHT; c8.number_format = '#,##0'

        c9 = ws_surplus.cell(curr_row, 9, inv['vat'])
        c9.alignment = ALIGN_RIGHT; c9.number_format = '#,##0'

        c10 = ws_surplus.cell(curr_row, 10, inv['tong_tien'])
        c10.alignment = ALIGN_RIGHT; c10.number_format = '#,##0'

        ws_surplus.cell(curr_row, 11, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_surplus.cell(curr_row, 12, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 13, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 14, inv['link']).alignment = ALIGN_LEFT
        ws_surplus.cell(curr_row, 15, inv['fkey']).alignment = ALIGN_CENTER
        ws_surplus.cell(curr_row, 16, inv['ghi_chu']).alignment = ALIGN_LEFT

        for c in range(1, 17):
            ws_surplus.cell(curr_row, c).font = FONT_MAIN
            ws_surplus.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    end_surplus_row = curr_row - 1
    tot_surplus_row = curr_row
    ws_surplus.row_dimensions[tot_surplus_row].height = 24
    ws_surplus.cell(tot_surplus_row, 1, "TỔNG CỘNG HÓA ĐƠN DƯ THỪA").font = FONT_BOLD
    ws_surplus.cell(tot_surplus_row, 1).alignment = ALIGN_LEFT
    ws_surplus.merge_cells(start_row=tot_surplus_row, start_column=1, end_row=tot_surplus_row, end_column=5)

    c6_s = ws_surplus.cell(tot_surplus_row, 6, f"=SUM(F5:F{end_surplus_row})")
    c6_s.font = FONT_BOLD; c6_s.alignment = ALIGN_RIGHT; c6_s.number_format = '#,##0.00'

    c8_s = ws_surplus.cell(tot_surplus_row, 8, f"=SUM(H5:H{end_surplus_row})")
    c8_s.font = FONT_BOLD; c8_s.alignment = ALIGN_RIGHT; c8_s.number_format = '#,##0'

    c9_s = ws_surplus.cell(tot_surplus_row, 9, f"=SUM(I5:I{end_surplus_row})")
    c9_s.font = FONT_BOLD; c9_s.alignment = ALIGN_RIGHT; c9_s.number_format = '#,##0'

    c10_s = ws_surplus.cell(tot_surplus_row, 10, f"=SUM(J5:J{end_surplus_row})")
    c10_s.font = FONT_BOLD; c10_s.alignment = ALIGN_RIGHT; c10_s.number_format = '#,##0'

    for c in range(1, 17):
        ws_surplus.cell(tot_surplus_row, c).fill = FILL_TOTAL
        ws_surplus.cell(tot_surplus_row, c).border = BORDER_TOTAL

    widths_surplus = {1: 8, 2: 14, 3: 14, 4: 30, 5: 16, 6: 16, 7: 14, 8: 22, 9: 16, 10: 22, 11: 35, 12: 16, 13: 14, 14: 30, 15: 18, 16: 45}
    for col_idx, w in widths_surplus.items():
        ws_surplus.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # -------------------------------------------------------------
    # STEP 8: SAVE TO DESKTOP AND DOWNLOADS
    # -------------------------------------------------------------
    print(">>> 8. Saving workbook...")
    wb.save(HSTT_PATH)
    shutil.copy2(HSTT_PATH, DOWNLOADS_PATH)
    print(f"SUCCESS: Saved {HSTT_PATH} and copied to {DOWNLOADS_PATH}")

if __name__ == '__main__':
    main()
