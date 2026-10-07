import os
import math
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel_a4():
    print("🚀 Đang khởi tạo Báo cáo Rà soát Gói 2, 3, 4 định dạng in A4...")

    # 1. Đọc dữ liệu từ Supabase
    env = {}
    with open('tvt3_v2/.env') as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith('#') and '=' in line:
                k, v = line.split('=', 1)
                env[k.strip()] = v.strip().strip('"').strip("'")

    url = env.get('VITE_SUPABASE_URL')
    key = env.get('VITE_SUPABASE_ANON_KEY')
    headers_sb = {'apikey': key, 'Authorization': f'Bearer {key}'}

    res = requests.get(f'{url}/rest/v1/infrastructure_projects?select=*&limit=1000', headers=headers_sb)
    sb_projects = {p['planning_id_new']: p for p in res.json()}

    def haversine(lat1, lon1, lat2, lon2):
        try:
            lat1, lon1, lat2, lon2 = float(lat1), float(lon1), float(lat2), float(lon2)
            R = 6371000.0
            phi1 = math.radians(lat1)
            phi2 = math.radians(lat2)
            delta_phi = math.radians(lat2 - lat1)
            delta_lambda = math.radians(lon2 - lon1)
            a = math.sin(delta_phi / 2)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(delta_lambda / 2)**2
            c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
            return round(R * c, 1)
        except Exception:
            return None

    # 2. Tạo Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ra_Soat_Goi_2_3_4"

    # Thiết lập in ấn khổ A4 ngang, vừa khít 1 trang chiều ngang
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Căn lề in hẹp tối ưu diện tích
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2

    # Lặp lại tiêu đề cột (Row 5) khi sang trang
    ws.print_title_rows = '5:5'
    ws.print_options.gridLines = True
    ws.sheet_view.showGridLines = True

    # Header / Footer trang in
    ws.oddHeader.left.text = "&B&8TỔ VIỄN THÔNG 3 - PHÒNG VIỄN THÔNG"
    ws.oddHeader.right.text = "&B&8BÁO CÁO RÀ SOÁT CSHT GÓI 2, 3, 4 (A4 LANDSCAPE)"
    ws.oddFooter.left.text = "&I&8Hệ thống Quản lý VHKT & CSHT - MobiFone Đồng Nai (10/2026)"
    ws.oddFooter.right.text = "&B&8Trang &P / &N"

    # 3. Định dạng Style
    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_sub = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_g2 = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    fill_g3 = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    fill_g4 = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    fill_htcs = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    fill_green_badge = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_amber_badge = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_rose_badge = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")

    font_main_title = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    font_sub_title = Font(name="Arial", size=9, italic=True, color="475569")
    font_kpi_num = Font(name="Arial", size=10.5, bold=True, color="1E3A8A")
    
    font_header = Font(name="Arial", size=8.5, bold=True, color="FFFFFF")
    font_group = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    
    font_cell = Font(name="Arial", size=8)
    font_cell_bold = Font(name="Arial", size=8, bold=True)
    font_cell_green = Font(name="Arial", size=8, bold=True, color="15803D")
    font_cell_amber = Font(name="Arial", size=8, bold=True, color="B45309")
    font_cell_rose = Font(name="Arial", size=8, bold=True, color="BE123C")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # 4. Ghi Header chính
    ws.merge_cells("A1:P1")
    ws["A1"] = "BÁO CÁO RÀ SOÁT CHI TIẾT DỮ LIỆU & ĐIỀU KIỆN TRÌNH KÝ HỢP ĐỒNG CÁC GÓI 2, 3, 4"
    ws["A1"].font = font_main_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:P2")
    ws["A2"] = "Đối soát Tọa độ Khảo sát vs Tọa độ Sở KH&CN Phê duyệt lại (100% Chấp thuận đầu tư mới) | Chuẩn hóa in khổ A4 Landscape"
    ws["A2"].font = font_sub_title
    ws["A2"].alignment = align_center

    # 5. Khối tóm tắt KPI (Dòng 3 - 4)
    ws.merge_cells("A3:C4")
    ws["A3"] = "TỔNG SỐ TRẠM\n20 Trạm MBF Đầu Tư"
    ws["A3"].font = font_kpi_num
    ws["A3"].fill = fill_g2
    ws["A3"].alignment = align_center

    ws.merge_cells("D3:F4")
    ws["D3"] = "ĐỦ ĐIỀU KIỆN TRÌNH HĐ\n🟢 11 Trạm (Đủ 100% hồ sơ)"
    ws["D3"].font = font_kpi_num
    ws["D3"].fill = fill_green_badge
    ws["D3"].alignment = align_center

    ws.merge_cells("G3:J4")
    ws["G3"] = "ĐÃ CÓ CHỦ - CHỜ LẤY STK\n🟡 7 Trạm (Đã có Thửa/Tờ, SĐT)"
    ws["G3"].font = font_kpi_num
    ws["G3"].fill = fill_amber_badge
    ws["G3"].alignment = align_center

    ws.merge_cells("K3:M4")
    ws["K3"] = "CHƯA CÓ CHỦ ĐẤT\n🔴 2 Trạm (26DNa250, 185)"
    ws["K3"].font = font_kpi_num
    ws["K3"].fill = fill_rose_badge
    ws["K3"].alignment = align_center

    ws.merge_cells("N3:P4")
    ws["N3"] = "PHÊ DUYỆT SỞ KH&CN\n100% Chấp Thuận Xây Mới"
    ws["N3"].font = font_kpi_num
    ws["N3"].fill = fill_g3
    ws["N3"].alignment = align_center

    for r in range(3, 5):
        for c in range(1, 17):
            ws.cell(row=r, column=c).border = thin_border

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 17
    ws.row_dimensions[4].height = 17

    # 6. Dòng tiêu đề cột chính (Dòng 5)
    headers = [
        "STT", "Gói", "Mã QH\nMới", "Mã QH\nCũ", "Huyện /\nXã Phường",
        "Tọa độ KS\n(Vĩ độ, Kinh độ)", "Tọa độ SKHCN duyệt\n(Vĩ độ, Kinh độ)", "Độ lệch\n(m)",
        "Họ Tên Chủ Đất", "Số Điện Thoại", "Số CCCD", "Thửa /\nTờ BĐ",
        "Diện Tích\n& Giá Thuê", "Tài Khoản Ngân Hàng", "Cột & Độ Cao\n(SKHCN Duyệt)", "Tình Trạng Hồ Sơ &\nViệc Cần Làm"
    ]
    
    header_row = 5
    ws.row_dimensions[header_row].height = 30

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_text)
        cell.fill = fill_navy
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border

    # 7. Danh sách dữ liệu các trạm
    sites_list = [
        # GÓI 2
        ('Gói 2', '26DNa175', '21DNTP200'),
        ('Gói 2', '26DNa187', '25DNTP002'),
        ('Gói 2', '26DNa242', '26DNI417'),
        ('Gói 2', '26DNa244', '26DNI425'),
        ('Gói 2', '26DNa250', '0'),
        # GÓI 3
        ('Gói 3', '26DNa162', '26DNI527'),
        ('Gói 3', '26DNa164', '21DNLK105'),
        ('Gói 3', '26DNa166', '25DNXL004'),
        ('Gói 3', '26DNa168', '26DNI605'),
        ('Gói 3', '26DNa245', '22DNI307'),
        ('Gói 3', '26DNa247', '26DNI415'),
        ('Gói 3', '26DNa258', '26DNa258'),
        # GÓI 4
        ('Gói 4', '26DNa179', '26DNa129 / 26DNI239'),
        ('Gói 4', '26DNa158', '26DNI385'),
        ('Gói 4', '26DNa163', 'DNI21031_3'),
        ('Gói 4', '26DNa165', '22DNLK005'),
        ('Gói 4', '26DNa167', '26DNI399'),
        ('Gói 4', '26DNa181', '21DNDQ002'),
        ('Gói 4', '26DNa185', '26DNI523'),
        ('Gói 4', '26DNa255', '26DNI601'),
        # TRẠM ĐẶC BIỆT CHUYỂN DÙNG CHUNG
        ('Chuyển HTCS', '26DNa246', '26DNI442')
    ]

    curr_row = 6
    last_pkg = None

    for idx, (pkg, sid, old_id) in enumerate(sites_list):
        sb = sb_projects.get(sid, {})
        
        # Nếu đổi sang nhóm gói mới, chèn tiêu đề nhóm
        if pkg != last_pkg:
            last_pkg = pkg
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=16)
            grp_cell = ws.cell(row=curr_row, column=1)
            if pkg == 'Gói 2':
                grp_cell.value = "🔹 KHỐI GÓI 2 — MOBIFONE TỰ ĐẦU TƯ (05 TRẠM) — SKHCN CHẤP THUẬN XÂY MỚI (VB 2965/SKHCN-CĐS)"
                grp_cell.fill = fill_blue_sub
            elif pkg == 'Gói 3':
                grp_cell.value = "🔹 KHỐI GÓI 3 — MOBIFONE TỰ ĐẦU TƯ (07 TRẠM) — SKHCN CHẤP THUẬN XÂY MỚI (VB 4370/SKHCN-CĐS)"
                grp_cell.fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
            elif pkg == 'Gói 4':
                grp_cell.value = "🔹 KHỐI GÓI 4 — DỰ KIẾN TRÌNH HỢP ĐỒNG (08 TRẠM) — SỞ KH&CN ĐÃ DUYỆT LẠI ĐẦU TƯ MỚI (MBF ĐẦU TƯ)"
                grp_cell.fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
            else:
                grp_cell.value = "⚠️ TRẠM ĐẶC BIỆT — CHUYỂN PHƯƠNG ÁN DÙNG CHUNG CSHT (HTCS) DO KHÔNG CÓ SỔ ĐỎ ĐẤT"
                grp_cell.fill = PatternFill(start_color="64748B", end_color="64748B", fill_type="solid")
            
            grp_cell.font = font_group
            grp_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[curr_row].height = 20
            for c in range(1, 17):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        huyen = sb.get('district') or ''
        xa = sb.get('ward') or ''
        huyen_xa = f"{huyen}\n{xa}".strip()

        lat_ks = sb.get('latitude_survey')
        lng_ks = sb.get('longitude_survey')
        coord_ks = f"{lat_ks:.5f}\n{lng_ks:.5f}" if lat_ks and lng_ks else "Chưa có"

        lat_sk = sb.get('latitude_skhcn') or sb.get('latitude_plan')
        lng_sk = sb.get('longitude_skhcn') or sb.get('longitude_plan')
        coord_sk = f"{float(lat_sk):.5f}\n{float(lng_sk):.5f}" if lat_sk and lng_sk else "Chưa có"

        dist = haversine(lat_ks, lng_ks, lat_sk, lng_sk) if lat_ks and lng_ks and lat_sk and lng_sk else None
        dist_str = f"{dist:.1f}" if dist is not None else "-"

        owner = sb.get('landowner_name') or "---"
        phone = sb.get('landlord_phone') or "---"
        cccd = sb.get('landlord_cccd') or "---"
        
        plot = sb.get('plot_number') or "-"
        sheet = sb.get('map_sheet') or "-"
        thua_to = f"T:{plot}\nBĐ:{sheet}" if (plot != "-" or sheet != "-") else "---"

        dt = sb.get('leased_area')
        dt_str = f"{dt} m²" if dt else "---"
        rent = sb.get('proposed_rent')
        rent_str = f"{int(rent):,}đ".replace(',', '.') if rent else "---"
        dt_rent = f"{dt_str}\n{rent_str}"

        bank = sb.get('bank_account')
        bname = sb.get('bank_name')
        stk_str = f"{bank}\n({bname})" if bank else ("Chưa có STK" if owner != "---" else "---")

        pole = f"{sb.get('antenna_type_survey') or 'Dây co'}\n{sb.get('antenna_height_survey') or '42'}m"

        # Đánh giá điều kiện trình hợp đồng
        is_ok = bool(owner != "---" and phone != "---" and bank and rent)
        if pkg == 'Chuyển HTCS':
            status_hd = "🔄 Chuyển HTCS"
            action = "Không có sổ đỏ, chuyển sang dùng chung CSHT"
        elif is_ok:
            status_hd = "🟢 Đủ ĐK Trình HĐ"
            action = "Sẵn sàng xuất tờ trình & ký HĐ mặt bằng"
        elif owner != "---" and not bank:
            status_hd = "🟡 Thiếu STK / DT"
            action = "Cần liên hệ chủ đất bổ sung STK & Tên NH"
        else:
            status_hd = "🔴 Thiếu Chủ Đất"
            action = "Cần khảo sát tìm chủ & giấy tờ thửa đất"

        row_vals = [
            idx + 1 if pkg != 'Chuyển HTCS' else "-",
            pkg,
            sid,
            old_id,
            huyen_xa,
            coord_ks,
            coord_sk,
            dist_str,
            owner,
            phone,
            cccd,
            thua_to,
            dt_rent,
            stk_str,
            pole,
            f"{status_hd}\n{action}"
        ]

        ws.row_dimensions[curr_row].height = 28

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.border = thin_border
            cell.font = font_cell

            # Căn lề
            if c_idx in [1, 2, 3, 4, 8, 12, 15]:
                cell.alignment = align_center
            elif c_idx in [6, 7]:
                cell.alignment = align_center
            elif c_idx in [9, 16]:
                cell.alignment = align_left
            elif c_idx in [10, 11, 13, 14]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            # Tô màu nhẹ theo gói
            if pkg == 'Gói 2':
                cell.fill = fill_g2
            elif pkg == 'Gói 3':
                cell.fill = fill_g3
            elif pkg == 'Gói 4':
                cell.fill = fill_g4
            else:
                cell.fill = fill_htcs

        # Format đặc biệt cho cột Mã mới và Tình trạng
        ws.cell(row=curr_row, column=3).font = font_cell_bold
        
        # Format màu cho trạng thái HĐ
        status_cell = ws.cell(row=curr_row, column=16)
        if "Đủ ĐK" in status_hd:
            status_cell.font = font_cell_green
        elif "Thiếu STK" in status_hd:
            status_cell.font = font_cell_amber
        elif "Thiếu Chủ" in status_hd:
            status_cell.font = font_cell_rose

        # Format màu độ lệch
        dist_cell = ws.cell(row=curr_row, column=8)
        if dist is not None:
            if dist < 50:
                dist_cell.font = font_cell_green
            elif dist < 350:
                dist_cell.font = font_cell_amber
            else:
                dist_cell.font = font_cell_rose

        curr_row += 1

    # 8. Căn chỉnh độ rộng các cột tối ưu chính xác cho A4 Landscape
    col_widths = {
        'A': 5,   # STT
        'B': 8,   # Gói
        'C': 11,  # Mã mới
        'D': 11,  # Mã cũ
        'E': 14,  # Huyện/Xã
        'F': 12,  # Tọa độ KS
        'G': 12,  # Tọa độ SKHCN
        'H': 8,   # Độ lệch (m)
        'I': 18,  # Chủ đất
        'J': 13,  # SĐT
        'K': 14,  # CCCD
        'L': 9,   # Thửa/Tờ
        'M': 11,  # DT & Giá
        'N': 16,  # STK
        'O': 11,  # Cột/Độ cao
        'P': 26   # Tình trạng & Việc cần làm
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Dòng chữ ký xác nhận
    curr_row += 1
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=16)
    sign_cell = ws.cell(row=curr_row, column=1)
    sign_cell.value = "Đồng Nai, ngày 06 tháng 10 năm 2026 — Đơn vị lập báo cáo: Tổ Kỹ Thuật Viễn Thông 3 (TVT3)"
    sign_cell.font = Font(name="Arial", size=8.5, italic=True, color="475569")
    sign_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[curr_row].height = 20

    # 9. Lưu các file đích
    destinations = [
        '/Users/cang_it/Desktop/Bao_Cao_Ra_Soat_Goi_2_3_4_A4.xlsx',
        '/Users/cang_it/Downloads/Bao_Cao_Ra_Soat_Goi_2_3_4_A4.xlsx',
        'tvt3_v2/public/reports/Bao_Cao_Ra_Soat_Goi_2_3_4_A4.xlsx'
    ]

    for p in destinations:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        print(f"  🎯 Đã xuất thành công: {p} ({os.path.getsize(p):,} bytes)")

    print("🎉 HOÀN THÀNH XUẤT FILE EXCEL IN KHỔ A4 CHUẨN ĐẸP!")

if __name__ == '__main__':
    build_excel_a4()
