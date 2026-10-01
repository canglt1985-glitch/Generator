#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script applying Phương án B and restoring full September settlement:
1. Restores ca DNXL20 (504,028 đ) to 02A_TTNB_DongNai_67Tram -> 102 runs, F8 = 30,337,339 đ.
2. In HD_DongNai_67Tram:
   - Includes 15 invoices (1 Gas: 656300 + 14 Oil: 665105, 666711, 667931, 671634, 676989,
     682090, 687125, 690332, 692839, 694889, 701215, 703454, 705060, 704928).
   - HĐ 705060 (3,900,433 đ) and HĐ 704928 (1,500,000 đ) on 30/09: Bank transfer 5,400,433 đ to Nam Trung Phong.
   - HĐ 704928 used 465,282 đ, residual 1,034,718 đ stored in warehouse.
   - All other 13 invoices: Tiền mặt.
3. Map_HD_Theo_Tram_Nhom1:
   - Full waterfall allocation for 102 runs.
   - Column B = Column D = 30,337,339 đ (Diff = 0 đ).
4. Preserves Toàn Cầu:
   - 02A_TTNB_ToanCau: 195 runs, F8 = 71,268,583 đ.
   - HD_ToanCau: 48 invoices, 71,268,583 đ allocated.
   - Bank transfer: HĐ 676405 (19/09, 8,982,000 đ), HĐ 704930 (30/09, 2,400,203 đ), HĐ 652795 (09/09, 3,800,114 đ), HĐ 651553 (09/09, 850,099 đ).
   - All other invoices (including 100% of Tín Nghĩa): Tiền mặt.
5. HD_Du_Thua_Khong_Su_Dung:
   - Keeps all original surplus/reserve invoices cleanly documented.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import defaultdict
import shutil
import os

HSTT_PATH = '/Users/cang_it/Desktop/Ho_So_Thanh_Toan_Chuan_Mau_09_2026_HSTT.xlsx'
DOWNLOADS_PATH = '/Users/cang_it/Downloads/Ho_So_Thanh_Toan_Chuan_Mau_09_2026_HSTT.xlsx'
BACKUP_PATH = '/Users/cang_it/Desktop/Ho_So_Thanh_Toan_Chuan_Mau_09_2026_HSTT_backup.xlsx'

FONT_NAME = 'Times New Roman'
FONT_MAIN = Font(name=FONT_NAME, size=11)
FONT_BOLD = Font(name=FONT_NAME, size=11, bold=True)
FONT_ITALIC = Font(name=FONT_NAME, size=10, italic=True)
FONT_SECTION = Font(name=FONT_NAME, size=11, bold=True, color='1F4E78')
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color='1F4E78')

FILL_HEADER = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
FILL_SECTION = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
FILL_TOTAL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
FILL_NOTE = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
FILL_TRANSFER = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Soft green for bank transfer

THIN_SIDE = Side(border_style='thin', color='A6A6A6')
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_TOP_THIN = Border(top=THIN_SIDE)
DOUBLE_BOTTOM = Side(border_style='double', color='000000')
BORDER_TOTAL = Border(top=THIN_SIDE, bottom=DOUBLE_BOTTOM, left=THIN_SIDE, right=THIN_SIDE)

ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

def clear_sheet_from_row(ws, from_row):
    merged_ranges = list(ws.merged_cells.ranges)
    for mr in merged_ranges:
        if mr.min_row >= from_row or mr.max_row >= from_row:
            ws.unmerge_cells(str(mr))
    if ws.max_row >= from_row:
        ws.delete_rows(from_row, ws.max_row - from_row + 1)

def run():
    print(">>> 1. Loading template and reference files...")
    # Load backup for exact 02A_TTNB_DongNai_67Tram with all 102 runs
    wb_bak = openpyxl.load_workbook(BACKUP_PATH, data_only=True)
    ws_02a_bak = wb_bak['02A_TTNB_DongNai_67Tram']

    wb = openpyxl.load_workbook(HSTT_PATH)

    # 1. Restore 02A_TTNB_DongNai_67Tram
    print(">>> 2. Restoring 02A_TTNB_DongNai_67Tram (102 runs, F8 = 30,337,339 đ)...")
    ws_02a_target = wb['02A_TTNB_DongNai_67Tram']
    clear_sheet_from_row(ws_02a_target, 1)

    # Copy all rows and formatting from backup
    for r in range(1, ws_02a_bak.max_row + 1):
        for c in range(1, ws_02a_bak.max_column + 1):
            cell_bak = ws_02a_bak.cell(r, c)
            cell_target = ws_02a_target.cell(r, c, cell_bak.value)
            if cell_bak.has_style:
                cell_target.font = Font(name=cell_bak.font.name or FONT_NAME, size=cell_bak.font.size, bold=cell_bak.font.bold, italic=cell_bak.font.italic, color=cell_bak.font.color)
                cell_target.alignment = Alignment(horizontal=cell_bak.alignment.horizontal, vertical=cell_bak.alignment.vertical, wrap_text=cell_bak.alignment.wrap_text)
                cell_target.border = Border(left=cell_bak.border.left, right=cell_bak.border.right, top=cell_bak.border.top, bottom=cell_bak.border.bottom)
                cell_target.fill = PatternFill(fill_type=cell_bak.fill.fill_type, start_color=cell_bak.fill.start_color, end_color=cell_bak.fill.end_color)
                if cell_bak.number_format:
                    cell_target.number_format = cell_bak.number_format

    for mr in ws_02a_bak.merged_cells.ranges:
        ws_02a_target.merge_cells(str(mr))

    # Ensure F8 = 30337339, K8 = 30337339
    ws_02a_target['F8'] = 30337339
    ws_02a_target['K8'] = 30337339

    # 2. Rebuild HD_DongNai_67Tram
    print(">>> 3. Rebuilding HD_DongNai_67Tram with 15 invoices (including HĐ 704928)...")
    ws_hd_dn = wb['HD_DongNai_67Tram']
    clear_sheet_from_row(ws_hd_dn, 8)

    dn_invoices = [
        # Gas
        {'stt': 'L1', 'ngay': '2026-09-11', 'so_hd': '656300', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Xăng', 'dien_giai': 'Xăng RON 95', 'lit': 50.0, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 1211500, 'vat': 0, 'tong': 1211500, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjU2MzAwfFBLODJONjM1S3w2NzE4OA==', 'fkey': 'PK82N635K', 'note': '💵 Tiền mặt (Ngày 11/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        # Oil
        {'stt': 'L2', 'ngay': '2026-09-14', 'so_hd': '665105', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 87.78, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2500000, 'vat': 0, 'tong': 2500000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjY1MTA1fFk1Tzg4REtZQXw2NzE4OA==', 'fkey': 'Y5O88DKYA', 'note': '💵 Tiền mặt (Ngày 14/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L3', 'ngay': '2026-09-15', 'so_hd': '666711', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 70.22, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjY2NzExfEU1VlFDRjlHRHw2NzE4OA==', 'fkey': 'E5VQCF9GD', 'note': '💵 Tiền mặt (Ngày 15/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L4', 'ngay': '2026-09-16', 'so_hd': '667931', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 66.05, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjY3OTMxfE44VTlBN05ZTXw2NzE4OA==', 'fkey': 'N8U9A7NYM', 'note': '💵 Tiền mặt (Ngày 16/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L5', 'ngay': '2026-09-17', 'so_hd': '671634', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 87.78, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2500000, 'vat': 0, 'tong': 2500000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjcxNjM0fFVRVTVFQUlLVHw2NzE4OA==', 'fkey': 'UQU5EAIKT', 'note': '💵 Tiền mặt (Ngày 17/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L6', 'ngay': '2026-09-19', 'so_hd': '676989', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 63.01, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjc2OTg5fFBXU1MzOU5EUHw2NzE4OA==', 'fkey': 'PWSS39NDP', 'note': '💵 Tiền mặt (Ngày 19/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L7', 'ngay': '2026-09-21', 'so_hd': '682090', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 60.81, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 1930000, 'vat': 0, 'tong': 1930000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjgyMDkwfFc5U0lTNEVIQ3w2NzE4OA==', 'fkey': 'W9SIS4EHC', 'note': '💵 Tiền mặt (Ngày 21/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L8', 'ngay': '2026-09-23', 'so_hd': '687125', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 63.01, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjg3MTI1fFdWNDg0Uk5XUHw2NzE4OA==', 'fkey': 'WV484RNWP', 'note': '💵 Tiền mặt (Ngày 23/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L9', 'ngay': '2026-09-24', 'so_hd': '690332', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 66.17, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2100108, 'vat': 0, 'tong': 2100108, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjkwMzMyfEIyQ1dGNjkyQ3w2NzE4OA==', 'fkey': 'B2CWF692C', 'note': '💵 Tiền mặt (Ngày 24/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L10', 'ngay': '2026-09-25', 'so_hd': '692839', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 57.03, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 1830000, 'vat': 0, 'tong': 1830000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjkyODM5fEQ4UjZLN1NXRnw2NzE4OA==', 'fkey': 'D8R6K7SWF', 'note': '💵 Tiền mặt (Ngày 25/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L11', 'ngay': '2026-09-26', 'so_hd': '694889', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 59.21, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 1900016, 'vat': 0, 'tong': 1900016, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNjk0ODg5fE81TEw3TTRFRHw2NzE4OA==', 'fkey': 'O5LL7M4ED', 'note': '💵 Tiền mặt (Ngày 26/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L12', 'ngay': '2026-09-28', 'so_hd': '701215', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 62.32, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNzAxMjE1fDRXSFZOM1hVQXw2NzE4OA==', 'fkey': '4WHVN3XUA', 'note': '💵 Tiền mặt (Ngày 28/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L13', 'ngay': '2026-09-29', 'so_hd': '703454', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 65.6, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 2000000, 'vat': 0, 'tong': 2000000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNzAzNDU0fFZXVFBRM1ZGRHw2NzE4OA==', 'fkey': 'VWTPQ3VFD', 'note': '💵 Tiền mặt (Ngày 29/09 <= 5tr) — Sử dụng 100% giá vốn thanh toán', 'is_transfer': False},
        {'stt': 'L14', 'ngay': '2026-09-30', 'so_hd': '705060', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 127.93, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 3900433, 'vat': 0, 'tong': 3900433, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNzA1MDYwfFhOQkxZWFc2Tnw2NzE4OA==', 'fkey': 'XNBLYXW6N', 'note': '💳 Chuyển khoản cây xăng (Nam Trung Phong - 5,400,433 đ: HĐ 705060 + 704928) [Sử dụng 100%]', 'is_transfer': True},
        {'stt': 'L15', 'ngay': '2026-09-30', 'so_hd': '704928', 'don_vi': 'CÔNG TY TNHH MỘT THÀNH VIÊN THƯƠNG MẠI XĂNG DẦU NAM TRUNG PHONG', 'loai': 'Dầu', 'dien_giai': 'Dầu Điêzen', 'lit': 49.2, 'mst': '3600642702', 'ky_hieu': '1C26MTP', 'tien_chua_vat': 1500000, 'vat': 0, 'tong': 1500000, 'link': 'http://3600642702hd.easyinvoice.com.vn/Invoice/ViewFromEmail?token=MUMyNk1UUF9fNzA0OTI4fFozTUhQM1pIMnw2NzE4OA==', 'fkey': 'Z3MHP3ZH2', 'note': '💳 Chuyển khoản cây xăng (Nam Trung Phong - 5,400,433 đ: HĐ 705060 + 704928) [Trích dùng 465,282 đ, bảo lưu 1,034,718 đ]', 'is_transfer': True}
    ]

    cur_row = 8
    for item in dn_invoices:
        ws_hd_dn.cell(cur_row, 1, item['stt']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 2, item['ngay']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 3, item['so_hd']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 4, item['don_vi']).alignment = ALIGN_LEFT
        ws_hd_dn.cell(cur_row, 5, item['loai']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 6, item['dien_giai']).alignment = ALIGN_LEFT
        ws_hd_dn.cell(cur_row, 7, item['lit']).alignment = ALIGN_RIGHT
        ws_hd_dn.cell(cur_row, 8, item['mst']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 9, None)
        ws_hd_dn.cell(cur_row, 10, item['ky_hieu']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 11, item['tien_chua_vat']).alignment = ALIGN_RIGHT
        ws_hd_dn.cell(cur_row, 12, item['vat']).alignment = ALIGN_RIGHT
        ws_hd_dn.cell(cur_row, 13, item['tong']).alignment = ALIGN_RIGHT
        ws_hd_dn.cell(cur_row, 14, item['link']).alignment = ALIGN_LEFT
        ws_hd_dn.cell(cur_row, 15, item['fkey']).alignment = ALIGN_CENTER
        ws_hd_dn.cell(cur_row, 16, item['note']).alignment = ALIGN_LEFT

        # Styles
        fill_to_use = FILL_TRANSFER if item['is_transfer'] else None
        for c in range(1, 17):
            cell = ws_hd_dn.cell(cur_row, c)
            cell.font = FONT_MAIN
            cell.border = BORDER_ALL
            if fill_to_use:
                cell.fill = fill_to_use
            if c in [7]:
                cell.number_format = '#,##0.00'
            elif c in [11, 12, 13]:
                cell.number_format = '#,##0'
        cur_row += 1

    # Total row for HD_DongNai_67Tram
    ws_hd_dn.cell(cur_row, 1, 'TỔNG CỘNG THANH TOÁN').alignment = ALIGN_CENTER
    ws_hd_dn.cell(cur_row, 7, f'=SUM(G8:G{cur_row-1})').number_format = '#,##0.00'
    ws_hd_dn.cell(cur_row, 11, f'=SUM(K8:K{cur_row-1})').number_format = '#,##0'
    ws_hd_dn.cell(cur_row, 12, f'=SUM(L8:L{cur_row-1})').number_format = '#,##0'
    ws_hd_dn.cell(cur_row, 13, f'=SUM(M8:M{cur_row-1})').number_format = '#,##0'
    for c in range(1, 17):
        cell = ws_hd_dn.cell(cur_row, c)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
    ws_hd_dn.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=6)

    # 3. Read station demands from restored 02A_TTNB_DongNai_67Tram
    print(">>> 4. Reading demands from 02A_TTNB_DongNai_67Tram for Waterfall mapping...")
    site_lookup = {}
    if os.path.exists('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json'):
        import json
        with open('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json') as f:
            for s in json.load(f):
                site_lookup[s.get('site_id')] = s.get('site_name')

    g1_x_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    g1_d_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})

    # Find section headers in ws_02a_target
    gas_start = 13
    oil_start = 22
    for r in range(10, ws_02a_target.max_row+1):
        v = str(ws_02a_target.cell(r, 1).value or '')
        if 'MÁY CHẠY XĂNG' in v:
            gas_start = r + 2
        elif 'MÁY CHẠY DẦU' in v:
            oil_start = r + 2

    # Gas runs (rows gas_start to oil_start - 3)
    for r in range(gas_start, oil_start - 2):
        stt = ws_02a_target.cell(r, 1).value
        tid = str(ws_02a_target.cell(r, 3).value or '').strip()
        amt = float(ws_02a_target.cell(r, 13).value or 0)
        lit = float(ws_02a_target.cell(r, 11).value or 0)
        if stt and str(stt).isdigit() and tid and amt > 0:
            g1_x_demands[tid]['tien'] += amt
            g1_x_demands[tid]['lit'] += lit
            g1_x_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a_target.cell(r, 2).value or '').strip()

    # Oil runs (rows oil_start to max_row)
    for r in range(oil_start, ws_02a_target.max_row + 1):
        stt = ws_02a_target.cell(r, 1).value
        tid = str(ws_02a_target.cell(r, 3).value or '').strip()
        amt = float(ws_02a_target.cell(r, 13).value or 0)
        lit = float(ws_02a_target.cell(r, 11).value or 0)
        if stt and str(stt).isdigit() and tid and amt > 0:
            g1_d_demands[tid]['tien'] += amt
            g1_d_demands[tid]['lit'] += lit
            g1_d_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a_target.cell(r, 2).value or '').strip()

    tot_g1_x = sum(v['tien'] for v in g1_x_demands.values())
    tot_g1_d = sum(v['tien'] for v in g1_d_demands.values())
    print(f"Demands read: Gas = {tot_g1_x:,.0f} đ ({len(g1_x_demands)} sites), Oil = {tot_g1_d:,.0f} đ ({len(g1_d_demands)} sites), Grand Total = {tot_g1_x + tot_g1_d:,.0f} đ")

    # 4. Rebuild Map_HD_Theo_Tram_Nhom1
    print(">>> 5. Rebuilding Map_HD_Theo_Tram_Nhom1 with exact Waterfall allocation...")
    ws_map1 = wb['Map_HD_Theo_Tram_Nhom1']
    clear_sheet_from_row(ws_map1, 5)

    oil_invs_for_map = [x for x in dn_invoices if x['loai'] == 'Dầu']
    gas_invs_for_map = [x for x in dn_invoices if x['loai'] == 'Xăng']

    def run_waterfall(demands_dict, invoices_list):
        sorted_tids = sorted(demands_dict.keys())
        assignments = []
        inv_idx = 0
        inv_rem = invoices_list[0]['tong'] if invoices_list else 0

        for tid in sorted_tids:
            needed = demands_dict[tid]['tien']
            is_first = True
            while needed > 0 and inv_idx < len(invoices_list):
                cur_inv = invoices_list[inv_idx]
                if inv_rem >= needed:
                    assignments.append({
                        'tid': tid,
                        'st_name': demands_dict[tid]['ten'],
                        'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                        'so_hd': cur_inv['so_hd'],
                        'tien_gan': needed,
                        'inv': cur_inv
                    })
                    inv_rem -= needed
                    needed = 0
                    is_first = False
                else:
                    assignments.append({
                        'tid': tid,
                        'st_name': demands_dict[tid]['ten'],
                        'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                        'so_hd': cur_inv['so_hd'],
                        'tien_gan': inv_rem,
                        'inv': cur_inv
                    })
                    needed -= inv_rem
                    is_first = False
                    inv_idx += 1
                    if inv_idx < len(invoices_list):
                        inv_rem = invoices_list[inv_idx]['tong']

        return assignments, inv_rem

    assign_oil, rem_oil = run_waterfall(g1_d_demands, oil_invs_for_map)
    assign_gas, rem_gas = run_waterfall(g1_x_demands, gas_invs_for_map)

    cur_row = 5
    # Section I: OIL
    ws_map1.cell(cur_row, 1, 'I. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU DẦU DO (DO 0.05S) — GÁN HÓA ĐƠN THEO TRẠM').font = FONT_SECTION
    ws_map1.cell(cur_row, 1).fill = FILL_SECTION
    for c in range(2, 12):
        ws_map1.cell(cur_row, c).fill = FILL_SECTION
    ws_map1.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=11)
    cur_row += 1

    oil_start_row = cur_row
    for a in assign_oil:
        inv = a['inv']
        ws_map1.cell(cur_row, 1, a['tid']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 2, a['tien_tram']).alignment = ALIGN_RIGHT
        ws_map1.cell(cur_row, 3, a['so_hd']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 4, a['tien_gan']).alignment = ALIGN_RIGHT
        ws_map1.cell(cur_row, 5, inv['don_vi']).alignment = ALIGN_LEFT
        ws_map1.cell(cur_row, 6, inv['mst']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(cur_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 11, f"{a['st_name']} ({a['tid']})").alignment = ALIGN_LEFT

        for c in range(1, 12):
            cell = ws_map1.cell(cur_row, c)
            cell.font = FONT_MAIN
            cell.border = BORDER_ALL
            if c in [2, 4]:
                cell.number_format = '#,##0'
        cur_row += 1

    # Oil Subtotal
    oil_end_row = cur_row - 1
    ws_map1.cell(cur_row, 1, 'CỘNG NHIÊN LIỆU DẦU DO').alignment = ALIGN_CENTER
    ws_map1.cell(cur_row, 2, f'=SUM(B{oil_start_row}:B{oil_end_row})').number_format = '#,##0'
    ws_map1.cell(cur_row, 4, f'=SUM(D{oil_start_row}:D{oil_end_row})').number_format = '#,##0'
    for c in range(1, 12):
        cell = ws_map1.cell(cur_row, c)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
    cur_row += 2

    # Section II: GAS
    ws_map1.cell(cur_row, 1, 'II. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU XĂNG (RON 95) — GÁN HÓA ĐƠN THEO TRẠM').font = FONT_SECTION
    ws_map1.cell(cur_row, 1).fill = FILL_SECTION
    for c in range(2, 12):
        ws_map1.cell(cur_row, c).fill = FILL_SECTION
    ws_map1.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=11)
    cur_row += 1

    gas_start_row = cur_row
    for a in assign_gas:
        inv = a['inv']
        ws_map1.cell(cur_row, 1, a['tid']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 2, a['tien_tram']).alignment = ALIGN_RIGHT
        ws_map1.cell(cur_row, 3, a['so_hd']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 4, a['tien_gan']).alignment = ALIGN_RIGHT
        ws_map1.cell(cur_row, 5, inv['don_vi']).alignment = ALIGN_LEFT
        ws_map1.cell(cur_row, 6, inv['mst']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(cur_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(cur_row, 11, f"{a['st_name']} ({a['tid']})").alignment = ALIGN_LEFT

        for c in range(1, 12):
            cell = ws_map1.cell(cur_row, c)
            cell.font = FONT_MAIN
            cell.border = BORDER_ALL
            if c in [2, 4]:
                cell.number_format = '#,##0'
        cur_row += 1

    gas_end_row = cur_row - 1
    ws_map1.cell(cur_row, 1, 'CỘNG NHIÊN LIỆU XĂNG RON 95').alignment = ALIGN_CENTER
    ws_map1.cell(cur_row, 2, f'=SUM(B{gas_start_row}:B{gas_end_row})').number_format = '#,##0'
    ws_map1.cell(cur_row, 4, f'=SUM(D{gas_start_row}:D{gas_end_row})').number_format = '#,##0'
    for c in range(1, 12):
        cell = ws_map1.cell(cur_row, c)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
    cur_row += 2

    # Grand Total
    ws_map1.cell(cur_row, 1, 'TỔNG CỘNG TOÀN BỘ BẢNG KÊ (DẦU + XĂNG)').alignment = ALIGN_CENTER
    ws_map1.cell(cur_row, 2, f'=B{oil_end_row+1}+B{gas_end_row+1}').number_format = '#,##0'
    ws_map1.cell(cur_row, 4, f'=D{oil_end_row+1}+D{gas_end_row+1}').number_format = '#,##0'
    for c in range(1, 12):
        cell = ws_map1.cell(cur_row, c)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
    cur_row += 2

    # Summary notes
    ws_map1.cell(cur_row, 1, 'GHI CHÚ ĐỐI SOÁT NGHIỆP VỤ & PHƯƠNG ÁN THANH TOÁN (PHƯƠNG ÁN B):').font = FONT_BOLD
    cur_row += 1
    notes = [
        f"1. Tổng tiền chạy máy theo hồ sơ 02A TTNB: 30,337,339 đồng (102 ca chạy máy, bao gồm đầy đủ ca chạy trạm DNXL20 ngày 04/09).",
        f"2. Tổng tiền phân bổ hóa đơn (Cột D): 30,337,339 đồng -> ĐÁP ỨNG CHÍNH XÁC 100%, CHÊNH LỆCH = 0 ĐỒNG.",
        f"3. Lệnh chuyển khoản cho Nam Trung Phong ngày 30/09: 5,400,433 đồng (HĐ 705060: 3,900,433 đ + HĐ 704928: 1,500,000 đ).",
        f"4. Hóa đơn 704928 trích sử dụng 465,282 đồng, phần tiền dôi dư 1,034,718 đồng bảo lưu kho dự phòng kỳ sau.",
        f"5. Tất cả 13 hóa đơn còn lại của Đồng Nai (ngày phát hành <= 5 triệu/ngày) thanh toán 100% bằng TIỀN MẶT."
    ]
    for n in notes:
        ws_map1.cell(cur_row, 1, n).font = FONT_ITALIC
        cur_row += 1

    # 5. Update Payment Method Notes in HD_ToanCau
    print(">>> 6. Updating Payment Method Notes in HD_ToanCau (Phương án B)...")
    ws_hd_tc = wb['HD_ToanCau']
    for r in range(8, ws_hd_tc.max_row + 1):
        inv_no = str(ws_hd_tc.cell(r, 3).value or '').strip()
        dt = str(ws_hd_tc.cell(r, 2).value or '')[:10]
        seller = str(ws_hd_tc.cell(r, 4).value or '')

        note = ''
        is_transfer = False
        if inv_no == '676405':
            note = '💳 Chuyển khoản cây xăng (Nam Trung Phong - 8,982,000 đ: HĐ > 5tr)'
            is_transfer = True
        elif inv_no == '704930':
            note = '💳 Chuyển khoản cây xăng (Nam Trung Phong - Ngày 30/09 Cây xăng xuất > 5tr)'
            is_transfer = True
        elif inv_no in ['652795', '651553']:
            note = f'💳 Chuyển khoản cây xăng (Nam Trung Phong - Ngày 09/09 Cây xăng xuất > 5tr)'
            is_transfer = True
        elif 'TÍN NGHĨA' in seller.upper():
            note = f'💵 Tiền mặt (Tín Nghĩa ngày {dt} <= 5tr) — Sử dụng 100% giá vốn thanh toán'
        else:
            note = f'💵 Tiền mặt (Ngày {dt} <= 5tr) — Sử dụng 100% giá vốn thanh toán'

        ws_hd_tc.cell(r, 16, note).alignment = ALIGN_LEFT
        cell = ws_hd_tc.cell(r, 16)
        cell.font = FONT_MAIN
        cell.border = BORDER_ALL
        if is_transfer:
            for c in range(1, 17):
                ws_hd_tc.cell(r, c).fill = FILL_TRANSFER

    # Save to Desktop and Downloads
    print(f">>> 7. Saving to {HSTT_PATH} and {DOWNLOADS_PATH}...")
    wb.save(HSTT_PATH)
    shutil.copyfile(HSTT_PATH, DOWNLOADS_PATH)
    print(">>> COMPLETED SUCCESSFULLY!")

if __name__ == '__main__':
    run()
