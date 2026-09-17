#!/usr/bin/env python3
"""
Xuất file Tọa độ & Cấu hình 4G / 5G khu vực huyện Tân Phú & Định Quán (Đồng Nai)
Dữ liệu tích hợp:
- datasites (Tọa độ GPS, Địa chỉ, QLT, Truyền dẫn, Vùng phủ)
- sran_5g_tracker (Cấu hình 4G/3G4G, Cấu hình 5G, Swap Solution, Thiết bị, Anten, PO, Cluster)
- datacells (Số lượng và chi tiết các cell 4G, 5G, 3G, 2G, PCI/PSC)
- infrastructure_projects (18 vị trí quy hoạch CSHT)
"""

import os
import sys
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from supabase import create_client

SUPABASE_URL = "https://lnmoczxjweuifacqujcu.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImxubW9jenhqd2V1aWZhY3F1amN1Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3Nzg2MzcxOTYsImV4cCI6MjA5NDIxMzE5Nn0.C0Si7ChY4T_mxLylSkDNJOUcj9D0uuGW_L4t7p9yONI"
supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

OUTPUT_FILE = "/Users/cang_it/Desktop/Toa_Do_Va_Cau_Hinh_4G_5G_Tan_Phu_Dinh_Quan.xlsx"

print("📡 1. Đang tải dữ liệu từ Supabase...")

# 1. Datasites
sites = supabase.table("datasites").select("*").execute().data or []
tp_dq_sites = [
    s for s in sites 
    if (s.get("site_id_old") or "").upper().startswith("DNTP") 
    or (s.get("site_id") or "").upper().startswith("DNTP") 
    or (s.get("site_id_old") or "").upper().startswith("DNDQ") 
    or (s.get("site_id") or "").upper().startswith("DNDQ")
]

# 2. SRAN Tracker
r1 = supabase.table("sran_5g_tracker").select("*").range(0, 999).execute()
r2 = supabase.table("sran_5g_tracker").select("*").range(1000, 1999).execute()
sran_list = (r1.data or []) + (r2.data or [])

sran_map = {}
for item in sran_list:
    sid = (item.get("site_id") or "").upper()
    sold = (item.get("site_id_old") or "").upper()
    if sid: sran_map[sid] = item
    if sold: sran_map[sold] = item
    raw = item.get("raw_data") or {}
    if raw.get("Radio_ID"): sran_map[str(raw["Radio_ID"]).upper()] = item
    if raw.get("Baseband_ID"): sran_map[str(raw["Baseband_ID"]).upper()] = item
    if raw.get("Site_ID (New)"): sran_map[str(raw["Site_ID (New)"]).upper()] = item

# 3. Datacells
cells = []
for offset in [0, 1000, 2000]:
    c_res = supabase.table("datacells").select("*").range(offset, offset + 999).execute().data or []
    cells.extend(c_res)

tp_dq_ids = set()
for s in tp_dq_sites:
    sid = (s.get("site_id") or "").upper()
    sold = (s.get("site_id_old") or "").upper()
    if sid: tp_dq_ids.add(sid)
    if sold: tp_dq_ids.add(sold)

matched_cells = [c for c in cells if (c.get("site_id") or "").upper() in tp_dq_ids]

cell_counts_by_site = {}
for c in matched_cells:
    sid = (c.get("site_id") or "").upper()
    ran = (c.get("ran") or "").upper()
    if sid not in cell_counts_by_site:
        cell_counts_by_site[sid] = {"2G": 0, "3G": 0, "4G": 0, "5G": 0, "total": 0}
    if ran in cell_counts_by_site[sid]:
        cell_counts_by_site[sid][ran] += 1
    cell_counts_by_site[sid]["total"] += 1

# 4. Infrastructure Projects (18 CSHT)
infra_res = supabase.table("infrastructure_projects").select("*").execute().data or []
tp_dq_infra = [
    p for p in infra_res 
    if any(k in (p.get("district") or "").lower() for k in ["tân phú", "định quán", "tan phu", "dinh quan"])
]

print(f"📊 Dữ liệu thu thập: {len(tp_dq_sites)} trạm, {len(matched_cells)} cells, {len(tp_dq_infra)} quy hoạch CSHT")

# Helper function to extract unified site record
def build_site_row(s, idx):
    sid = (s.get("site_id") or "").upper()
    sold = (s.get("site_id_old") or "").upper()
    loc = s.get("location_info") or {}
    mgmt = s.get("management_info") or {}
    tech = s.get("technical_info") or {}
    
    is_tp = sold.startswith("DNTP") or sid.startswith("DNTP")
    huyen = "Tân Phú" if is_tp else "Định Quán"
    
    sran = sran_map.get(sid) or sran_map.get(sold) or {}
    raw = sran.get("raw_data") or {}
    
    lat = float(loc.get("vi_do") or 0)
    lng = float(loc.get("kinh_do") or 0)
    
    # 5G & 4G Configs
    cfg_5g = sran.get("config_5g") or raw.get("5G_Config") or ""
    scope_5g = sran.get("scope_5g") or raw.get("5G Scope") or raw.get("5G_Scope") or ""
    has_5g = bool((cfg_5g and cfg_5g != "-" and cfg_5g != "0") or "5g" in scope_5g.lower() or sran.get("onair_date"))
    
    cfg_4g = sran.get("config_3g4g") or raw.get("3G4G Config") or ""
    scope_4g = sran.get("scope_3g4g") or raw.get("4G Scope") or raw.get("3G4G_Scope") or ""
    
    # Cell counts
    c_info = cell_counts_by_site.get(sid) or cell_counts_by_site.get(sold) or {"2G": 0, "3G": 0, "4G": 0, "5G": 0, "total": 0}
    
    # Category
    onair_date = sran.get("onair_date") or ""
    cluster = raw.get("Cluster_New") or raw.get("Cluster_Name") or sran.get("district") or ""
    
    if onair_date or (has_5g and ("_00_PILOT" in cluster or "_01_LT" in cluster)):
        category = "5G Onair"
    elif has_5g:
        category = "5G Quy hoạch"
    else:
        category = "4G Hiện hữu"
        
    return {
        "stt": idx,
        "huyen": huyen,
        "site_id_old": sold,
        "site_id_new": sid,
        "name": s.get("name") or raw.get("Site_Name") or "",
        "lat": lat,
        "lng": lng,
        "gmaps": f"https://www.google.com/maps?q={lat:.6f},{lng:.6f}" if lat and lng else "",
        "phuong_xa": loc.get("phuong_xa") or raw.get("Ward") or "",
        "dia_chi": loc.get("dia_chi") or raw.get("Address") or "",
        "vung_phu": mgmt.get("vung_phu") or raw.get("Type") or "MACRO",
        "tram_main": mgmt.get("tram_main") if mgmt.get("tram_main") != "KHÔNG" else "",
        "category": category,
        "has_5g": has_5g,
        "cfg_5g": cfg_5g if has_5g else "Không có",
        "scope_5g": scope_5g if has_5g else "",
        "cfg_4g": cfg_4g,
        "scope_4g": scope_4g,
        "swap_solution": sran.get("swap_solution") or raw.get("Swap_Solution") or "",
        "equip_solution": sran.get("equip_solution") or raw.get("Equip_Solution") or "",
        "antenna_solution": sran.get("antenna_solution") or raw.get("3G4G_Antenna_Solution") or "",
        "power_solution": sran.get("power_solution") or raw.get("Power_Solution") or "",
        "baseband": raw.get("Baseband_Scenario") or raw.get("RP6655_Site_Type") or ("BB6631" if raw.get("BB6631_Contract") == "1" else ""),
        "radio": raw.get("Air26_PO") or raw.get("5G_Air_Solution") or "",
        "cluster": cluster,
        "survey_date": sran.get("survey_date") or raw.get("Survey_Actual_Date") or "",
        "onair_date": onair_date,
        "cell_4g": c_info["4G"],
        "cell_5g": c_info["5G"],
        "cell_3g": c_info["3G"],
        "cell_total": c_info["total"],
        "qlt": mgmt.get("qlt") or "",
        "sdt_qlt": mgmt.get("sdt_qlt") or "",
        "loai_cap": tech.get("loai_ket_noi") or "",
        "don_vi_cap": tech.get("don_vi_van_hanh_cap") or tech.get("chu_dau_tu_cap") or "",
        "tuyen_cap": tech.get("last_mile_primary") or tech.get("huong_ket_noi") or ""
    }

# Build rows sorted by Huyện (Tân Phú -> Định Quán) then site_id_old
sorted_sites = sorted(
    tp_dq_sites, 
    key=lambda x: (
        0 if (x.get("site_id_old") or "").startswith("DNTP") else 1,
        x.get("site_id_old") or ""
    )
)

site_rows = [build_site_row(s, i + 1) for i, s in enumerate(sorted_sites)]

# Styling Tokens
font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="475569")
font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=9.5, color="0F172A")
font_data_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
font_coord = Font(name="Consolas", size=9.5, color="0284C7")
font_5g = Font(name="Calibri", size=9.5, bold=True, color="DB2777")
font_link = Font(name="Calibri", size=9.5, color="2563EB", underline="single")

fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_header_sub = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
fill_header_pink = PatternFill(start_color="BE185D", end_color="BE185D", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_5g_tint = PatternFill(start_color="FDF2F8", end_color="FDF2F8", fill_type="solid")
fill_tp_tint = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
fill_dq_tint = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
align_right = Alignment(horizontal='right', vertical='center', wrap_text=False)

print("📝 2. Đang tạo Workbook Excel chuyên nghiệp...")
wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Columns Definition for Site Sheets
COLUMNS = [
    ("STT", 6, align_center),
    ("Huyện", 12, align_center),
    ("Mã trạm cũ", 13, align_center),
    ("Mã trạm mới (5G/SRAN)", 22, align_center),
    ("Tên trạm / Vị trí", 24, align_left),
    ("Vĩ độ (Lat)", 13, align_right),
    ("Kinh độ (Lng)", 13, align_right),
    ("Bản đồ", 12, align_center),
    ("Xã / Thị trấn", 20, align_left),
    ("Địa chỉ chi tiết", 38, align_left),
    ("Vùng phủ", 11, align_center),
    ("Trạm Main", 12, align_center),
    ("Phân loại 5G/4G", 15, align_center),
    ("Cấu hình 5G", 18, align_center),
    ("Phạm vi 5G", 22, align_left),
    ("Cấu hình 4G / 3G4G", 32, align_left),
    ("Phạm vi 4G", 15, align_left),
    ("Giải pháp Swap", 15, align_center),
    ("Giải pháp Thiết bị", 18, align_left),
    ("Giải pháp Anten", 35, align_left),
    ("Giải pháp Nguồn", 25, align_left),
    ("Kịch bản Baseband", 22, align_left),
    ("Giải pháp Radio / Air26", 24, align_left),
    ("Cụm SRAN (Cluster)", 16, align_center),
    ("Ngày khảo sát", 14, align_center),
    ("Ngày Onair 5G", 14, align_center),
    ("Cell 4G", 9, align_right),
    ("Cell 5G", 9, align_right),
    ("Cell 3G", 9, align_right),
    ("Tổng Cell", 10, align_right),
    ("Người QLT", 18, align_left),
    ("SĐT QLT", 13, align_center),
    ("Loại cáp", 10, align_center),
    ("Đơn vị cáp", 16, align_left),
    ("Tuyến Lastmile", 28, align_left)
]

def render_site_sheet(ws, title, subtitle, rows, custom_header_fill=fill_header):
    # Title Block
    ws.merge_cells("A1:I1")
    ws["A1"] = title
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:I2")
    ws["A2"] = subtitle
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    # Header Row
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, (col_name, width, alignment) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = custom_header_fill
        cell.alignment = align_center
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(width, len(col_name) + 3)

    # Freeze panes below headers
    ws.freeze_panes = "E5"

    # Data Rows
    for r_idx, r_data in enumerate(rows, 5):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        row_fill = fill_5g_tint if r_data["has_5g"] else (fill_zebra if is_even else None)

        values = [
            r_data["stt"],
            r_data["huyen"],
            r_data["site_id_old"],
            r_data["site_id_new"],
            r_data["name"],
            r_data["lat"],
            r_data["lng"],
            "Google Maps",
            r_data["phuong_xa"],
            r_data["dia_chi"],
            r_data["vung_phu"],
            r_data["tram_main"],
            r_data["category"],
            r_data["cfg_5g"],
            r_data["scope_5g"],
            r_data["cfg_4g"],
            r_data["scope_4g"],
            r_data["swap_solution"],
            r_data["equip_solution"],
            r_data["antenna_solution"],
            r_data["power_solution"],
            r_data["baseband"],
            r_data["radio"],
            r_data["cluster"],
            r_data["survey_date"],
            r_data["onair_date"],
            r_data["cell_4g"],
            r_data["cell_5g"],
            r_data["cell_3g"],
            r_data["cell_total"],
            r_data["qlt"],
            r_data["sdt_qlt"],
            r_data["loai_cap"],
            r_data["don_vi_cap"],
            r_data["tuyen_cap"]
        ]

        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            # Alignment and specific formatting
            _, _, col_align = COLUMNS[c_idx - 1]
            cell.alignment = col_align

            # Coordinates format
            if c_idx in [6, 7]: # Lat, Lng
                cell.font = font_coord
                cell.number_format = '0.000000'
            elif c_idx == 8: # Google Maps Link
                if r_data["gmaps"]:
                    cell.hyperlink = r_data["gmaps"]
                    cell.font = font_link
            elif c_idx in [3, 4]: # Site IDs
                cell.font = font_data_bold
            elif c_idx == 13: # Category
                if "5G" in str(val):
                    cell.font = font_5g
                else:
                    cell.font = font_data
            elif c_idx == 14 and r_data["has_5g"]: # Cấu hình 5G
                cell.font = font_5g
            else:
                cell.font = font_data

# Sheet 1: Toàn bộ 119 trạm Tân Phú & Định Quán
ws1 = wb.create_sheet(title="TOAN_MANG_TP_DQ")
render_site_sheet(
    ws1,
    "BẢNG TỌA ĐỘ VÀ CẤU HÌNH 4G / 5G TOÀN MẠNG HUYỆN TÂN PHÚ & ĐỊNH QUÁN",
    f"Hệ thống cơ sở dữ liệu VHKT Tổ 3 | Tổng số: {len(site_rows)} trạm (50 trạm Tân Phú, 69 trạm Định Quán) | Xuất lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
    site_rows,
    fill_header
)

# Sheet 2: Huyện Tân Phú (50 trạm)
tp_rows = [r for r in site_rows if r["huyen"] == "Tân Phú"]
for idx, r in enumerate(tp_rows, 1):
    r_copy = dict(r)
    r_copy["stt"] = idx
    tp_rows[idx-1] = r_copy

ws2 = wb.create_sheet(title="HUYEN_TAN_PHU")
render_site_sheet(
    ws2,
    "DANH SÁCH TỌA ĐỘ VÀ CẤU HÌNH 4G / 5G - HUYỆN TÂN PHÚ",
    f"Địa bàn huyện Tân Phú | Tổng cộng: {len(tp_rows)} trạm phát sóng (Mã trạm cũ DNTP00 - DNTP54) | Tổ Viễn Thông 3",
    tp_rows,
    PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
)

# Sheet 3: Huyện Định Quán (69 trạm)
dq_rows = [r for r in site_rows if r["huyen"] == "Định Quán"]
for idx, r in enumerate(dq_rows, 1):
    r_copy = dict(r)
    r_copy["stt"] = idx
    dq_rows[idx-1] = r_copy

ws3 = wb.create_sheet(title="HUYEN_DINH_QUAN")
render_site_sheet(
    ws3,
    "DANH SÁCH TỌA ĐỘ VÀ CẤU HÌNH 4G / 5G - HUYỆN ĐỊNH QUÁN",
    f"Địa bàn huyện Định Quán | Tổng cộng: {len(dq_rows)} trạm phát sóng (Mã trạm cũ DNDQ00 - DNDQ71) | Tổ Viễn Thông 3",
    dq_rows,
    PatternFill(start_color="047857", end_color="047857", fill_type="solid")
)

# Sheet 4: Danh sách riêng 28 trạm có cấu hình 5G
s5g_rows = [r for r in site_rows if r["has_5g"]]
for idx, r in enumerate(s5g_rows, 1):
    r_copy = dict(r)
    r_copy["stt"] = idx
    s5g_rows[idx-1] = r_copy

ws4 = wb.create_sheet(title="TRAM_CO_CAU_HINH_5G")
render_site_sheet(
    ws4,
    "DANH SÁCH CÁC TRẠM CÓ CẤU HÌNH & QUY HOẠCH 5G (TÂN PHÚ - ĐỊNH QUÁN)",
    f"Các trạm trang bị thiết bị 5G NR26 / AIR 3265 | Tổng số: {len(s5g_rows)} trạm trọng điểm | Tổ Viễn Thông 3",
    s5g_rows,
    fill_header_pink
)

# Sheet 5: Chi tiết 864 Cell 4G/5G/3G từ datacells
print("📝 3. Đang tạo Sheet Chi tiết Cell...")
ws5 = wb.create_sheet(title="CHI_TIET_CELL_4G5G")
ws5.merge_cells("A1:G1")
ws5["A1"] = "DANH SÁCH CHI TIẾT CÁC CELL MẠNG (4G, 5G, 3G) - TÂN PHÚ & ĐỊNH QUÁN"
ws5["A1"].font = font_title
ws5["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws5.row_dimensions[1].height = 32

ws5.merge_cells("A2:G2")
ws5["A2"] = f"Dữ liệu chi tiết từ Datacells | Tổng cộng: {len(matched_cells)} cells (405 cell 4G, 6 cell 5G, 453 cell 3G) | Trích xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
ws5["A2"].font = font_subtitle
ws5["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws5.row_dimensions[2].height = 20
ws5.row_dimensions[3].height = 8

CELL_COLS = [
    ("STT", 6, align_center),
    ("Mã trạm (Site ID)", 15, align_center),
    ("Mã Cell (Cell ID)", 20, align_center),
    ("Tên Cell cũ", 16, align_center),
    ("Tên Cell mới (Canonical)", 22, align_center),
    ("Công nghệ (RAN)", 14, align_center),
    ("Hãng thiết bị (Vendor)", 14, align_center),
    ("Mã PCI / PSC", 12, align_right),
    ("Node ID", 12, align_right),
    ("Cell Index", 12, align_right),
    ("Vĩ độ (Latitude)", 13, align_right),
    ("Kinh độ (Longitude)", 13, align_right),
    ("Trạng thái", 12, align_center),
    ("Vùng phủ / Zone", 16, align_left),
    ("Lưu lượng (GB)", 14, align_right)
]

header_row = 4
ws5.row_dimensions[header_row].height = 28
for c_idx, (cname, width, alignment) in enumerate(CELL_COLS, 1):
    c = ws5.cell(row=header_row, column=c_idx, value=cname)
    c.font = font_header
    c.fill = fill_header_sub
    c.alignment = align_center
    c.border = thin_border
    col_letter = get_column_letter(c_idx)
    ws5.column_dimensions[col_letter].width = max(width, len(cname) + 3)

ws5.freeze_panes = "C5"

# Sort cells: 5G first, then 4G, then 3G, then by site_id
sorted_cells = sorted(
    matched_cells,
    key=lambda x: (
        0 if x.get("ran") == "5G" else (1 if x.get("ran") == "4G" else 2),
        x.get("site_id") or "",
        x.get("cell_id") or ""
    )
)

for r_idx, cell_data in enumerate(sorted_cells, 5):
    ws5.row_dimensions[r_idx].height = 20
    is_even = (r_idx % 2 == 0)
    ran = cell_data.get("ran", "")
    
    row_fill = None
    if ran == "5G": row_fill = fill_5g_tint
    elif is_even: row_fill = fill_zebra
    
    meta = cell_data.get("metadata") or {}
    lat = float(cell_data.get("latitude") or 0)
    lng = float(cell_data.get("longitude") or 0)
    traffic = float(meta.get("traffic") or 0)

    cell_vals = [
        r_idx - 4,
        cell_data.get("site_id"),
        cell_data.get("cell_id"),
        cell_data.get("cell_name_old"),
        cell_data.get("cell_name_new"),
        ran,
        cell_data.get("vendor"),
        cell_data.get("pci_psc"),
        cell_data.get("node_id"),
        cell_data.get("cell_idx"),
        lat,
        lng,
        cell_data.get("status"),
        f"{meta.get('vung_phu', '')} ({meta.get('zone', '')})".strip(),
        traffic
    ]

    for c_idx, val in enumerate(cell_vals, 1):
        cell_item = ws5.cell(row=r_idx, column=c_idx, value=val)
        cell_item.border = thin_border
        if row_fill: cell_item.fill = row_fill
        
        _, _, col_align = CELL_COLS[c_idx - 1]
        cell_item.alignment = col_align
        
        if c_idx in [11, 12]:
            cell_item.font = font_coord
            cell_item.number_format = '0.000000'
        elif c_idx == 6: # RAN
            if val == "5G": cell_item.font = font_5g
            elif val == "4G": cell_item.font = font_data_bold
            else: cell_item.font = font_data
        elif c_idx == 15:
            cell_item.number_format = '#,##0.00'
            cell_item.font = font_data
        else:
            cell_item.font = font_data

# Sheet 6: CSHT Quy hoạch (18 Vị trí)
print("📝 4. Đang tạo Sheet CSHT Quy hoạch...")
ws6 = wb.create_sheet(title="CSHT_QUY_HOACH_18_TRAM")
ws6.merge_cells("A1:G1")
ws6["A1"] = "DANH SÁCH 18 VỊ TRÍ HẠ TẦNG QUY HOẠCH CSHT - HUYỆN TÂN PHÚ & ĐỊNH QUÁN"
ws6["A1"].font = font_title
ws6["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws6.row_dimensions[1].height = 32

ws6.merge_cells("A2:G2")
ws6["A2"] = f"Hạ tầng quy hoạch cột trạm phát sóng (Sở KH&CN / Mobifone) | Tổng số: {len(tp_dq_infra)} vị trí | Xuất lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
ws6["A2"].font = font_subtitle
ws6["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws6.row_dimensions[2].height = 20
ws6.row_dimensions[3].height = 8

INFRA_COLS = [
    ("STT", 6, align_center),
    ("Huyện", 14, align_center),
    ("Mã quy hoạch mới", 18, align_center),
    ("Mã quy hoạch cũ", 16, align_center),
    ("Xã / Thị trấn", 20, align_left),
    ("Vĩ độ khảo sát", 14, align_right),
    ("Kinh độ khảo sát", 14, align_right),
    ("Bản đồ khảo sát", 14, align_center),
    ("Vĩ độ quy hoạch", 14, align_right),
    ("Kinh độ quy hoạch", 14, align_right),
    ("Tiến độ khảo sát", 16, align_center),
    ("Trạng thái phê duyệt", 18, align_center),
    ("Ý kiến Sở KH&CN", 22, align_left),
    ("Đơn vị dùng chung", 18, align_center),
    ("Mã trạm dùng chung", 18, align_center),
    ("Ghi chú hạ tầng", 35, align_left)
]

header_row = 4
ws6.row_dimensions[header_row].height = 28
for c_idx, (cname, width, alignment) in enumerate(INFRA_COLS, 1):
    c = ws6.cell(row=header_row, column=c_idx, value=cname)
    c.font = font_header
    c.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    c.alignment = align_center
    c.border = thin_border
    col_letter = get_column_letter(c_idx)
    ws6.column_dimensions[col_letter].width = max(width, len(cname) + 3)

ws6.freeze_panes = "E5"

sorted_infra = sorted(
    tp_dq_infra, 
    key=lambda x: (
        0 if "tân phú" in (x.get("district") or "").lower() else 1,
        x.get("planning_id_new") or ""
    )
)

for r_idx, proj in enumerate(sorted_infra, 5):
    ws6.row_dimensions[r_idx].height = 20
    is_even = (r_idx % 2 == 0)
    row_fill = fill_zebra if is_even else None
    
    lat_s = float(proj.get("latitude_survey") or 0)
    lng_s = float(proj.get("longitude_survey") or 0)
    lat_p = float(proj.get("latitude_plan") or 0)
    lng_p = float(proj.get("longitude_plan") or 0)
    gmaps_s = f"https://www.google.com/maps?q={lat_s:.6f},{lng_s:.6f}" if lat_s and lng_s else ""

    infra_vals = [
        r_idx - 4,
        proj.get("district"),
        proj.get("planning_id_new"),
        proj.get("planning_id_old"),
        proj.get("ward"),
        lat_s if lat_s else None,
        lng_s if lng_s else None,
        "Google Maps" if gmaps_s else "",
        lat_p if lat_p else None,
        lng_p if lng_p else None,
        proj.get("survey_status"),
        proj.get("overall_status"),
        proj.get("skhcn_status"),
        proj.get("sharing_partner"),
        proj.get("shared_site_id"),
        proj.get("notes") or proj.get("conflict_notes") or ""
    ]

    for c_idx, val in enumerate(infra_vals, 1):
        cell_item = ws6.cell(row=r_idx, column=c_idx, value=val)
        cell_item.border = thin_border
        if row_fill: cell_item.fill = row_fill
        
        _, _, col_align = INFRA_COLS[c_idx - 1]
        cell_item.alignment = col_align
        
        if c_idx in [6, 7, 9, 10]:
            if val:
                cell_item.font = font_coord
                cell_item.number_format = '0.000000'
        elif c_idx == 8 and gmaps_s:
            cell_item.hyperlink = gmaps_s
            cell_item.font = font_link
        elif c_idx in [3, 4]:
            cell_item.font = font_data_bold
        else:
            cell_item.font = font_data

print(f"💾 5. Đang lưu file Excel ra: {OUTPUT_FILE}...")
wb.save(OUTPUT_FILE)
print(f"🎉 XUẤT FILE HOÀN TẤT THÀNH CÔNG!")
