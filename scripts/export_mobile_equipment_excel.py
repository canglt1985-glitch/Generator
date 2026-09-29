from collections import defaultdict
import os
import json
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from supabase import create_client

load_dotenv('tvt3_v2/.env')
supabase = create_client(os.getenv('VITE_SUPABASE_URL'), os.getenv('VITE_SUPABASE_ANON_KEY'))

print("📡 Đang lấy dữ liệu thiết bị lưu động từ Supabase...")

# 1. Fetch mobile_equipment
equips = supabase.table('mobile_equipment').select('*').order('type').order('equipment_code').execute().data or []
# 2. Fetch equipment_transfers
transfers = supabase.table('equipment_transfers').select('*').order('transfer_date', desc=True).execute().data or []
# 3. Fetch datasites for station mapping
sites = supabase.table('datasites').select('site_id, site_id_old, name, location_info, infrastructure_info').execute().data or []

site_map = {}
for s in sites:
    sid = (s.get('site_id') or '').strip().upper()
    sold = (s.get('site_id_old') or '').strip().upper()
    if sid: site_map[sid] = s
    if sold: site_map[sold] = s

eq_map = {e['id']: e for e in equips}

def format_location(loc_code):
    if not loc_code: return '—'
    u = str(loc_code).strip().upper()
    if u == 'KHO':
        return 'Kho TVT3'
    elif 'NHÀ' in u:
        return loc_code
    
    st = site_map.get(u)
    if st:
        sid_old = st.get('site_id_old')
        sid_new = st.get('site_id')
        return sid_old if sid_old else sid_new
    return loc_code

def format_dt(dt_str):
    if not dt_str: return '—'
    try:
        # e.g. 2026-08-16T03:27:54.837+00:00 or 2026-03-04 14:23:33
        dt_clean = str(dt_str).replace('T', ' ').split('.')[0].split('+')[0]
        dt = datetime.strptime(dt_clean, '%Y-%m-%d %H:%M:%S')
        return dt.strftime('%d/%m/%Y %H:%M')
    except:
        try:
            dt = datetime.strptime(str(dt_str)[:10], '%Y-%m-%d')
            return dt.strftime('%d/%m/%Y')
        except:
            return str(dt_str)

# Extract mobile generators from datasites
mobile_gens_sites = []
for s in sites:
    infra = s.get('infrastructure_info') or {}
    loc = s.get('location_info') or {}
    mpds = (infra.get('may_phat_dien') or {}).get('mpd') or []
    for m in mpds:
        lld = str(m.get('loai_lap_dat') or '').strip().lower()
        ten = str(m.get('ten') or '').strip().upper()
        nh = str(m.get('nhan_hieu') or '').strip().upper()
        nl = str(m.get('nhien_lieu') or '').strip().lower()
        if 'lưu động' in lld or 'di động' in lld or 'luu dong' in lld or 'lưu động' in ten.lower() or 'lưu động' in nh.lower() or 'xăng' in nl:
            mobile_gens_sites.append({
                'site_id': s.get('site_id'),
                'site_id_old': s.get('site_id_old') or '—',
                'name': s.get('name'),
                'huyen': loc.get('huyen_cu') or loc.get('huyen') or '',
                'xa': loc.get('xa_cu') or loc.get('xa_moi') or '',
                'nhan_hieu': m.get('nhan_hieu') or 'MLĐ KYO POWER / KiBii',
                'cong_suat': f"{m.get('cong_suat')} kVA" if m.get('cong_suat') else '—',
                'nhien_lieu': m.get('nhien_lieu') or 'Xăng',
                'dinh_muc': float(m.get('dinh_muc') or m.get('dinh_muc_quy_chuan') or 0),
                'ghi_chu': m.get('ghi_chu') or m.get('tinh_trang') or 'Máy xăng lưu động tại trạm'
            })

mobile_gens_sites.sort(key=lambda x: (x['huyen'], x['site_id']))

print(f"✅ Đã tải: {len(equips)} thiết bị lưu động, {len(transfers)} lượt điều chuyển, {len(mobile_gens_sites)} máy xăng tại trạm.")

# ----------------- TẠO WORKBOOK -----------------
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Styles
font_title = Font(name='Segoe UI', size=15, bold=True, color='1F4E78')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='595959')
font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_data = Font(name='Segoe UI', size=9)
font_bold = Font(name='Segoe UI', size=9, bold=True)
font_highlight = Font(name='Segoe UI', size=9, bold=True, color='1F4E78')
font_green = Font(name='Segoe UI', size=9, bold=True, color='006100')
font_red = Font(name='Segoe UI', size=9, bold=True, color='9C0006')

fill_header_navy = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
fill_header_purple = PatternFill(start_color='7030A0', end_color='7030A0', fill_type='solid')
fill_header_teal = PatternFill(start_color='008080', end_color='008080', fill_type='solid')
fill_alt = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
fill_accent = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

fill_status_good = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
fill_status_bad = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

fill_loc_kho = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
fill_loc_site = PatternFill(start_color='EBF1F5', end_color='EBF1F5', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
border_header = Border(
    left=Side(style='thin', color='FFFFFF'),
    right=Side(style='thin', color='FFFFFF'),
    top=Side(style='medium', color='1F4E78'),
    bottom=Side(style='medium', color='1F4E78')
)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')

# ==================== SHEET 1: VỊ TRÍ THIẾT BỊ LƯU ĐỘNG ====================
ws1 = wb.create_sheet(title='Vị trí Thiết bị Lưu động')
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:J1')
ws1['A1'] = "DANH SÁCH & VỊ TRÍ HIỆN TẠI CỦA THIẾT BỊ LƯU ĐỘNG"
ws1['A1'].font = font_title
ws1['A1'].alignment = align_center

ws1.merge_cells('A2:J2')
ws1['A2'] = f"Đài Viễn thông 3 — Theo dõi thời gian thực | Cập nhật lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
ws1['A2'].font = font_subtitle
ws1['A2'].alignment = align_center

s1_headers = [
    'STT', 'Mã thiết bị', 'Phân loại', 'Thông số kỹ thuật / Model',
    'Ngày đưa vào SD', 'Số Serial',
    'Vị trí hiện tại', 'Tình trạng', 'Tồn nhiên liệu (Lít)', 'Ghi chú vận hành'
]
row_hdr1 = 4
ws1.row_dimensions[row_hdr1].height = 28
for c_idx, h in enumerate(s1_headers, 1):
    c = ws1.cell(row=row_hdr1, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header_navy
    c.alignment = align_center
    c.border = border_header

for idx, eq in enumerate(equips, 1):
    r = row_hdr1 + idx
    row_fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
    
    loc_val = eq.get('current_location') or 'KHO'
    loc_display = format_location(loc_val)
    status_val = eq.get('status') or 'Tốt'
    
    vals = [
        idx,
        eq.get('equipment_code'),
        eq.get('type'),
        eq.get('specifications') or 'N/A',
        eq.get('commissioning_date') or '—',
        eq.get('serial_number') or '—',
        loc_display,
        status_val,
        eq.get('fuel_balance') if eq.get('fuel_balance') is not None else 0,
        eq.get('notes') or '—'
    ]
    
    for c_idx, val in enumerate(vals, 1):
        c = ws1.cell(row=r, column=c_idx, value=val)
        c.font = font_data
        c.border = border_thin
        if row_fill.fill_type: c.fill = row_fill
        
        if c_idx in [1, 5]: c.alignment = align_center
        elif c_idx == 2: c.alignment = align_left; c.font = font_bold
        elif c_idx == 3: c.alignment = align_center
        elif c_idx in [4, 6, 10]: c.alignment = align_left
        elif c_idx == 7:
            c.alignment = align_left
            c.font = font_highlight
            if loc_val == 'KHO':
                c.fill = fill_loc_kho
            else:
                c.fill = fill_loc_site
        elif c_idx == 8:
            c.alignment = align_center
            if status_val == 'Tốt':
                c.font = font_green
                c.fill = fill_status_good
            else:
                c.font = font_red
                c.fill = fill_status_bad
        elif c_idx == 9:
            c.alignment = align_center

s1_widths = {1: 6, 2: 24, 3: 15, 4: 35, 5: 18, 6: 22, 7: 35, 8: 15, 9: 20, 10: 30}
for col_idx, width in s1_widths.items():
    ws1.column_dimensions[get_column_letter(col_idx)].width = width

ws1.freeze_panes = 'C5'
ws1.auto_filter.ref = f"A{row_hdr1}:J{row_hdr1 + len(equips)}"

# ==================== SHEET 2: NHẬT KÝ LỊCH SỬ ĐIỀU CHUYỂN ====================
ws2 = wb.create_sheet(title='Nhật ký Lịch sử Điều chuyển')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:G1')
ws2['A1'] = "NHẬT KÝ LỊCH SỬ ĐIỀU CHUYỂN & BÀN GIAO THIẾT BỊ LƯU ĐỘNG"
ws2['A1'].font = font_title
ws2['A1'].alignment = align_center

ws2.merge_cells('A2:G2')
ws2['A2'] = f"Ghi nhận chi tiết {len(transfers)} lượt điều chuyển giữa Kho và các trạm BTS | Sắp xếp mới nhất trước"
ws2['A2'].font = font_subtitle
ws2['A2'].alignment = align_center

s2_headers = [
    'STT', 'Thời gian điều chuyển', 'Mã thiết bị', 'Loại thiết bị',
    'Từ vị trí (Xuất phát)', 'Đến vị trí (Điểm đến)', 'Người thực hiện điều chuyển', 'Ghi chú / Lý do điều động'
]
row_hdr2 = 4
ws2.row_dimensions[row_hdr2].height = 28
for c_idx, h in enumerate(s2_headers, 1):
    c = ws2.cell(row=row_hdr2, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header_purple
    c.alignment = align_center
    c.border = border_header

for idx, tr in enumerate(transfers, 1):
    r = row_hdr2 + idx
    row_fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
    
    eq = eq_map.get(tr.get('equipment_id')) or {}
    dt_formatted = format_dt(tr.get('transfer_date'))
    from_loc = format_location(tr.get('from_location'))
    to_loc = format_location(tr.get('to_location'))
    
    vals = [
        idx,
        dt_formatted,
        eq.get('equipment_code') or '—',
        eq.get('type') or '—',
        from_loc,
        to_loc,
        tr.get('operator') or '—',
        tr.get('notes') or '—'
    ]
    
    for c_idx, val in enumerate(vals, 1):
        c = ws2.cell(row=r, column=c_idx, value=val)
        c.font = font_data
        c.border = border_thin
        if row_fill.fill_type: c.fill = row_fill
        
        if c_idx == 1: c.alignment = align_center
        elif c_idx == 2: c.alignment = align_center; c.font = font_bold
        elif c_idx == 3: c.alignment = align_left; c.font = font_highlight
        elif c_idx == 4: c.alignment = align_center
        elif c_idx in [5, 6]: c.alignment = align_left
        elif c_idx == 7: c.alignment = align_center; c.font = font_bold
        elif c_idx == 8: c.alignment = align_left

s2_widths = {1: 6, 2: 22, 3: 24, 4: 15, 5: 35, 6: 35, 7: 24, 8: 30}
for col_idx, width in s2_widths.items():
    ws2.column_dimensions[get_column_letter(col_idx)].width = width

ws2.freeze_panes = 'C5'
ws2.auto_filter.ref = f"A{row_hdr2}:H{row_hdr2 + len(transfers)}"

# ==================== SHEET 3: MÁY NỔ LƯU ĐỘNG GÁN TẠI TRẠM ====================
ws3 = wb.create_sheet(title='Máy nổ xăng lưu động tại Trạm')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:J1')
ws3['A1'] = "DANH SÁCH CÁC TRẠM SỬ DỤNG MÁY NỔ XĂNG LƯU ĐỘNG (THEO HẠ TẦNG TRẠM)"
ws3['A1'].font = font_title
ws3['A1'].alignment = align_center

ws3.merge_cells('A2:J2')
ws3['A2'] = f"Gồm {len(mobile_gens_sites)} trạm viễn thông đang ứng trực bằng máy xăng lưu động 5.5KVA - 7KVA thay thế máy dầu"
ws3['A2'].font = font_subtitle
ws3['A2'].alignment = align_center

s3_headers = [
    'STT', 'Mã trạm (Mới)', 'Mã trạm (Cũ)', 'Tên trạm', 'Huyện / Thị xã', 'Xã / Phường',
    'Nhãn hiệu máy nổ', 'Công suất', 'Nhiên liệu', 'Định mức tiêu hao (L/h)', 'Ghi chú điều chuyển / Cấu hình'
]
row_hdr3 = 4
ws3.row_dimensions[row_hdr3].height = 28
for c_idx, h in enumerate(s3_headers, 1):
    c = ws3.cell(row=row_hdr3, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header_teal
    c.alignment = align_center
    c.border = border_header

for idx, mg in enumerate(mobile_gens_sites, 1):
    r = row_hdr3 + idx
    row_fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
    
    vals = [
        idx,
        mg['site_id'],
        mg['site_id_old'],
        mg['name'],
        mg['huyen'],
        mg['xa'],
        mg['nhan_hieu'],
        mg['cong_suat'],
        mg['nhien_lieu'],
        mg['dinh_muc'] if mg['dinh_muc'] > 0 else '—',
        mg['ghi_chu']
    ]
    
    for c_idx, val in enumerate(vals, 1):
        c = ws3.cell(row=r, column=c_idx, value=val)
        c.font = font_data
        c.border = border_thin
        if row_fill.fill_type: c.fill = row_fill
        
        if c_idx in [1, 2, 3, 8, 9]: c.alignment = align_center
        elif c_idx in [4, 5, 6, 7, 11]: c.alignment = align_left
        elif c_idx == 10: c.alignment = align_right

s3_widths = {1: 6, 2: 15, 3: 15, 4: 24, 5: 18, 6: 20, 7: 22, 8: 14, 9: 14, 10: 22, 11: 45}
for col_idx, width in s3_widths.items():
    ws3.column_dimensions[get_column_letter(col_idx)].width = width

ws3.freeze_panes = 'C5'
ws3.auto_filter.ref = f"A{row_hdr3}:K{row_hdr3 + len(mobile_gens_sites)}"

# ==================== SHEET 4: BÁO CÁO TỔNG HỢP & KPI ====================
ws4 = wb.create_sheet(title='Tổng hợp & KPI Lưu động')
ws4.views.sheetView[0].showGridLines = True

ws4.merge_cells('A1:F1')
ws4['A1'] = "BÁO CÁO TỔNG HỢP & PHÂN PHỐI THIẾT BỊ LƯU ĐỘNG"
ws4['A1'].font = font_title
ws4['A1'].alignment = align_center

ws4.merge_cells('A2:F2')
ws4['A2'] = f"Thống kê thiết bị tại Kho vs Ứng trực tại trạm | Cập nhật: {datetime.now().strftime('%d/%m/%Y')}"
ws4['A2'].font = font_subtitle
ws4['A2'].alignment = align_center

# KPI block
ws4.cell(row=4, column=1, value="1. TỔNG QUAN PHÂN BỔ THIẾT BỊ LƯU ĐỘNG").font = Font(name='Segoe UI', size=11, bold=True, color='1F4E78')

mpd_equips = [e for e in equips if e.get('type') == 'MPĐ']
pin_equips = [e for e in equips if e.get('type') == 'Pin']

at_kho = [e for e in equips if (e.get('current_location') or '').upper() == 'KHO']
at_sites = [e for e in equips if (e.get('current_location') or '').upper() != 'KHO']

status_good = [e for e in equips if (e.get('status') or 'Tốt') == 'Tốt']
status_bad = [e for e in equips if (e.get('status') or 'Tốt') != 'Tốt']

kpi_table = [
    ('Chỉ số theo dõi', 'Số lượng', 'Tỷ lệ %', 'Ghi chú vận hành'),
    ('Tổng số thiết bị lưu động quản lý', len(equips), '100%', 'Gồm MPĐ lưu động và Pin Lithium lưu động'),
    ('• Máy phát điện lưu động (MPĐ 5KVA - 7KVA)', len(mpd_equips), f"{len(mpd_equips)/len(equips)*100:.1f}%", 'Phục vụ ứng cứu cúp điện di động'),
    ('• Pin lưu động (Lithium Postef 48V-100Ah)', len(pin_equips), f"{len(pin_equips)/len(equips)*100:.1f}%", 'Dự phòng nguồn DC cho trạm nghẽn tải'),
    ('Thiết bị đang dự phòng tại KHO TVT3', len(at_kho), f"{len(at_kho)/len(equips)*100:.1f}%", 'Sẵn sàng điều động ứng cứu khẩn cấp'),
    ('Thiết bị đang ứng trực thực tế tại TRẠM', len(at_sites), f"{len(at_sites)/len(equips)*100:.1f}%", 'Được điều chuyển cắm trực tiếp tại trạm'),
    ('Tình trạng hoạt động TỐT', len(status_good), f"{len(status_good)/len(equips)*100:.1f}%", 'Sẵn sàng nổ máy / cấp tải'),
    ('Tình trạng HƯ HỎNG (Cần sửa chữa)', len(status_bad), f"{len(status_bad)/len(equips)*100:.1f}%", f"Gồm {len(status_bad)} máy: " + ", ".join([f"{e.get('equipment_code')} ({e.get('specifications') or ''})" for e in status_bad]))
]

for idx, (c1, c2, c3, c4) in enumerate(kpi_table, 5):
    for col_i, val in enumerate([c1, c2, c3, c4], 1):
        c = ws4.cell(row=idx, column=col_i, value=val)
        c.font = font_header if idx == 5 else font_data
        c.border = border_header if idx == 5 else border_thin
        if idx == 5:
            c.fill = fill_header_navy
            c.alignment = align_center
        else:
            if col_i == 1:
                c.alignment = align_left
                if '•' in str(val) or 'Tổng' in str(val): c.font = font_bold
            elif col_i in [2, 3]:
                c.alignment = align_center
                c.font = font_bold
            else:
                c.alignment = align_left
            if 'HƯ HỎNG' in str(c1):
                c.font = font_red
                c.fill = fill_status_bad

# Operator transfers stats
ws4.cell(row=15, column=1, value="2. THỐNG KÊ LƯỢT ĐIỀU CHUYỂN THEO NHÂN SỰ").font = Font(name='Segoe UI', size=11, bold=True, color='1F4E78')

operator_counts = defaultdict(int)
for tr in transfers:
    op = tr.get('operator') or 'Chưa xác định'
    operator_counts[op] += 1

ws4.cell(row=16, column=1, value="STT").font = font_header; ws4.cell(row=16, column=1).fill = fill_header_navy; ws4.cell(row=16, column=1).alignment = align_center; ws4.cell(row=16, column=1).border = border_header
ws4.cell(row=16, column=2, value="Nhân sự thực hiện điều chuyển").font = font_header; ws4.cell(row=16, column=2).fill = fill_header_navy; ws4.cell(row=16, column=2).alignment = align_center; ws4.cell(row=16, column=2).border = border_header
ws4.cell(row=16, column=3, value="Số lượt điều chuyển").font = font_header; ws4.cell(row=16, column=3).fill = fill_header_navy; ws4.cell(row=16, column=3).alignment = align_center; ws4.cell(row=16, column=3).border = border_header
ws4.cell(row=16, column=4, value="Tỷ lệ %").font = font_header; ws4.cell(row=16, column=4).fill = fill_header_navy; ws4.cell(row=16, column=4).alignment = align_center; ws4.cell(row=16, column=4).border = border_header

for op_idx, (op_name, count) in enumerate(sorted(operator_counts.items(), key=lambda x: x[1], reverse=True), 1):
    r_op = 16 + op_idx
    pct = f"{count / len(transfers) * 100:.1f}%"
    for c_i, v in enumerate([op_idx, op_name, f"{count} lượt", pct], 1):
        cell = ws4.cell(row=r_op, column=c_i, value=v)
        cell.font = font_data
        cell.border = border_thin
        if c_i in [1, 3, 4]: cell.alignment = align_center
        else: cell.alignment = align_left; cell.font = font_bold

s4_widths = {1: 35, 2: 32, 3: 22, 4: 45}
for col_idx, width in s4_widths.items():
    ws4.column_dimensions[get_column_letter(col_idx)].width = width

# Save files
desktop_path = "/Users/cang_it/Desktop/Quan_Ly_Thiet_Bi_Luu_Dong_TVT3.xlsx"
drive_path = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/datasite/Quan_Ly_Thiet_Bi_Luu_Dong_TVT3.xlsx"
scratch_path = "/Users/cang_it/Antigravity/TVT3/scratch/Quan_Ly_Thiet_Bi_Luu_Dong_TVT3.xlsx"

wb.save(desktop_path)
print(f"✅ Đã lưu file Excel trên Desktop: {desktop_path}")

try:
    wb.save(drive_path)
    print(f"✅ Đã lưu bản sao trên Google Drive: {drive_path}")
except Exception as e:
    print(f"Lỗi lưu Google Drive: {e}")

wb.save(scratch_path)
print(f"✅ Đã lưu bản sao tại scratch: {scratch_path}")
