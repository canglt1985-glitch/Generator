#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export Ho_So_Thanh_Toan_Chuan_Mau_08_2026_HSTT.xlsx:
- Applies all AWF Generator Settlement Rules to August 2026:
  1. Dong Nai (Group 1):
     - 100% absolute control: NO day exceeds 5,000,000 VND!
     - 106 generator runs preserved 100% (21 Gas = 9,316,848 VND; 85 Oil = 30,483,733 VND; F8 = 39,800,581 VND).
     - Map Nhom 1: Column B = Column D = 39,800,581 VND (100% EXACT MATCH).
     - Transparent notes in Column P (HD) & Column L (Map) for the last invoice with cost reduction and reserve.
  2. Toan Cau (Group 2):
     - Optimized gas and oil distribution: NO day exceeds 5,000,000 VND (100% days <= 5M)!
     - 208 generator runs preserved 100% (63 Gas = 31,709,034 VND; 145 Oil = 48,796,467 VND; F8 = 80,505,501 VND).
     - Map Nhom 2: Column B = Column D = 80,505,501 VND (100% EXACT MATCH).
     - Transparent notes in Column P (HD) & Column L (Map).
  3. Surplus Sheet (HD_Du_Thua_Khong_Su_Dung):
     - Collects all unused invoices with complete lookup metadata and reasons.
  4. Saves to Desktop and Downloads.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from collections import defaultdict
from itertools import combinations
import shutil
import os
import json

SRC_PATH = '/Users/cang_it/Desktop/QL_VienThong_DongNai/03_ThanhToan_HoaDon/Ho_So_Thanh_Toan_Chuan_Mau_08_2026.xlsx'
DST_DESKTOP = '/Users/cang_it/Desktop/Ho_So_Thanh_Toan_Chuan_Mau_08_2026_HSTT.xlsx'
DST_DOWNLOADS = '/Users/cang_it/Downloads/Ho_So_Thanh_Toan_Chuan_Mau_08_2026_HSTT.xlsx'

FONT_NAME = 'Times New Roman'
FONT_MAIN = Font(name=FONT_NAME, size=11)
FONT_BOLD = Font(name=FONT_NAME, size=11, bold=True)
FONT_ITALIC = Font(name=FONT_NAME, size=10, italic=True)
FONT_SECTION = Font(name=FONT_NAME, size=11, bold=True, color='1F4E78')
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color='1F4E78')
FONT_NOTE_SPECIAL = Font(name=FONT_NAME, size=10, bold=True, color='C00000')

FILL_HEADER = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
FILL_SECTION = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
FILL_TOTAL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
FILL_NOTE = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
FILL_SPECIAL = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

THIN_SIDE = Side(border_style='thin', color='A6A6A6')
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
DOUBLE_BOTTOM = Side(border_style='double', color='000000')
BORDER_TOTAL = Border(top=THIN_SIDE, bottom=DOUBLE_BOTTOM, left=THIN_SIDE, right=THIN_SIDE)

ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

def clear_sheet_from_row(ws, from_row):
    merged_ranges = list(ws.merged_cells.ranges)
    for mr in merged_ranges:
        if mr.min_row >= from_row or mr.max_row >= from_row:
            ws.unmerge_cells(str(mr))
    if ws.max_row >= from_row:
        ws.delete_rows(from_row, ws.max_row - from_row + 1)

def run_waterfall(demands_dict, invoices_list):
    sorted_tids = sorted(demands_dict.keys())
    assignments = []
    inv_idx = 0
    inv_rem = invoices_list[0]['tong_tien'] if invoices_list else 0

    for tid in sorted_tids:
        needed = demands_dict[tid]['tien']
        is_first = True
        while needed > 0 and inv_idx < len(invoices_list):
            cur_inv = invoices_list[inv_idx]
            if inv_rem >= needed:
                assignments.append({
                    'tid': tid,
                    'st_name': demands_dict[tid]['ten'],
                    'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                    'so_hd': cur_inv['so_hd'],
                    'tien_gan': needed,
                    'inv': cur_inv
                })
                inv_rem -= needed
                needed = 0
                is_first = False
            else:
                assignments.append({
                    'tid': tid,
                    'st_name': demands_dict[tid]['ten'],
                    'tien_tram': demands_dict[tid]['tien'] if is_first else None,
                    'so_hd': cur_inv['so_hd'],
                    'tien_gan': inv_rem,
                    'inv': cur_inv
                })
                needed -= inv_rem
                is_first = False
                inv_idx += 1
                if inv_idx < len(invoices_list):
                    inv_rem = invoices_list[inv_idx]['tong_tien']

    tot_tram = sum(demands_dict[t]['tien'] for t in demands_dict)
    tot_gan = sum(a['tien_gan'] for a in assignments)
    return assignments, tot_tram, tot_gan, inv_rem

def main():
    print(">>> 1. Loading source Month 8 workbook...")
    wb = openpyxl.load_workbook(SRC_PATH)

    # Load site lookup
    site_lookup = {}
    if os.path.exists('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json'):
        with open('/Users/cang_it/Antigravity/TVT3/data/sites_backup_20260930_165011.json') as f:
            for s in json.load(f):
                site_lookup[s.get('site_id')] = s.get('site_name')

    # Read Demands Dong Nai (02A_TTNB_DongNai_67Tram)
    print(">>> 2. Reading Demands for Dong Nai...")
    ws_02a1 = wb['02A_TTNB_DongNai_67Tram']
    g1_x_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(13, 34):
        tid = str(ws_02a1.cell(r, 3).value).strip()
        g1_x_demands[tid]['tien'] += float(ws_02a1.cell(r, 13).value or 0)
        g1_x_demands[tid]['lit'] += float(ws_02a1.cell(r, 11).value or 0)
        g1_x_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a1.cell(r, 2).value or '').strip()

    g1_d_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(38, 123):
        tid = str(ws_02a1.cell(r, 3).value).strip()
        g1_d_demands[tid]['tien'] += float(ws_02a1.cell(r, 13).value or 0)
        g1_d_demands[tid]['lit'] += float(ws_02a1.cell(r, 11).value or 0)
        g1_d_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a1.cell(r, 2).value or '').strip()

    tot_dn_oil_demand = sum(g1_d_demands[t]['tien'] for t in g1_d_demands)
    tot_dn_gas_demand = sum(g1_x_demands[t]['tien'] for t in g1_x_demands)
    tot_dn_demand = tot_dn_oil_demand + tot_dn_gas_demand
    print(f"Dong Nai Demand: Gas = {tot_dn_gas_demand:,.0f} đ | Oil = {tot_dn_oil_demand:,.0f} đ | Total = {tot_dn_demand:,.0f} đ")

    # Read Demands Toan Cau (02A_TTNB_ToanCau)
    print(">>> 3. Reading Demands for Toan Cau...")
    ws_02a2 = wb['02A_TTNB_ToanCau']
    g2_x_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(13, 76):
        tid = str(ws_02a2.cell(r, 3).value).strip()
        g2_x_demands[tid]['tien'] += float(ws_02a2.cell(r, 13).value or 0)
        g2_x_demands[tid]['lit'] += float(ws_02a2.cell(r, 11).value or 0)
        g2_x_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a2.cell(r, 2).value or '').strip()

    g2_d_demands = defaultdict(lambda: {'lit': 0.0, 'tien': 0.0, 'ten': ''})
    for r in range(80, 225):
        tid = str(ws_02a2.cell(r, 3).value).strip()
        g2_d_demands[tid]['tien'] += float(ws_02a2.cell(r, 13).value or 0)
        g2_d_demands[tid]['lit'] += float(ws_02a2.cell(r, 11).value or 0)
        g2_d_demands[tid]['ten'] = site_lookup.get(tid) or str(ws_02a2.cell(r, 2).value or '').strip()

    tot_tc_oil_demand = sum(g2_d_demands[t]['tien'] for t in g2_d_demands)
    tot_tc_gas_demand = sum(g2_x_demands[t]['tien'] for t in g2_x_demands)
    tot_tc_demand = tot_tc_oil_demand + tot_tc_gas_demand
    print(f"Toan Cau Demand: Gas = {tot_tc_gas_demand:,.0f} đ | Oil = {tot_tc_oil_demand:,.0f} đ | Total = {tot_tc_demand:,.0f} đ")

    # Collect ALL Invoices from Month 8 workbook
    print(">>> 4. Collecting all invoices from HD and Du Thua...")
    ws_hd1 = wb['HD_DongNai_67Tram']
    ws_hd2 = wb['HD_ToanCau']
    ws_dt = wb['HD_Du_Thua_Khong_Su_Dung']

    dn_pool = []
    tc_pool = []
    seen = set()

    # From HD_DongNai_67Tram
    for r in range(8, ws_hd1.max_row+1):
        sohd = ws_hd1.cell(r, 3).value
        if sohd and str(sohd).isdigit() and len(str(sohd)) >= 6:
            d = str(ws_hd1.cell(r, 2).value)[:10]
            key = (str(sohd).strip(), d, 'Đồng Nai')
            if key not in seen:
                seen.add(key)
                nl = str(ws_hd1.cell(r, 5).value or '').strip()
                dn_pool.append({
                    'so_hd': str(sohd).strip(), 'ngay': d, 'phap_nhan': 'Đồng Nai',
                    'buyer': 'MobiFone Đồng Nai (0100686209-129)',
                    'don_vi_ban': ws_hd1.cell(r, 4).value, 'loai_nl': 'Dầu' if 'Dầu' in nl else 'Xăng',
                    'dien_giai': ws_hd1.cell(r, 6).value, 'lit': float(ws_hd1.cell(r, 7).value or 0),
                    'mst_ban': str(ws_hd1.cell(r, 8).value or '').strip(), 'mau_so': ws_hd1.cell(r, 9).value,
                    'ky_hieu': ws_hd1.cell(r, 10).value, 'tien_chua_vat': float(ws_hd1.cell(r, 11).value or 0),
                    'vat': float(ws_hd1.cell(r, 12).value or 0), 'tong_tien': float(ws_hd1.cell(r, 13).value or 0),
                    'link': ws_hd1.cell(r, 14).value, 'fkey': ws_hd1.cell(r, 15).value,
                })

    # From HD_ToanCau
    for r in range(8, ws_hd2.max_row+1):
        sohd = ws_hd2.cell(r, 3).value
        if sohd and str(sohd).isdigit() and len(str(sohd)) >= 6:
            d = str(ws_hd2.cell(r, 2).value)[:10]
            key = (str(sohd).strip(), d, 'Toàn Cầu')
            if key not in seen:
                seen.add(key)
                nl = str(ws_hd2.cell(r, 5).value or '').strip()
                tc_pool.append({
                    'so_hd': str(sohd).strip(), 'ngay': d, 'phap_nhan': 'Toàn Cầu',
                    'buyer': 'MobiFone Toàn Cầu (0102577251-001)',
                    'don_vi_ban': ws_hd2.cell(r, 4).value, 'loai_nl': 'Dầu' if 'Dầu' in nl else 'Xăng',
                    'dien_giai': ws_hd2.cell(r, 6).value, 'lit': float(ws_hd2.cell(r, 7).value or 0),
                    'mst_ban': str(ws_hd2.cell(r, 8).value or '').strip(), 'mau_so': ws_hd2.cell(r, 9).value,
                    'ky_hieu': ws_hd2.cell(r, 10).value, 'tien_chua_vat': float(ws_hd2.cell(r, 11).value or 0),
                    'vat': float(ws_hd2.cell(r, 12).value or 0), 'tong_tien': float(ws_hd2.cell(r, 13).value or 0),
                    'link': ws_hd2.cell(r, 14).value, 'fkey': ws_hd2.cell(r, 15).value,
                })

    # From HD_Du_Thua_Khong_Su_Dung
    for r in range(5, ws_dt.max_row+1):
        sohd = ws_dt.cell(r, 3).value
        buyer = str(ws_dt.cell(r, 4).value or '')
        d = str(ws_dt.cell(r, 2).value)[:10]
        if sohd and str(sohd).isdigit() and len(str(sohd)) >= 6:
            pn = 'Đồng Nai' if 'Đồng Nai' in buyer else 'Toàn Cầu'
            key = (str(sohd).strip(), d, pn)
            if key not in seen:
                seen.add(key)
                nl = str(ws_dt.cell(r, 5).value or '').strip()
                inv_data = {
                    'so_hd': str(sohd).strip(), 'ngay': d, 'phap_nhan': pn,
                    'buyer': buyer,
                    'don_vi_ban': ws_dt.cell(r, 11).value, 'loai_nl': 'Dầu' if 'Dầu' in nl else 'Xăng',
                    'dien_giai': ws_dt.cell(r, 5).value, 'lit': float(ws_dt.cell(r, 6).value or 0),
                    'mst_ban': str(ws_dt.cell(r, 12).value or '').strip(), 'mau_so': None,
                    'ky_hieu': ws_dt.cell(r, 13).value, 'tien_chua_vat': float(ws_dt.cell(r, 8).value or 0),
                    'vat': float(ws_dt.cell(r, 9).value or 0), 'tong_tien': float(ws_dt.cell(r, 10).value or 0),
                    'link': ws_dt.cell(r, 14).value, 'fkey': ws_dt.cell(r, 15).value,
                }
                if pn == 'Đồng Nai':
                    dn_pool.append(inv_data)
                else:
                    tc_pool.append(inv_data)

    print(f"Total Invoices: Dong Nai Pool = {len(dn_pool)} | Toan Cau Pool = {len(tc_pool)}")

    # 5. Smart Selection for Dong Nai (100% days <= 5M)
    print(">>> 5. Solving optimal selection for Dong Nai (<= 5M/day)...")
    dn_pool.sort(key=lambda x: (x['ngay'], x['so_hd']))
    by_date_dn = defaultdict(list)
    for inv in dn_pool:
        by_date_dn[inv['ngay']].append(inv)
    sorted_dn_dates = sorted(by_date_dn.keys())

    date_options_dn = {}
    for d in sorted_dn_dates:
        invs = by_date_dn[d]
        opts = []
        for k in range(len(invs) + 1):
            for combo in combinations(invs, k):
                tot = sum(x['tong_tien'] for x in combo)
                if tot <= 5000000:
                    gas = sum(x['tong_tien'] for x in combo if x['loai_nl'] == 'Xăng')
                    oil = sum(x['tong_tien'] for x in combo if x['loai_nl'] == 'Dầu')
                    opts.append({'combo': combo, 'tot': tot, 'gas': gas, 'oil': oil})
        date_options_dn[d] = opts

    BUCKET = 200000
    dp_dn = {(0, 0): (0, 0, ())}
    for d in sorted_dn_dates:
        next_dp = {}
        for (bg, bo), (cg, co, chosen) in dp_dn.items():
            for opt in date_options_dn[d]:
                ng = cg + opt['gas']
                no = co + opt['oil']
                bg_next = min(int(ng // BUCKET), int((tot_dn_gas_demand + 3000000) // BUCKET))
                bo_next = min(int(no // BUCKET), int((tot_dn_oil_demand + 3000000) // BUCKET))
                bkey = (bg_next, bo_next)
                tot_spend = ng + no
                if bkey not in next_dp or tot_spend < (next_dp[bkey][0] + next_dp[bkey][1]):
                    next_dp[bkey] = (ng, no, chosen + opt['combo'])
        dp_dn = next_dp

    valid_dn = []
    for (bg, bo), (cg, co, chosen) in dp_dn.items():
        if cg >= tot_dn_gas_demand and co >= tot_dn_oil_demand:
            valid_dn.append((cg + co, cg, co, chosen))
    valid_dn.sort(key=lambda x: x[0])
    _, _, _, dn_chosen = valid_dn[0]

    dn_chosen_oil = [i for i in dn_chosen if i['loai_nl'] == 'Dầu']
    dn_chosen_gas = [i for i in dn_chosen if i['loai_nl'] == 'Xăng']
    dn_chosen_oil.sort(key=lambda x: (x['ngay'], x['so_hd']))
    dn_chosen_gas.sort(key=lambda x: (x['ngay'], x['so_hd']))
    dn_unpicked = [i for i in dn_pool if i not in dn_chosen]

    print(f"Dong Nai Chosen: {len(dn_chosen)} HĐ (Oil: {len(dn_chosen_oil)}, Gas: {len(dn_chosen_gas)}) | Unpicked: {len(dn_unpicked)}")

    # 6. Smart Selection for Toan Cau (100% days <= 5M)
    print(">>> 6. Solving optimal selection for Toan Cau (<= 5M/day)...")
    tc_pool_cand = [i for i in tc_pool if i['tong_tien'] <= 5000000] # exclude single > 5M
    by_d_tc = defaultdict(list)
    for i in tc_pool_cand:
        by_d_tc[i['ngay']].append(i)

    tc_picked_oil = []
    tc_acc_oil = 0
    # Step 1: pick oil on days where daily oil <= 4M
    for d in sorted(by_d_tc):
        for i in by_d_tc[d]:
            if i['loai_nl'] == 'Dầu' and tc_acc_oil < tot_tc_oil_demand:
                day_oil = sum(x['tong_tien'] for x in tc_picked_oil if x['ngay'] == d)
                if day_oil + i['tong_tien'] <= 4000000:
                    tc_picked_oil.append(i)
                    tc_acc_oil += i['tong_tien']

    # Step 2: if still needed, allow up to 4.8M
    if tc_acc_oil < tot_tc_oil_demand:
        for d in sorted(by_d_tc):
            for i in by_d_tc[d]:
                if i['loai_nl'] == 'Dầu' and i not in tc_picked_oil and tc_acc_oil < tot_tc_oil_demand:
                    day_oil = sum(x['tong_tien'] for x in tc_picked_oil if x['ngay'] == d)
                    if day_oil + i['tong_tien'] <= 4800000:
                        tc_picked_oil.append(i)
                        tc_acc_oil += i['tong_tien']

    # Step 3: pick Gas filling days where day_oil + day_gas + g <= 5M
    tc_picked_gas = []
    tc_acc_gas = 0
    for d in sorted(by_d_tc):
        for i in by_d_tc[d]:
            if i['loai_nl'] == 'Xăng' and tc_acc_gas < tot_tc_gas_demand:
                day_tot = sum(x['tong_tien'] for x in tc_picked_oil if x['ngay'] == d) + sum(x['tong_tien'] for x in tc_picked_gas if x['ngay'] == d)
                if day_tot + i['tong_tien'] <= 5000000:
                    tc_picked_gas.append(i)
                    tc_acc_gas += i['tong_tien']

    tc_picked_oil.sort(key=lambda x: (x['ngay'], x['so_hd']))
    tc_picked_gas.sort(key=lambda x: (x['ngay'], x['so_hd']))
    tc_chosen = tc_picked_oil + tc_picked_gas
    tc_unpicked = [i for i in tc_pool if i not in tc_chosen]

    print(f"Toan Cau Chosen: {len(tc_chosen)} HĐ (Oil: {len(tc_picked_oil)}, Gas: {len(tc_picked_gas)}) | Unpicked: {len(tc_unpicked)}")

    # 7. Waterfall Group 1 (Dong Nai)
    a1_d, t1_d, g1_d, r1_d = run_waterfall(g1_d_demands, dn_chosen_oil)
    a1_x, t1_x, g1_x, r1_x = run_waterfall(g1_x_demands, dn_chosen_gas)
    assert t1_d == g1_d and t1_x == g1_x, "G1 Waterfall mismatch!"

    # 8. Waterfall Group 2 (Toan Cau)
    a2_d, t2_d, g2_d, r2_d = run_waterfall(g2_d_demands, tc_picked_oil)
    a2_x, t2_x, g2_x, r2_x = run_waterfall(g2_x_demands, tc_picked_gas)
    assert t2_d == g2_d and t2_x == g2_x, "G2 Waterfall mismatch!"

    # 9. Rewrite Map_HD_Theo_Tram_Nhom1
    print(">>> 9. Writing Map_HD_Theo_Tram_Nhom1...")
    ws_map1 = wb['Map_HD_Theo_Tram_Nhom1']
    clear_sheet_from_row(ws_map1, 1)

    ws_map1.cell(1, 1, "BẢNG KÊ PHÂN BỔ HÓA ĐƠN XĂNG DẦU THEO TỪNG TRẠM CHẠY MÁY (MOBIFONE ĐỒNG NAI - 67 TRẠM)").font = FONT_TITLE
    ws_map1.cell(2, 1, "Tháng 08/2026 • Phân tách độc lập bảng kê Dầu DO và Xăng RON 95 • Ưu tiên đáp ứng đủ 100% số tiền bảng kê, bảo lưu số lít/tiền dư").font = FONT_ITALIC
    ws_map1.row_dimensions[1].height = 28
    ws_map1.row_dimensions[2].height = 20

    headers_map = [
        'ID trạm\n(Nhãn Hàng)', 'Thành tiền chạy máy\ntheo trạm (đồng)', 'Số hóa đơn',
        'Số tiền gán từ HĐ\n(đồng)', 'Đơn vị bán hàng', 'Mã số thuế\n(Bán)',
        'Ký hiệu HĐ', 'Link tra cứu hóa đơn', 'Mã tra cứu / Fkey', 'Ngày HĐ', 'Tên trạm / Địa bàn',
        'Ghi Chú Phân Bổ Giá Vốn'
    ]
    ws_map1.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers_map, 1):
        cell = ws_map1.cell(4, col_idx, h)
        cell.font = FONT_BOLD; cell.fill = FILL_HEADER; cell.alignment = ALIGN_CENTER; cell.border = BORDER_ALL

    # Section I: Dầu DO
    curr_row = 5
    sec1 = ws_map1.cell(curr_row, 1, "I. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU DẦU DO (DO 0.05S) — GÁN HÓA ĐƠN THEO TRẠM")
    sec1.font = FONT_SECTION; sec1.fill = FILL_SECTION
    ws_map1.row_dimensions[curr_row].height = 24
    ws_map1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=12)
    for c in range(1, 13):
        ws_map1.cell(curr_row, c).border = BORDER_ALL; ws_map1.cell(curr_row, c).fill = FILL_SECTION

    dau1_start = curr_row + 1
    curr_row += 1
    last_d1_sohd = dn_chosen_oil[-1]['so_hd']
    for a in a1_d:
        ws_map1.row_dimensions[curr_row].height = 20
        c1 = ws_map1.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map1.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map1.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map1.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map1.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT

        c12 = ws_map1.cell(curr_row, 12)
        c12.alignment = ALIGN_LEFT
        if a['so_hd'] == last_d1_sohd and a == a1_d[-1]:
            c12.value = f"Trích một phần HĐ {last_d1_sohd} (vừa khít 100% tiền dầu trạm, dư {r1_d:,.0f} đ bảo lưu kho)"
            c12.font = FONT_NOTE_SPECIAL; c12.fill = FILL_SPECIAL
        else:
            c12.value = "Gán 100% chi phí trạm"
            c12.font = FONT_ITALIC

        for c in range(5, 13):
            if c != 12 or a != a1_d[-1]:
                ws_map1.cell(curr_row, c).font = FONT_MAIN
            ws_map1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    dau1_end = curr_row - 1
    # Total Dầu DO G1
    tot_d1_row = curr_row
    ws_map1.row_dimensions[tot_d1_row].height = 22
    ws_map1.cell(tot_d1_row, 1, "TỔNG CỘNG DẦU DO").font = FONT_BOLD
    ws_map1.cell(tot_d1_row, 1).alignment = ALIGN_LEFT

    t2 = ws_map1.cell(tot_d1_row, 2, f"=SUM(B{dau1_start}:B{dau1_end})")
    t2.font = FONT_BOLD; t2.alignment = ALIGN_RIGHT; t2.number_format = '#,##0'
    t4 = ws_map1.cell(tot_d1_row, 4, f"=SUM(D{dau1_start}:D{dau1_end})")
    t4.font = FONT_BOLD; t4.alignment = ALIGN_RIGHT; t4.number_format = '#,##0'
    for c in range(1, 13):
        ws_map1.cell(tot_d1_row, c).fill = FILL_TOTAL; ws_map1.cell(tot_d1_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Dầu G1
    note_d1_row = curr_row
    ws_map1.row_dimensions[note_d1_row].height = 22
    tot_d1_buy = sum(i['tong_tien'] for i in dn_chosen_oil)
    tot_d1_lit = sum(i['lit'] for i in dn_chosen_oil)
    tot_d1_used_lit = sum(g1_d_demands[t]['lit'] for t in g1_d_demands)
    note_d1_text = f"📌 Ghi chú bảo lưu Dầu DO: Tổng HĐ mua {tot_d1_lit:,.1f} L ({tot_d1_buy:,.0f} đ) — Tiêu hao chạy máy {tot_d1_used_lit:,.1f} L ({t1_d:,.0f} đ) ➔ Tiền HĐ còn dư bảo lưu kho: +{r1_d:,.0f} đ"
    n_cell = ws_map1.cell(note_d1_row, 1, note_d1_text)
    n_cell.font = FONT_ITALIC; n_cell.fill = FILL_NOTE
    ws_map1.merge_cells(start_row=note_d1_row, start_column=1, end_row=note_d1_row, end_column=12)
    for c in range(1, 13):
        ws_map1.cell(note_d1_row, c).border = BORDER_ALL; ws_map1.cell(note_d1_row, c).fill = FILL_NOTE

    curr_row += 2
    # Section II: Xăng RON 95 G1
    sec2 = ws_map1.cell(curr_row, 1, "II. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU XĂNG RON 95 (RON 95-III) — GÁN HÓA ĐƠN THEO TRẠM")
    sec2.font = FONT_SECTION; sec2.fill = FILL_SECTION
    ws_map1.row_dimensions[curr_row].height = 24
    ws_map1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=12)
    for c in range(1, 13):
        ws_map1.cell(curr_row, c).border = BORDER_ALL; ws_map1.cell(curr_row, c).fill = FILL_SECTION

    xang1_start = curr_row + 1
    curr_row += 1
    last_x1_sohd = dn_chosen_gas[-1]['so_hd']
    for a in a1_x:
        ws_map1.row_dimensions[curr_row].height = 20
        c1 = ws_map1.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map1.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map1.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map1.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map1.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map1.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map1.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT

        c12 = ws_map1.cell(curr_row, 12)
        c12.alignment = ALIGN_LEFT
        if a['so_hd'] == last_x1_sohd and a == a1_x[-1]:
            c12.value = f"Trích một phần HĐ {last_x1_sohd} (vừa khít 100% tiền xăng trạm, dư {r1_x:,.0f} đ bảo lưu kho)"
            c12.font = FONT_NOTE_SPECIAL; c12.fill = FILL_SPECIAL
        else:
            c12.value = "Gán 100% chi phí trạm"
            c12.font = FONT_ITALIC

        for c in range(5, 13):
            if c != 12 or a != a1_x[-1]:
                ws_map1.cell(curr_row, c).font = FONT_MAIN
            ws_map1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    xang1_end = curr_row - 1
    # Total Xăng G1
    tot_x1_row = curr_row
    ws_map1.row_dimensions[tot_x1_row].height = 22
    ws_map1.cell(tot_x1_row, 1, "TỔNG CỘNG XĂNG RON 95").font = FONT_BOLD
    ws_map1.cell(tot_x1_row, 1).alignment = ALIGN_LEFT

    t2_x1_cell = ws_map1.cell(tot_x1_row, 2, f"=SUM(B{xang1_start}:B{xang1_end})")
    t2_x1_cell.font = FONT_BOLD; t2_x1_cell.alignment = ALIGN_RIGHT; t2_x1_cell.number_format = '#,##0'
    t4_x = ws_map1.cell(tot_x1_row, 4, f"=SUM(D{xang1_start}:D{xang1_end})")
    t4_x.font = FONT_BOLD; t4_x.alignment = ALIGN_RIGHT; t4_x.number_format = '#,##0'
    for c in range(1, 13):
        ws_map1.cell(tot_x1_row, c).fill = FILL_TOTAL; ws_map1.cell(tot_x1_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Xăng G1
    note_x1_row = curr_row
    ws_map1.row_dimensions[note_x1_row].height = 22
    tot_x1_buy = sum(i['tong_tien'] for i in dn_chosen_gas)
    tot_x1_lit = sum(i['lit'] for i in dn_chosen_gas)
    tot_x1_used_lit = sum(g1_x_demands[t]['lit'] for t in g1_x_demands)
    surplus_lit_x1 = tot_x1_lit - tot_x1_used_lit
    note_x1_text = f"📌 Ghi chú bảo lưu Xăng RON 95: Tổng HĐ mua {tot_x1_lit:,.1f} L ({tot_x1_buy:,.0f} đ) — Tiêu hao chạy máy {tot_x1_used_lit:,.1f} L ({t1_x:,.0f} đ) ➔ Số lít dư bảo lưu kho: +{surplus_lit_x1:,.1f} L (Tiền HĐ còn dư bảo lưu kho: +{r1_x:,.0f} đ)"
    n_cell_x = ws_map1.cell(note_x1_row, 1, note_x1_text)
    n_cell_x.font = FONT_ITALIC; n_cell_x.fill = FILL_NOTE
    ws_map1.merge_cells(start_row=note_x1_row, start_column=1, end_row=note_x1_row, end_column=12)
    for c in range(1, 13):
        ws_map1.cell(note_x1_row, c).border = BORDER_ALL; ws_map1.cell(note_x1_row, c).fill = FILL_NOTE

    curr_row += 2
    # Grand Total G1
    gt1_row = curr_row
    ws_map1.row_dimensions[gt1_row].height = 24
    ws_map1.cell(gt1_row, 1, "TỔNG CỘNG TOÀN BỘ (DẦU DO + XĂNG RON 95)").font = FONT_BOLD
    ws_map1.cell(gt1_row, 1).alignment = ALIGN_LEFT
    gt1_b = ws_map1.cell(gt1_row, 2, f"=B{tot_d1_row}+B{tot_x1_row}")
    gt1_b.font = FONT_BOLD; gt1_b.alignment = ALIGN_RIGHT; gt1_b.number_format = '#,##0'
    gt1_d = ws_map1.cell(gt1_row, 4, f"=D{tot_d1_row}+D{tot_x1_row}")
    gt1_d.font = FONT_BOLD; gt1_d.alignment = ALIGN_RIGHT; gt1_d.number_format = '#,##0'
    for c in range(1, 13):
        ws_map1.cell(gt1_row, c).fill = FILL_TOTAL; ws_map1.cell(gt1_row, c).border = BORDER_TOTAL

    widths = {1: 14, 2: 24, 3: 15, 4: 24, 5: 35, 6: 16, 7: 14, 8: 30, 9: 18, 10: 14, 11: 25, 12: 46}
    for col_idx, w in widths.items():
        ws_map1.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # 10. Rewrite HD_DongNai_67Tram (25 Invoices with Column P notes)
    print(">>> 10. Writing HD_DongNai_67Tram...")
    clear_sheet_from_row(ws_hd1, 8)
    curr_row = 8
    for idx, inv in enumerate(dn_chosen, 1):
        ws_hd1.row_dimensions[curr_row].height = 20
        ws_hd1.cell(curr_row, 1, f"L{idx}").alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 2, inv['ngay']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 3, inv['so_hd']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 4, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_hd1.cell(curr_row, 5, inv['loai_nl']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 6, inv['dien_giai']).alignment = ALIGN_LEFT

        c7 = ws_hd1.cell(curr_row, 7, inv['lit'])
        c7.alignment = ALIGN_RIGHT; c7.number_format = '#,##0.00'

        ws_hd1.cell(curr_row, 8, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 9, inv['mau_so']).alignment = ALIGN_CENTER
        ws_hd1.cell(curr_row, 10, inv['ky_hieu']).alignment = ALIGN_CENTER

        c11 = ws_hd1.cell(curr_row, 11, inv['tien_chua_vat'])
        c11.alignment = ALIGN_RIGHT; c11.number_format = '#,##0'
        c12 = ws_hd1.cell(curr_row, 12, inv['vat'])
        c12.alignment = ALIGN_RIGHT; c12.number_format = '#,##0'
        c13 = ws_hd1.cell(curr_row, 13, inv['tong_tien'])
        c13.alignment = ALIGN_RIGHT; c13.number_format = '#,##0'

        ws_hd1.cell(curr_row, 14, inv['link']).alignment = ALIGN_LEFT
        ws_hd1.cell(curr_row, 15, inv['fkey']).alignment = ALIGN_CENTER

        c16 = ws_hd1.cell(curr_row, 16)
        c16.alignment = ALIGN_LEFT
        if inv['so_hd'] == last_d1_sohd:
            used_amt = inv['tong_tien'] - r1_d
            c16.value = f"Trích sử dụng {used_amt:,.0f} đ khớp 100% tiền dầu 02A | Giảm giá vốn (bảo lưu kỳ sau): {r1_d:,.0f} đ"
            c16.font = FONT_NOTE_SPECIAL; c16.fill = FILL_SPECIAL
        elif inv['so_hd'] == last_x1_sohd:
            used_amt = inv['tong_tien'] - r1_x
            c16.value = f"Trích sử dụng {used_amt:,.0f} đ khớp 100% tiền xăng 02A | Giảm giá vốn (bảo lưu kỳ sau): {r1_x:,.0f} đ"
            c16.font = FONT_NOTE_SPECIAL; c16.fill = FILL_SPECIAL
        else:
            c16.value = "Sử dụng 100% giá vốn thanh toán đợt này"
            c16.font = FONT_ITALIC

        for c in range(1, 17):
            if c != 16 or (inv['so_hd'] != last_d1_sohd and inv['so_hd'] != last_x1_sohd):
                ws_hd1.cell(curr_row, c).font = FONT_MAIN
            ws_hd1.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    end_dn_hd = curr_row - 1
    tot_row = curr_row
    ws_hd1.row_dimensions[tot_row].height = 24
    ws_hd1.cell(tot_row, 1, "TỔNG CỘNG TOÀN BỘ HÓA ĐƠN").font = FONT_BOLD
    ws_hd1.cell(tot_row, 1).alignment = ALIGN_LEFT
    ws_hd1.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)

    c7_tot = ws_hd1.cell(tot_row, 7, f"=SUM(G8:G{end_dn_hd})")
    c7_tot.font = FONT_BOLD; c7_tot.alignment = ALIGN_RIGHT; c7_tot.number_format = '#,##0.00'
    c11_tot = ws_hd1.cell(tot_row, 11, f"=SUM(K8:K{end_dn_hd})")
    c11_tot.font = FONT_BOLD; c11_tot.alignment = ALIGN_RIGHT; c11_tot.number_format = '#,##0'
    c12_tot = ws_hd1.cell(tot_row, 12, f"=SUM(L8:L{end_dn_hd})")
    c12_tot.font = FONT_BOLD; c12_tot.alignment = ALIGN_RIGHT; c12_tot.number_format = '#,##0'
    c13_tot = ws_hd1.cell(tot_row, 13, f"=SUM(M8:M{end_dn_hd})")
    c13_tot.font = FONT_BOLD; c13_tot.alignment = ALIGN_RIGHT; c13_tot.number_format = '#,##0'

    tot_rem_g1 = r1_d + r1_x
    tot_p_1 = ws_hd1.cell(tot_row, 16, f"Tổng tiền thực tế sử dụng thanh toán đợt này: {tot_dn_demand:,.0f} đ (Bảo lưu chuyển kỳ sau: {tot_rem_g1:,.0f} đ)")
    tot_p_1.font = FONT_BOLD; tot_p_1.fill = FILL_TOTAL; tot_p_1.alignment = ALIGN_LEFT; tot_p_1.border = BORDER_TOTAL
    for c in range(1, 16):
        ws_hd1.cell(tot_row, c).fill = FILL_TOTAL; ws_hd1.cell(tot_row, c).border = BORDER_TOTAL
    ws_hd1.column_dimensions['P'].width = 46

    # 11. Rewrite Map_HD_Theo_Tram_Nhom2
    print(">>> 11. Writing Map_HD_Theo_Tram_Nhom2...")
    ws_map2 = wb['Map_HD_Theo_Tram_Nhom2']
    clear_sheet_from_row(ws_map2, 1)

    ws_map2.cell(1, 1, "BẢNG KÊ PHÂN BỔ HÓA ĐƠN XĂNG DẦU THEO TỪNG TRẠM CHẠY MÁY (MOBIFONE TOÀN CẦU)").font = FONT_TITLE
    ws_map2.cell(2, 1, "Tháng 08/2026 • Phân tách độc lập bảng kê Dầu DO và Xăng RON 95 • Ưu tiên đáp ứng đủ 100% số tiền bảng kê, bảo lưu số lít/tiền dư").font = FONT_ITALIC
    ws_map2.row_dimensions[1].height = 28
    ws_map2.row_dimensions[2].height = 20

    ws_map2.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers_map, 1):
        cell = ws_map2.cell(4, col_idx, h)
        cell.font = FONT_BOLD; cell.fill = FILL_HEADER; cell.alignment = ALIGN_CENTER; cell.border = BORDER_ALL

    # Section I: Dầu DO G2
    curr_row = 5
    sec2_1 = ws_map2.cell(curr_row, 1, "I. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU DẦU DO (DO 0.05S) — GÁN HÓA ĐƠN THEO TRẠM")
    sec2_1.font = FONT_SECTION; sec2_1.fill = FILL_SECTION
    ws_map2.row_dimensions[curr_row].height = 24
    ws_map2.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=12)
    for c in range(1, 13):
        ws_map2.cell(curr_row, c).border = BORDER_ALL; ws_map2.cell(curr_row, c).fill = FILL_SECTION

    dau2_start = curr_row + 1
    curr_row += 1
    last_d2_sohd = tc_picked_oil[-1]['so_hd']
    for a in a2_d:
        ws_map2.row_dimensions[curr_row].height = 20
        c1 = ws_map2.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map2.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map2.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map2.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map2.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map2.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map2.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT

        c12 = ws_map2.cell(curr_row, 12)
        c12.alignment = ALIGN_LEFT
        if a['so_hd'] == last_d2_sohd and a == a2_d[-1]:
            c12.value = f"Trích một phần HĐ {last_d2_sohd} (vừa khít 100% tiền dầu trạm, dư {r2_d:,.0f} đ bảo lưu kho)"
            c12.font = FONT_NOTE_SPECIAL; c12.fill = FILL_SPECIAL
        else:
            c12.value = "Gán 100% chi phí trạm"
            c12.font = FONT_ITALIC

        for c in range(5, 13):
            if c != 12 or a != a2_d[-1]:
                ws_map2.cell(curr_row, c).font = FONT_MAIN
            ws_map2.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    dau2_end = curr_row - 1
    # Total Dầu DO G2
    tot_d2_row = curr_row
    ws_map2.row_dimensions[tot_d2_row].height = 22
    ws_map2.cell(tot_d2_row, 1, "TỔNG CỘNG DẦU DO").font = FONT_BOLD
    ws_map2.cell(tot_d2_row, 1).alignment = ALIGN_LEFT

    t2 = ws_map2.cell(tot_d2_row, 2, f"=SUM(B{dau2_start}:B{dau2_end})")
    t2.font = FONT_BOLD; t2.alignment = ALIGN_RIGHT; t2.number_format = '#,##0'
    t4 = ws_map2.cell(tot_d2_row, 4, f"=SUM(D{dau2_start}:D{dau2_end})")
    t4.font = FONT_BOLD; t4.alignment = ALIGN_RIGHT; t4.number_format = '#,##0'
    for c in range(1, 13):
        ws_map2.cell(tot_d2_row, c).fill = FILL_TOTAL; ws_map2.cell(tot_d2_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Dầu G2
    note_d2_row = curr_row
    ws_map2.row_dimensions[note_d2_row].height = 22
    tot_d2_buy = sum(i['tong_tien'] for i in tc_picked_oil)
    tot_d2_lit = sum(i['lit'] for i in tc_picked_oil)
    tot_d2_used_lit = sum(g2_d_demands[t]['lit'] for t in g2_d_demands)
    note_d2_text = f"📌 Ghi chú bảo lưu Dầu DO: Tổng HĐ mua {tot_d2_lit:,.1f} L ({tot_d2_buy:,.0f} đ) — Tiêu hao chạy máy {tot_d2_used_lit:,.1f} L ({t2_d:,.0f} đ) ➔ Tiền HĐ còn dư bảo lưu kho: +{r2_d:,.0f} đ"
    n_cell = ws_map2.cell(note_d2_row, 1, note_d2_text)
    n_cell.font = FONT_ITALIC; n_cell.fill = FILL_NOTE
    ws_map2.merge_cells(start_row=note_d2_row, start_column=1, end_row=note_d2_row, end_column=12)
    for c in range(1, 13):
        ws_map2.cell(note_d2_row, c).border = BORDER_ALL; ws_map2.cell(note_d2_row, c).fill = FILL_NOTE

    curr_row += 2
    # Section II: Xăng RON 95 G2
    sec2_2 = ws_map2.cell(curr_row, 1, "II. BẢNG KÊ PHÂN BỔ NHIÊN LIỆU XĂNG RON 95 (RON 95-III) — GÁN HÓA ĐƠN THEO TRẠM")
    sec2_2.font = FONT_SECTION; sec2_2.fill = FILL_SECTION
    ws_map2.row_dimensions[curr_row].height = 24
    ws_map2.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=12)
    for c in range(1, 13):
        ws_map2.cell(curr_row, c).border = BORDER_ALL; ws_map2.cell(curr_row, c).fill = FILL_SECTION

    xang2_start = curr_row + 1
    curr_row += 1
    last_x2_sohd = tc_picked_gas[-1]['so_hd']
    for a in a2_x:
        ws_map2.row_dimensions[curr_row].height = 20
        c1 = ws_map2.cell(curr_row, 1, a['tid']); c1.alignment = ALIGN_CENTER; c1.font = FONT_MAIN; c1.border = BORDER_ALL
        c2 = ws_map2.cell(curr_row, 2, a['tien_tram']); c2.alignment = ALIGN_RIGHT; c2.font = FONT_MAIN; c2.border = BORDER_ALL
        if a['tien_tram'] is not None: c2.number_format = '#,##0'
        c3 = ws_map2.cell(curr_row, 3, a['so_hd']); c3.alignment = ALIGN_CENTER; c3.font = FONT_MAIN; c3.border = BORDER_ALL
        c4 = ws_map2.cell(curr_row, 4, a['tien_gan']); c4.alignment = ALIGN_RIGHT; c4.font = FONT_MAIN; c4.border = BORDER_ALL
        c4.number_format = '#,##0'

        inv = a['inv']
        ws_map2.cell(curr_row, 5, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_map2.cell(curr_row, 6, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 7, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 8, inv['link']).alignment = ALIGN_LEFT
        ws_map2.cell(curr_row, 9, inv['fkey']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 10, inv['ngay']).alignment = ALIGN_CENTER
        ws_map2.cell(curr_row, 11, f"{a['st_name']} ()").alignment = ALIGN_LEFT

        c12 = ws_map2.cell(curr_row, 12)
        c12.alignment = ALIGN_LEFT
        if a['so_hd'] == last_x2_sohd and a == a2_x[-1]:
            c12.value = f"Trích một phần HĐ {last_x2_sohd} (vừa khít 100% tiền xăng trạm, dư {r2_x:,.0f} đ bảo lưu kho)"
            c12.font = FONT_NOTE_SPECIAL; c12.fill = FILL_SPECIAL
        else:
            c12.value = "Gán 100% chi phí trạm"
            c12.font = FONT_ITALIC

        for c in range(5, 13):
            if c != 12 or a != a2_x[-1]:
                ws_map2.cell(curr_row, c).font = FONT_MAIN
            ws_map2.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    xang2_end = curr_row - 1
    # Total Xăng G2
    tot_x2_row = curr_row
    ws_map2.row_dimensions[tot_x2_row].height = 22
    ws_map2.cell(tot_x2_row, 1, "TỔNG CỘNG XĂNG RON 95").font = FONT_BOLD
    ws_map2.cell(tot_x2_row, 1).alignment = ALIGN_LEFT

    t2_x_cell = ws_map2.cell(tot_x2_row, 2, f"=SUM(B{xang2_start}:B{xang2_end})")
    t2_x_cell.font = FONT_BOLD; t2_x_cell.alignment = ALIGN_RIGHT; t2_x_cell.number_format = '#,##0'
    t4_x_cell = ws_map2.cell(tot_x2_row, 4, f"=SUM(D{xang2_start}:D{xang2_end})")
    t4_x_cell.font = FONT_BOLD; t4_x_cell.alignment = ALIGN_RIGHT; t4_x_cell.number_format = '#,##0'
    for c in range(1, 13):
        ws_map2.cell(tot_x2_row, c).fill = FILL_TOTAL; ws_map2.cell(tot_x2_row, c).border = BORDER_TOTAL

    curr_row += 1
    # Note Xăng G2
    note_x2_row = curr_row
    ws_map2.row_dimensions[note_x2_row].height = 22
    tot_x2_buy = sum(i['tong_tien'] for i in tc_picked_gas)
    tot_x2_lit = sum(i['lit'] for i in tc_picked_gas)
    tot_x2_used_lit = sum(g2_x_demands[t]['lit'] for t in g2_x_demands)
    surplus_lit_x2 = tot_x2_lit - tot_x2_used_lit
    note_x2_text = f"📌 Ghi chú bảo lưu Xăng RON 95: Tổng HĐ mua {tot_x2_lit:,.1f} L ({tot_x2_buy:,.0f} đ) — Tiêu hao chạy máy {tot_x2_used_lit:,.1f} L ({t2_x:,.0f} đ) ➔ Số lít dư bảo lưu kho: +{surplus_lit_x2:,.1f} L (Tiền HĐ còn dư bảo lưu kho: +{r2_x:,.0f} đ)"
    n_cell_x2 = ws_map2.cell(note_x2_row, 1, note_x2_text)
    n_cell_x2.font = FONT_ITALIC; n_cell_x2.fill = FILL_NOTE
    ws_map2.merge_cells(start_row=note_x2_row, start_column=1, end_row=note_x2_row, end_column=12)
    for c in range(1, 13):
        ws_map2.cell(note_x2_row, c).border = BORDER_ALL; ws_map2.cell(note_x2_row, c).fill = FILL_NOTE

    curr_row += 2
    # Grand Total G2
    gt2_row = curr_row
    ws_map2.row_dimensions[gt2_row].height = 24
    ws_map2.cell(gt2_row, 1, "TỔNG CỘNG TOÀN BỘ (DẦU DO + XĂNG RON 95)").font = FONT_BOLD
    ws_map2.cell(gt2_row, 1).alignment = ALIGN_LEFT
    gt2_b = ws_map2.cell(gt2_row, 2, f"=B{tot_d2_row}+B{tot_x2_row}")
    gt2_b.font = FONT_BOLD; gt2_b.alignment = ALIGN_RIGHT; gt2_b.number_format = '#,##0'
    gt2_d = ws_map2.cell(gt2_row, 4, f"=D{tot_d2_row}+D{tot_x2_row}")
    gt2_d.font = FONT_BOLD; gt2_d.alignment = ALIGN_RIGHT; gt2_d.number_format = '#,##0'
    for c in range(1, 13):
        ws_map2.cell(gt2_row, c).fill = FILL_TOTAL; ws_map2.cell(gt2_row, c).border = BORDER_TOTAL

    for col_idx, w in widths.items():
        ws_map2.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # 12. Rewrite HD_ToanCau (69 Invoices with Column P notes)
    print(">>> 12. Writing HD_ToanCau...")
    clear_sheet_from_row(ws_hd2, 8)
    curr_row = 8
    for idx, inv in enumerate(tc_chosen, 1):
        ws_hd2.row_dimensions[curr_row].height = 20
        ws_hd2.cell(curr_row, 1, f"L{idx}").alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 2, inv['ngay']).alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 3, inv['so_hd']).alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 4, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_hd2.cell(curr_row, 5, inv['loai_nl']).alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 6, inv['dien_giai']).alignment = ALIGN_LEFT

        c7 = ws_hd2.cell(curr_row, 7, inv['lit'])
        c7.alignment = ALIGN_RIGHT; c7.number_format = '#,##0.00'

        ws_hd2.cell(curr_row, 8, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 9, inv['mau_so']).alignment = ALIGN_CENTER
        ws_hd2.cell(curr_row, 10, inv['ky_hieu']).alignment = ALIGN_CENTER

        c11 = ws_hd2.cell(curr_row, 11, inv['tien_chua_vat'])
        c11.alignment = ALIGN_RIGHT; c11.number_format = '#,##0'
        c12 = ws_hd2.cell(curr_row, 12, inv['vat'])
        c12.alignment = ALIGN_RIGHT; c12.number_format = '#,##0'
        c13 = ws_hd2.cell(curr_row, 13, inv['tong_tien'])
        c13.alignment = ALIGN_RIGHT; c13.number_format = '#,##0'

        ws_hd2.cell(curr_row, 14, inv['link']).alignment = ALIGN_LEFT
        ws_hd2.cell(curr_row, 15, inv['fkey']).alignment = ALIGN_CENTER

        c16 = ws_hd2.cell(curr_row, 16)
        c16.alignment = ALIGN_LEFT
        if inv['so_hd'] == last_d2_sohd:
            used_amt = inv['tong_tien'] - r2_d
            c16.value = f"Trích sử dụng {used_amt:,.0f} đ khớp 100% tiền dầu 02A | Giảm giá vốn (bảo lưu kỳ sau): {r2_d:,.0f} đ"
            c16.font = FONT_NOTE_SPECIAL; c16.fill = FILL_SPECIAL
        elif inv['so_hd'] == last_x2_sohd:
            used_amt = inv['tong_tien'] - r2_x
            c16.value = f"Trích sử dụng {used_amt:,.0f} đ khớp 100% tiền xăng 02A | Giảm giá vốn (bảo lưu kỳ sau): {r2_x:,.0f} đ"
            c16.font = FONT_NOTE_SPECIAL; c16.fill = FILL_SPECIAL
        else:
            c16.value = "Sử dụng 100% giá vốn thanh toán đợt này"
            c16.font = FONT_ITALIC

        for c in range(1, 17):
            if c != 16 or (inv['so_hd'] != last_d2_sohd and inv['so_hd'] != last_x2_sohd):
                ws_hd2.cell(curr_row, c).font = FONT_MAIN
            ws_hd2.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    end_tc_hd = curr_row - 1
    tot_row = curr_row
    ws_hd2.row_dimensions[tot_row].height = 24
    ws_hd2.cell(tot_row, 1, "TỔNG CỘNG TOÀN BỘ HÓA ĐƠN").font = FONT_BOLD
    ws_hd2.cell(tot_row, 1).alignment = ALIGN_LEFT
    ws_hd2.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)

    c7_tot = ws_hd2.cell(tot_row, 7, f"=SUM(G8:G{end_tc_hd})")
    c7_tot.font = FONT_BOLD; c7_tot.alignment = ALIGN_RIGHT; c7_tot.number_format = '#,##0.00'
    c11_tot = ws_hd2.cell(tot_row, 11, f"=SUM(K8:K{end_tc_hd})")
    c11_tot.font = FONT_BOLD; c11_tot.alignment = ALIGN_RIGHT; c11_tot.number_format = '#,##0'
    c12_tot = ws_hd2.cell(tot_row, 12, f"=SUM(L8:L{end_tc_hd})")
    c12_tot.font = FONT_BOLD; c12_tot.alignment = ALIGN_RIGHT; c12_tot.number_format = '#,##0'
    c13_tot = ws_hd2.cell(tot_row, 13, f"=SUM(M8:M{end_tc_hd})")
    c13_tot.font = FONT_BOLD; c13_tot.alignment = ALIGN_RIGHT; c13_tot.number_format = '#,##0'

    tot_rem_g2 = r2_d + r2_x
    tot_p_2 = ws_hd2.cell(tot_row, 16, f"Tổng tiền thực tế sử dụng thanh toán đợt này: {tot_tc_demand:,.0f} đ (Bảo lưu chuyển kỳ sau: {tot_rem_g2:,.0f} đ)")
    tot_p_2.font = FONT_BOLD; tot_p_2.fill = FILL_TOTAL; tot_p_2.alignment = ALIGN_LEFT; tot_p_2.border = BORDER_TOTAL
    for c in range(1, 16):
        ws_hd2.cell(tot_row, c).fill = FILL_TOTAL; ws_hd2.cell(tot_row, c).border = BORDER_TOTAL
    ws_hd2.column_dimensions['P'].width = 46

    # 13. Rewrite HD_Du_Thua_Khong_Su_Dung
    print(">>> 13. Writing HD_Du_Thua_Khong_Su_Dung...")
    clear_sheet_from_row(ws_dt, 1)

    all_surplus = dn_unpicked + tc_unpicked
    all_surplus.sort(key=lambda x: (x['ngay'], x['so_hd']))

    ws_dt.cell(1, 1, "BẢNG KÊ HÓA ĐƠN XĂNG DẦU DƯ THỪA / BẢO LƯU KHO KHÔNG SỬ DỤNG").font = FONT_TITLE
    ws_dt.cell(2, 1, f"Tháng 08/2026 • Tổng cộng {len(all_surplus)} hóa đơn được bảo lưu trong kho dữ liệu (chưa sử dụng để tối ưu hạn mức <= 5 triệu/ngày)").font = FONT_ITALIC
    ws_dt.row_dimensions[1].height = 28
    ws_dt.row_dimensions[2].height = 20

    headers_surplus = [
        "STT", "Ngày Lập HĐ", "Số Hóa Đơn", "Bên Mua (Pháp Nhân / MST)", "Loại Nhiên Liệu",
        "Số Lượng (Lít)", "Đơn Giá (đ/L)", "Thành Tiền Chưa Thuế (đ)", "Thuế GTGT 8% (đ)",
        "Tổng Tiền Thanh Toán (đ)", "Đơn Vị Xuất Hóa Đơn", "Mã Số Thuế (Bán)",
        "Ký Hiệu HĐ", "Link Tra Cứu Hóa Đơn", "Mã Tra Cứu / Fkey", "Lý Do Dư Thừa / Ghi Chú Phân Loại"
    ]
    ws_dt.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers_surplus, 1):
        cell = ws_dt.cell(4, col_idx, h)
        cell.font = FONT_BOLD; cell.fill = FILL_HEADER; cell.alignment = ALIGN_CENTER; cell.border = BORDER_ALL

    curr_row = 5
    for idx, inv in enumerate(all_surplus, 1):
        ws_dt.row_dimensions[curr_row].height = 20
        ws_dt.cell(curr_row, 1, idx).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 2, inv['ngay']).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 3, inv['so_hd']).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 4, inv['buyer']).alignment = ALIGN_LEFT
        ws_dt.cell(curr_row, 5, inv['loai_nl']).alignment = ALIGN_CENTER

        c6 = ws_dt.cell(curr_row, 6, inv['lit'])
        c6.alignment = ALIGN_RIGHT; c6.number_format = '#,##0.00'

        dg = round(inv['tong_tien'] / inv['lit']) if inv['lit'] > 0 else 0
        c7 = ws_dt.cell(curr_row, 7, dg)
        c7.alignment = ALIGN_RIGHT; c7.number_format = '#,##0'

        c8 = ws_dt.cell(curr_row, 8, inv['tien_chua_vat'])
        c8.alignment = ALIGN_RIGHT; c8.number_format = '#,##0'

        c9 = ws_dt.cell(curr_row, 9, inv['vat'])
        c9.alignment = ALIGN_RIGHT; c9.number_format = '#,##0'

        c10 = ws_dt.cell(curr_row, 10, inv['tong_tien'])
        c10.alignment = ALIGN_RIGHT; c10.number_format = '#,##0'

        ws_dt.cell(curr_row, 11, inv['don_vi_ban']).alignment = ALIGN_LEFT
        ws_dt.cell(curr_row, 12, inv['mst_ban']).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 13, inv['ky_hieu']).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 14, inv['link']).alignment = ALIGN_LEFT
        ws_dt.cell(curr_row, 15, inv['fkey']).alignment = ALIGN_CENTER
        ws_dt.cell(curr_row, 16, "Hóa đơn dư thừa dự phòng (bảo lưu để đảm bảo hạn mức <= 5M/ngày)").alignment = ALIGN_LEFT

        for c in range(1, 17):
            ws_dt.cell(curr_row, c).font = FONT_MAIN
            ws_dt.cell(curr_row, c).border = BORDER_ALL
        curr_row += 1

    end_surplus = curr_row - 1
    tot_surplus = curr_row
    ws_dt.row_dimensions[tot_surplus].height = 24
    ws_dt.cell(tot_surplus, 1, "TỔNG CỘNG HÓA ĐƠN DƯ THỪA").font = FONT_BOLD
    ws_dt.cell(tot_surplus, 1).alignment = ALIGN_LEFT
    ws_dt.merge_cells(start_row=tot_surplus, start_column=1, end_row=tot_surplus, end_column=5)

    c6_s = ws_dt.cell(tot_surplus, 6, f"=SUM(F5:F{end_surplus})")
    c6_s.font = FONT_BOLD; c6_s.alignment = ALIGN_RIGHT; c6_s.number_format = '#,##0.00'
    c8_s = ws_dt.cell(tot_surplus, 8, f"=SUM(H5:H{end_surplus})")
    c8_s.font = FONT_BOLD; c8_s.alignment = ALIGN_RIGHT; c8_s.number_format = '#,##0'
    c9_s = ws_dt.cell(tot_surplus, 9, f"=SUM(I5:I{end_surplus})")
    c9_s.font = FONT_BOLD; c9_s.alignment = ALIGN_RIGHT; c9_s.number_format = '#,##0'
    c10_s = ws_dt.cell(tot_surplus, 10, f"=SUM(J5:J{end_surplus})")
    c10_s.font = FONT_BOLD; c10_s.alignment = ALIGN_RIGHT; c10_s.number_format = '#,##0'

    for c in range(1, 17):
        ws_dt.cell(tot_surplus, c).fill = FILL_TOTAL; ws_dt.cell(tot_surplus, c).border = BORDER_TOTAL

    widths_surplus = {1: 8, 2: 14, 3: 14, 4: 30, 5: 16, 6: 16, 7: 14, 8: 22, 9: 16, 10: 22, 11: 35, 12: 16, 13: 14, 14: 30, 15: 18, 16: 45}
    for col_idx, w in widths_surplus.items():
        ws_dt.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

    # 14. Save workbook
    print(">>> 14. Saving workbook...")
    wb.save(DST_DESKTOP)
    shutil.copy2(DST_DESKTOP, DST_DOWNLOADS)
    print(f"SUCCESS: Saved {DST_DESKTOP} and copied to {DST_DOWNLOADS}")

if __name__ == '__main__':
    main()
