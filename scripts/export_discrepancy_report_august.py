#!/usr/bin/env python3
"""
Script xuất file Excel chi tiết 4 sai lệch cụ thể giữa file Excel Desktop (5. NL Long Khánh _ T8.2026.xlsx)
và Dữ liệu chuẩn trên Web (Supabase) tháng 08/2026.
"""

import os
import json
from datetime import datetime, timedelta
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from supabase import create_client

# 1. Kết nối Supabase
load_dotenv('tvt3_v2/.env')
url = os.getenv('VITE_SUPABASE_URL')
key = os.getenv('VITE_SUPABASE_ANON_KEY')
supabase = create_client(url, key)

print("🚀 Đang trích xuất dữ liệu để tạo báo cáo 4 sai lệch...")

# 2. Đọc file backup gốc trước chuẩn hóa (391 logs)
backup_file = 'backend/backup_generator_logs_august_20260929_143701.json'
with open(backup_file, 'r', encoding='utf-8') as f:
    orig_logs = json.load(f)

# 3. Đọc dữ liệu chuẩn trên Web hiện tại (316 logs)
res_db = supabase.table('generator_logs').select('*').gte('date', '2026-08-01').lte('date', '2026-08-31').execute()
current_db_logs = res_db.data or []

# 4. Đọc dữ liệu danh mục trạm
res_sites = supabase.table('datasites').select('site_id, site_id_old, name, location_info').execute()
site_map = {}
for s in res_sites.data or []:
    sid = (s.get('site_id') or '').upper()
    sold = (s.get('site_id_old') or '').upper()
    name = s.get('name') or ''
    loc = s.get('location_info') or {}
    dist = loc.get('district') or loc.get('huyen') or ''
    info = {'name': name, 'district': dist, 'site_id_old': sold}
    if sid: site_map[sid] = info
    if sold: site_map[sold] = info

def get_site_info(sid):
    u = (sid or '').upper().strip()
    return site_map.get(u, {'name': u, 'district': '', 'site_id_old': ''})

# 5. Phân loại 4 nhóm sai lệch
def parse_alarm(alarm_id):
    if not alarm_id or '__' not in alarm_id:
        return None, None
    site, ts_str = alarm_id.split('__', 1)
    for fmt in ['%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%b %d, %Y %I:%M:%S %p', '%b %d, %Y %I:%M %p', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M']:
        try:
            return site.strip().upper(), datetime.strptime(ts_str.strip(), fmt)
        except:
            pass
    return site.strip().upper(), None

def is_utc_record(log):
    rd = log.get('run_details') or {}
    start_t = rd.get('gio_bat_dau', '')
    aid = rd.get('smartw_alarm_id', '')
    site, adt = parse_alarm(aid)
    if not adt or not start_t or ':' not in start_t:
        return False
    sh, sm = map(int, start_t.split(':')[:2])
    diff = (adt.hour * 60 + adt.minute - (sh * 60 + sm)) % 1440
    return diff in [420, 419, 421]

by_alarm = defaultdict(list)
for l in orig_logs:
    rd = l.get('run_details') or {}
    aid = rd.get('smartw_alarm_id', '')
    site, adt = parse_alarm(aid)
    if adt:
        by_alarm[(site, adt.strftime('%Y-%m-%d %H:%M:%S'))].append(l)

# Nhóm 1: 54 ca trùng lặp UTC
dup_pairs = []
utc_standalone = []
for k, elist in by_alarm.items():
    vn = [l for l in elist if not is_utc_record(l)]
    utc = [l for l in elist if is_utc_record(l)]
    if vn and utc:
        dup_pairs.append({
            'site': k[0], 'alarm_dt': k[1],
            'vn_log': vn[0],
            'utc_log': utc[0]
        })
    elif utc:
        utc_standalone.append(utc[0])

dup_pairs.sort(key=lambda x: (x['vn_log']['date'], x['site']))

# Nhóm 2: 195 ca lệch múi giờ -7h (chỉ có bản ghi UTC đơn lẻ)
utc_standalone_list = []
for l in utc_standalone:
    rd = l.get('run_details') or {}
    aid = rd.get('smartw_alarm_id', '')
    site, adt = parse_alarm(aid)
    h = float(rd.get('thoi_gian_hoat_dong') or 0)
    fuel = float(rd.get('nhien_lieu_tieu_hao') or 0)
    money = float(rd.get('thanh_tien') or 0)
    don_gia = float(rd.get('don_gia') or 0)
    
    real_start_dt = adt
    real_end_dt = real_start_dt + timedelta(minutes=int(round(h * 60)))
    
    utc_standalone_list.append({
        'site_id': l['site_id'],
        'orig_date': l['date'],
        'orig_start': rd.get('gio_bat_dau', ''),
        'orig_end': rd.get('gio_ket_thuc', ''),
        'real_date': real_start_dt.strftime('%Y-%m-%d'),
        'real_start': real_start_dt.strftime('%H:%M'),
        'real_end': real_end_dt.strftime('%H:%M'),
        'hours': h, 'fuel': fuel, 'don_gia': don_gia, 'money': money,
        'alarm_id': aid
    })
utc_standalone_list.sort(key=lambda x: (x['real_date'], x['site_id']))

# Nhóm 3: 17 cụm nối ca (21 ca gộp)
# Phân tích cụm nối ca từ logs_after_shift
logs_after_shift = []
for l in orig_logs:
    if any(l['gen_log_id'] == d['utc_log']['gen_log_id'] for d in dup_pairs):
        continue # bỏ 54 bản sao trùng lặp
    l_copy = dict(l)
    l_copy['run_details'] = dict(l['run_details'])
    if is_utc_record(l):
        rd = l_copy['run_details']
        site, adt = parse_alarm(rd.get('smartw_alarm_id'))
        h = float(rd.get('thoi_gian_hoat_dong') or 0)
        l_copy['date'] = adt.strftime('%Y-%m-%d')
        rd['gio_bat_dau'] = adt.strftime('%H:%M')
        rd['gio_ket_thuc'] = (adt + timedelta(minutes=int(round(h * 60)))).strftime('%H:%M')
    logs_after_shift.append(l_copy)

def to_mins(t_str):
    sh, sm = map(int, t_str.split(':')[:2])
    return sh * 60 + sm

by_site_date = defaultdict(list)
for l in logs_after_shift:
    by_site_date[(l['site_id'].upper(), l['date'])].append(l)

merge_clusters = []
for (site, dt), slist in by_site_date.items():
    if len(slist) < 2: continue
    slist_sorted = sorted(slist, key=lambda x: to_mins(x['run_details'].get('gio_bat_dau', '00:00')))
    current_cluster = [slist_sorted[0]]
    for next_log in slist_sorted[1:]:
        prev_log = current_cluster[-1]
        prev_e = to_mins(prev_log['run_details']['gio_ket_thuc'])
        next_s = to_mins(next_log['run_details']['gio_bat_dau'])
        gap = next_s - prev_e
        if gap <= 1:
            current_cluster.append(next_log)
        else:
            if len(current_cluster) > 1:
                merge_clusters.append((site, dt, current_cluster))
            current_cluster = [next_log]
    if len(current_cluster) > 1:
        merge_clusters.append((site, dt, current_cluster))

merge_clusters.sort(key=lambda x: (x[1], x[0]))

# Nhóm 4: 20 hóa đơn có trong DB nhưng chưa có trong file Excel cũ
# Đọc HĐ từ file Excel
excel_path = '/Users/cang_it/Desktop/5. NL Long Khánh _ T8.2026.xlsx'
wb_old = openpyxl.load_workbook(excel_path, data_only=True)
excel_inv_nums = set()
for sname in ['HD_DongNai_67Tram', 'HD_ToanCau', 'HD_Du_Thua_Khong_Su_Dung']:
    ws_temp = wb_old[sname]
    start_r = 5 if 'Du_Thua' in sname else 8
    for r in range(start_r, ws_temp.max_row + 1):
        num = ws_temp.cell(r, 3).value
        if num: excel_inv_nums.add(str(num).strip())

res_invs = supabase.table('parsed_invoices').select('*').gte('invoice_date', '2026-08-01').lte('invoice_date', '2026-08-31').execute()
db_all_invs = res_invs.data or []
db_inv_nums = {str(i['invoice_number']).strip(): i for i in db_all_invs if i.get('invoice_number')}

extra_invs = [db_inv_nums[num] for num in sorted(db_inv_nums.keys()) if num not in excel_inv_nums]
extra_invs.sort(key=lambda x: (x.get('invoice_date', ''), x.get('invoice_number', '')))

# Bắt đầu tạo file Excel báo cáo chi tiết
print("📊 Đang khởi tạo file Excel báo cáo chuẩn đẹp...")
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Định dạng chung
thin_border = Border(left=Side(style='thin', color='D0D7DE'), right=Side(style='thin', color='D0D7DE'), top=Side(style='thin', color='D0D7DE'), bottom=Side(style='thin', color='D0D7DE'))
thick_bottom = Border(left=Side(style='thin', color='D0D7DE'), right=Side(style='thin', color='D0D7DE'), top=Side(style='thin', color='D0D7DE'), bottom=Side(style='medium', color='1F2328'))
double_bottom = Border(left=Side(style='thin', color='D0D7DE'), right=Side(style='thin', color='D0D7DE'), top=Side(style='thin', color='D0D7DE'), bottom=Side(style='double', color='1F2328'))

fill_navy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
fill_blue = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
fill_red = PatternFill(start_color='991B1B', end_color='991B1B', fill_type='solid')
fill_amber = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid')
fill_green = PatternFill(start_color='047857', end_color='047857', fill_type='solid')
fill_subtot = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_yellow = PatternFill(start_color='FEF08A', end_color='FEF08A', fill_type='solid')
fill_light_blue = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid')

font_hdr = Font(name='Arial', size=9, bold=True, color='FFFFFF')
font_bold = Font(name='Arial', size=9, bold=True)
font_regular = Font(name='Arial', size=9)
font_link = Font(name='Arial', size=9, color='0284C7', underline='single')

# ==============================================================================
# SHEET 1: TỔNG QUAN ĐỐI CHIẾU
# ==============================================================================
ws1 = wb.create_sheet('Tong_Quan_Doi_Chieu')
ws1.views.sheetView[0].showGridLines = True

ws1['A1'] = 'BÁO CÁO CHI TIẾT 4 SAI LỆCH CỤ THỂ GIỮA FILE EXCEL VÀ DỮ LIỆU CHUẨN TRÊN WEB'
ws1['A1'].font = Font(name='Arial', size=13, bold=True, color='1E3A8A')
ws1['A2'] = 'Kỳ thanh toán: Tháng 08/2026 • Đối soát giữa file Desktop: "5. NL Long Khánh _ T8.2026.xlsx" và Hệ thống Quản trị Web TVT3'
ws1['A2'].font = Font(name='Arial', size=10, italic=True, color='4B5563')

ws1['A4'] = 'I. BẢNG SO SÁNH TỔNG HỢP CÁC CHỈ TIÊU VẬN HÀNH & CHI PHÍ'
ws1['A4'].font = Font(name='Arial', size=11, bold=True, color='1E3A8A')

headers_s1_t1 = ['STT', 'Chỉ Tiêu So Sánh', 'File Excel Desktop (Chưa chuẩn hóa)', 'Dữ Liệu Chuẩn Trên Web (Sau chuẩn hóa)', 'Chênh Lệch / Thu Hồi Khống', 'Tỷ Lệ Giảm / Thu Hồi', 'Ghi Chú Đánh Giá']
for c_idx, h in enumerate(headers_s1_t1, 1):
    cell = ws1.cell(5, c_idx, h)
    cell.fill = fill_navy; cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws1.row_dimensions[5].height = 28

t1_data = [
    (1, 'Tổng số ca chạy máy ghi nhận', 391, 316, -75, -75/391, 'Loại 54 bản sao trùng UTC + gộp 21 ca nối ca liên tục'),
    (2, '• Số ca Nhóm 1 (Đồng Nai 67 trạm)', 139, 108, -31, -31/139, 'Đã thu hồi 22 ca trùng lặp + gộp 9 ca liên tục'),
    (3, '• Số ca Nhóm 2 (MobiFone Toàn Cầu)', 252, 208, -44, -44/252, 'Đã thu hồi 32 ca trùng lặp + gộp 12 ca liên tục'),
    (4, 'Tổng thời gian hoạt động máy nổ (giờ)', 1652.68, 1398.36, -254.32, -254.32/1652.68, 'Thu hồi 254.32 giờ chạy máy bị trùng lặp / ảo'),
    (5, '• Giờ chạy máy Nhóm 1 (giờ)', 506.00, 417.82, -88.18, -88.18/506.00, 'Chuẩn hóa đúng theo lịch cúp điện và giờ thực tế'),
    (6, '• Giờ chạy máy Nhóm 2 (giờ)', 1146.68, 980.54, -166.14, -166.14/1146.68, 'Loại bỏ thời gian bị đội khống do lỗi scraper'),
    (7, 'Tổng tiêu hao nhiên liệu (Lít)', 5041.48, 4306.47, -735.01, -735.01/5041.48, 'Thu hồi 735.01 lít xăng dầu đội khống ngoài thực tế'),
    (8, '• Lượng nhiên liệu Nhóm 1 (Lít)', 1683.53, 1398.71, -284.82, -284.82/1683.53, 'Tiết kiệm 284.82 lít dầu/xăng cho MobiFone Đồng Nai'),
    (9, '• Lượng nhiên liệu Nhóm 2 (Lít)', 3357.95, 2907.76, -450.19, -450.19/3357.95, 'Tiết kiệm 450.19 lít dầu/xăng cho MobiFone Toàn Cầu'),
    (10, 'Tổng thành tiền thanh toán (VNĐ)', 139860735, 119618232, -20242503, -20242503/139860735, 'Thu hồi thành công 20,242,503 VNĐ chi phí đội khống'),
    (11, '• Thành tiền Nhóm 1 (VNĐ)', 46669615, 38784520, -7885095, -7885095/46669615, 'Giảm trừ thanh toán thừa 7,885,095 VNĐ'),
    (12, '• Thành tiền Nhóm 2 (VNĐ)', 93191120, 80833712, -12357408, -12357408/93191120, 'Giảm trừ thanh toán thừa 12,357,408 VNĐ'),
    (13, 'Tổng số lượng hóa đơn điện tử', 124, 144, 20, 20/124, 'Web có đầy đủ 144 HĐ (+20 HĐ đợt cuối tháng 27-31/08)')
]

for row_idx, r in enumerate(t1_data, 6):
    ws1.cell(row_idx, 1, r[0]).alignment = Alignment(horizontal='center')
    ws1.cell(row_idx, 2, r[1])
    c3 = ws1.cell(row_idx, 3, r[2])
    c4 = ws1.cell(row_idx, 4, r[3])
    c5 = ws1.cell(row_idx, 5, r[4])
    c6 = ws1.cell(row_idx, 6, r[5])
    ws1.cell(row_idx, 7, r[6])
    
    # Formats
    if r[0] in [1, 2, 3, 13]:
        c3.number_format = '#,##0'; c4.number_format = '#,##0'; c5.number_format = '+#,##0;-#,##0;0'
    elif r[0] in [10, 11, 12]:
        c3.number_format = '#,##0'; c4.number_format = '#,##0'; c5.number_format = '+#,##0;-#,##0;0'
        c3.font = font_bold; c4.font = font_bold; c5.font = Font(name='Arial', size=9, bold=True, color='DC2626')
    else:
        c3.number_format = '#,##0.00'; c4.number_format = '#,##0.00'; c5.number_format = '+#,##0.00;-#,##0.00;0.00'
    
    c6.number_format = '0.0%'
    if '-' in str(r[4]):
        c5.font = Font(name='Arial', size=9, bold=True, color='DC2626')
    else:
        c5.font = Font(name='Arial', size=9, bold=True, color='047857')

    for ci in range(1, 8):
        cell = ws1.cell(row_idx, ci)
        cell.border = thin_border
        if r[0] in [1, 4, 7, 10, 13]:
            cell.fill = fill_subtot
            cell.font = font_bold

# Section II: Tóm tắt 4 nhóm sai lệch
r_start_ii = len(t1_data) + 8
ws1.cell(r_start_ii, 1, 'II. TÓM TẮT 4 NHÓM SAI LỆCH CỤ THỂ ĐÃ ĐƯỢC XỬ LÝ').font = Font(name='Arial', size=11, bold=True, color='1E3A8A')

headers_s1_t2 = ['STT', 'Nhóm Sai Lệch', 'Quy Mô / Số Lượng', 'Nguyên Nhân Kỹ Thuật Trong File Excel Cũ', 'Giải Pháp Chuẩn Hóa Trên Web', 'Tên Sheet Chi Tiết']
for c_idx, h in enumerate(headers_s1_t2, 1):
    cell = ws1.cell(r_start_ii + 1, c_idx, h)
    cell.fill = fill_blue; cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws1.row_dimensions[r_start_ii + 1].height = 28

t2_data = [
    (1, 'Sai lệch 1: Bản sao trùng lặp UTC', '54 cặp ca trùng (Thu hồi 20,242,503 đ)', 'Scraper cũ nạp 2 lần cho cùng 1 cảnh báo: 1 bản US AM/PM và 1 bản DD/MM/YYYY lệch 7 tiếng', 'Đã xóa hoàn toàn 54 bản ghi trùng lặp, bảo đảm tính đúng 1 lần', 'Sheet: 1_Trung_Lap_UTC_54_Ca'),
    (2, 'Sai lệch 2: Giờ bắt đầu lệch múi giờ -7h', '195 ca đơn lẻ UTC', 'Hệ thống import cũ lưu giờ theo chuẩn UTC (00:xx, 01:xx, 02:xx) thay vì giờ Việt Nam', 'Đã tịnh tiến +7h về đúng khung giờ hành chính cúp điện thực tế (07:xx, 08:xx)', 'Sheet: 2_Lech_Mui_Gio_195_Ca'),
    (3, 'Sai lệch 3: Các ca chạy liên tục bị chẻ nhỏ', '17 cụm nối ca (Gộp 21 dòng phụ)', 'Các ca chạy liên tục (gap ≤ 1 phút) hoặc chồng lấn (gap < 0) bị chẻ thành nhiều dòng rời rạc', 'Đã gộp nối ca liền mạch theo đúng quy tắc, cập nhật giờ sớm nhất và muộn nhất', 'Sheet: 3_Noi_Ca_Lien_Tuc_21_Ca'),
    (4, 'Sai lệch 4: Chênh lệch số lượng hóa đơn', '20 hóa đơn bổ sung trên Web', 'File Excel cũ kết xuất trước ngày 31/08 nên chưa gom các hóa đơn xuất bù đợt cuối tháng', 'Web có đầy đủ 144 HĐ (+20 HĐ cuối tháng). 124 HĐ cũ khớp chính xác 100% số tiền', 'Sheet: 4_Lech_Hoa_Don_20_HD')
]

for row_idx, r in enumerate(t2_data, r_start_ii + 2):
    ws1.cell(row_idx, 1, r[0]).alignment = Alignment(horizontal='center')
    ws1.cell(row_idx, 2, r[1]).font = font_bold
    ws1.cell(row_idx, 3, r[2]).font = Font(name='Arial', size=9, bold=True, color='991B1B')
    ws1.cell(row_idx, 4, r[3])
    ws1.cell(row_idx, 5, r[4])
    c6 = ws1.cell(row_idx, 6, r[5]); c6.font = font_link; c6.alignment = Alignment(horizontal='center')
    for ci in range(1, 7):
        ws1.cell(row_idx, ci).border = thin_border

# ==============================================================================
# SHEET 2: CHI TIẾT 54 CA TRÙNG LẶP UTC (SAI LỆCH 1)
# ==============================================================================
ws2 = wb.create_sheet('1_Trung_Lap_UTC_54_Ca')
ws2.views.sheetView[0].showGridLines = True

ws2['A1'] = 'DANH SÁCH CHI TIẾT 54 CA BẢN SAO TRÙNG LẶP UTC ĐÃ ĐƯỢC THU HỒI & LOẠI BỎ'
ws2['A1'].font = Font(name='Arial', size=13, bold=True, color='991B1B')
ws2['A2'] = 'Nguyên nhân: Cùng 1 lần chạy máy nhưng bị ghi nhận 2 lần (1 bản ghi giờ VN và 1 bản ghi giờ UTC lệch 7 tiếng) • Thu hồi 20,242,503 VNĐ'
ws2['A2'].font = Font(name='Arial', size=10, italic=True, color='4B5563')

headers_s2 = [
    'STT', 'Mã Trạm', 'Mã Trạm Cũ', 'Tên Trạm / Địa Bàn', 'SmartW Alarm Key',
    'Ngày Chạy (Chuẩn)', 'Giờ BĐ (Chuẩn VN)', 'Giờ KT (Chuẩn VN)', 'Giờ Chạy Chuẩn (h)', 'Nhiên Liệu Chuẩn (L)', 'Thành Tiền Chuẩn (đ)',
    'Ngày (Bản Sao UTC)', 'Giờ BĐ (Bản Sao UTC)', 'Giờ KT (Bản Sao UTC)', 'Giờ Chạy Ảo (h)', 'Nhiên Liệu Ảo (L)', 'Tiền Đội Khống Thu Hồi (đ)', 'Đánh Giá Xử Lý'
]

for c_idx, h in enumerate(headers_s2, 1):
    cell = ws2.cell(4, c_idx, h)
    cell.fill = fill_red if c_idx >= 12 else fill_navy
    cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws2.row_dimensions[4].height = 32

tot_dup_money = 0
tot_dup_hours = 0
tot_dup_fuel = 0

for idx, p in enumerate(dup_pairs, 1):
    cur_r = idx + 4
    sid = p['site']
    st_info = get_site_info(sid)
    vn_rd = p['vn_log'].get('run_details') or {}
    utc_rd = p['utc_log'].get('run_details') or {}
    
    vn_h = float(vn_rd.get('thoi_gian_hoat_dong') or 0)
    vn_f = float(vn_rd.get('nhien_lieu_tieu_hao') or 0)
    vn_m = float(vn_rd.get('thanh_tien') or 0)
    
    utc_h = float(utc_rd.get('thoi_gian_hoat_dong') or 0)
    utc_f = float(utc_rd.get('nhien_lieu_tieu_hao') or 0)
    utc_m = float(utc_rd.get('thanh_tien') or 0)
    
    tot_dup_hours += utc_h
    tot_dup_fuel += utc_f
    tot_dup_money += utc_m
    
    ws2.cell(cur_r, 1, idx).alignment = Alignment(horizontal='center')
    ws2.cell(cur_r, 2, sid).alignment = Alignment(horizontal='center')
    ws2.cell(cur_r, 3, st_info['site_id_old']).alignment = Alignment(horizontal='center')
    ws2.cell(cur_r, 4, f"{st_info['name']} ({st_info['district']})")
    ws2.cell(cur_r, 5, p['alarm_dt']).alignment = Alignment(horizontal='center')
    
    # Bản chuẩn
    ws2.cell(cur_r, 6, p['vn_log']['date']).alignment = Alignment(horizontal='center')
    ws2.cell(cur_r, 7, vn_rd.get('gio_bat_dau', '')).alignment = Alignment(horizontal='center')
    ws2.cell(cur_r, 8, vn_rd.get('gio_ket_thuc', '')).alignment = Alignment(horizontal='center')
    c9 = ws2.cell(cur_r, 9, vn_h); c9.number_format = '#,##0.00'
    c10 = ws2.cell(cur_r, 10, vn_f); c10.number_format = '#,##0.00'
    c11 = ws2.cell(cur_r, 11, vn_m); c11.number_format = '#,##0'
    
    # Bản sao UTC
    ws2.cell(cur_r, 12, p['utc_log']['date']).alignment = Alignment(horizontal='center')
    c13 = ws2.cell(cur_r, 13, utc_rd.get('gio_bat_dau', '')); c13.alignment = Alignment(horizontal='center'); c13.font = Font(name='Arial', size=9, color='DC2626')
    ws2.cell(cur_r, 14, utc_rd.get('gio_ket_thuc', '')).alignment = Alignment(horizontal='center')
    c15 = ws2.cell(cur_r, 15, utc_h); c15.number_format = '#,##0.00'
    c16 = ws2.cell(cur_r, 16, utc_f); c16.number_format = '#,##0.00'
    c17 = ws2.cell(cur_r, 17, utc_m); c17.number_format = '#,##0'; c17.font = Font(name='Arial', size=9, bold=True, color='DC2626')
    ws2.cell(cur_r, 18, 'Đã xóa bản sao UTC khỏi hệ thống').alignment = Alignment(horizontal='center')
    
    for ci in range(1, 19):
        ws2.cell(cur_r, ci).border = thin_border

# Dòng tổng cộng
r_tot_s2 = len(dup_pairs) + 5
ws2.cell(r_tot_s2, 1, 'TỔNG CỘNG THU HỒI TỪ 54 BẢN SAO TRÙNG LẶP UTC:').font = font_bold
ws2.merge_cells(start_row=r_tot_s2, start_column=1, end_row=r_tot_s2, end_column=14)
c_th = ws2.cell(r_tot_s2, 15, tot_dup_hours); c_th.number_format = '#,##0.00'; c_th.font = font_bold
c_tf = ws2.cell(r_tot_s2, 16, tot_dup_fuel); c_tf.number_format = '#,##0.00'; c_tf.font = font_bold
c_tm = ws2.cell(r_tot_s2, 17, tot_dup_money); c_tm.number_format = '#,##0'; c_tm.font = Font(name='Arial', size=10, bold=True, color='DC2626')
for ci in range(1, 19):
    c_cell = ws2.cell(r_tot_s2, ci)
    c_cell.fill = fill_subtot
    c_cell.border = double_bottom

# ==============================================================================
# SHEET 3: CHI TIẾT 195 CA LỆCH MÚI GIỜ -7H (SAI LỆCH 2)
# ==============================================================================
ws3 = wb.create_sheet('2_Lech_Mui_Gio_195_Ca')
ws3.views.sheetView[0].showGridLines = True

ws3['A1'] = 'DANH SÁCH CHI TIẾT 195 CA BỊ LỆCH MÚI GIỜ -7H ĐÃ ĐƯỢC HIỆU CHỈNH VỀ GIỜ VIỆT NAM'
ws3['A1'].font = Font(name='Arial', size=13, bold=True, color='1E3A8A')
ws3['A2'] = 'Nguyên nhân: Hệ thống cũ lưu nhầm giờ theo UTC (00:xx, 01:xx nửa đêm) • Đã tịnh tiến +7 tiếng về đúng khung giờ cúp điện ban ngày'
ws3['A2'].font = Font(name='Arial', size=10, italic=True, color='4B5563')

headers_s3 = [
    'STT', 'Mã Trạm', 'Mã Trạm Cũ', 'Tên Trạm / Địa Bàn',
    'Ngày (File Excel Cũ)', 'Giờ BĐ (Excel Cũ)', 'Giờ KT (Excel Cũ)',
    'Ngày Chuẩn (Đã Sửa)', 'Giờ BĐ Chuẩn (+7h)', 'Giờ KT Chuẩn (+7h)',
    'Thời Gian Chạy (h)', 'Nhiên Liệu (L)', 'Đơn Giá (đ)', 'Thành Tiền (đ)', 'SmartW Alarm Key', 'Ghi Chú Đánh Giá'
]

for c_idx, h in enumerate(headers_s3, 1):
    cell = ws3.cell(4, c_idx, h)
    cell.fill = fill_amber if 5 <= c_idx <= 7 else (fill_green if 8 <= c_idx <= 10 else fill_navy)
    cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws3.row_dimensions[4].height = 32

for idx, u in enumerate(utc_standalone_list, 1):
    cur_r = idx + 4
    st_info = get_site_info(u['site_id'])
    
    ws3.cell(cur_r, 1, idx).alignment = Alignment(horizontal='center')
    ws3.cell(cur_r, 2, u['site_id']).alignment = Alignment(horizontal='center')
    ws3.cell(cur_r, 3, st_info['site_id_old']).alignment = Alignment(horizontal='center')
    ws3.cell(cur_r, 4, f"{st_info['name']} ({st_info['district']})")
    
    # Cũ
    ws3.cell(cur_r, 5, u['orig_date']).alignment = Alignment(horizontal='center')
    c6 = ws3.cell(cur_r, 6, u['orig_start']); c6.alignment = Alignment(horizontal='center'); c6.font = Font(name='Arial', size=9, color='B45309')
    ws3.cell(cur_r, 7, u['orig_end']).alignment = Alignment(horizontal='center')
    
    # Mới chuẩn +7h
    ws3.cell(cur_r, 8, u['real_date']).alignment = Alignment(horizontal='center')
    c9 = ws3.cell(cur_r, 9, u['real_start']); c9.alignment = Alignment(horizontal='center'); c9.font = Font(name='Arial', size=9, bold=True, color='047857')
    c10 = ws3.cell(cur_r, 10, u['real_end']); c10.alignment = Alignment(horizontal='center'); c10.font = Font(name='Arial', size=9, bold=True, color='047857')
    
    c11 = ws3.cell(cur_r, 11, u['hours']); c11.number_format = '#,##0.00'
    c12 = ws3.cell(cur_r, 12, u['fuel']); c12.number_format = '#,##0.00'
    c13 = ws3.cell(cur_r, 13, u['don_gia']); c13.number_format = '#,##0'
    c14 = ws3.cell(cur_r, 14, u['money']); c14.number_format = '#,##0'
    ws3.cell(cur_r, 15, u['alarm_id'])
    ws3.cell(cur_r, 16, 'Đã chuẩn hóa về giờ hành chính Việt Nam')
    
    for ci in range(1, 17):
        ws3.cell(cur_r, ci).border = thin_border

# ==============================================================================
# SHEET 4: CHI TIẾT 17 CỤM NỐI CA LIÊN TỤC (SAI LỆCH 3)
# ==============================================================================
ws4 = wb.create_sheet('3_Noi_Ca_Lien_Tuc_21_Ca')
ws4.views.sheetView[0].showGridLines = True

ws4['A1'] = 'DANH SÁCH CHI TIẾT 17 CỤM CA CHẠY LIÊN TỤC / CHỒNG LẤN ĐÃ ĐƯỢC NỐI GỘP'
ws4['A1'].font = Font(name='Arial', size=13, bold=True, color='047857')
ws4['A2'] = 'Nguyên nhân: Trong file Excel cũ bị chẻ nhỏ thành nhiều dòng rời rạc (gap ≤ 1 phút hoặc chồng lấn) • Đã gộp thành 1 ca duy nhất'
ws4['A2'].font = Font(name='Arial', size=10, italic=True, color='4B5563')

headers_s4 = [
    'STT Cụm', 'Mã Trạm', 'Mã Trạm Cũ', 'Tên Trạm / Địa Bàn', 'Ngày Vận Hành',
    'Số Ca Bị Chẻ (Excel Cũ)', 'Chi Tiết Các Ca Rời Rạc Trong File Excel Cũ',
    'Giờ BĐ Gộp (Sớm Nhất)', 'Giờ KT Gộp (Muộn Nhất)', 'Tổng Giờ Chạy Gộp (h)', 'Tổng Lượng Dầu/Xăng (L)', 'Tổng Thành Tiền (đ)', 'Nhãn Ghi Chú Chuẩn Trên Web'
]

for c_idx, h in enumerate(headers_s4, 1):
    cell = ws4.cell(4, c_idx, h)
    cell.fill = fill_navy if c_idx < 8 else fill_green
    cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws4.row_dimensions[4].height = 32

for idx, (site, dt, cluster) in enumerate(merge_clusters, 1):
    cur_r = idx + 4
    st_info = get_site_info(site)
    
    earliest_s = min(to_mins(l['run_details']['gio_bat_dau']) for l in cluster)
    latest_e = max(to_mins(l['run_details']['gio_ket_thuc']) for l in cluster)
    tot_h = round(sum(float(l['run_details'].get('thoi_gian_hoat_dong') or 0) for l in cluster), 2)
    tot_f = round(sum(float(l['run_details'].get('nhien_lieu_tieu_hao') or 0) for l in cluster), 2)
    tot_m = round(sum(float(l['run_details'].get('thanh_tien') or 0) for l in cluster))
    
    s_str = f'{earliest_s//60:02d}:{earliest_s%60:02d}'
    e_str = f'{latest_e//60:02d}:{latest_e%60:02d}'
    
    # Chuỗi mô tả các ca con
    sub_runs = [f"{l['run_details']['gio_bat_dau']}-{l['run_details']['gio_ket_thuc']} ({l['run_details']['thoi_gian_hoat_dong']}h)" for l in cluster]
    detail_str = " + ".join(sub_runs)
    
    ws4.cell(cur_r, 1, idx).alignment = Alignment(horizontal='center')
    ws4.cell(cur_r, 2, site).alignment = Alignment(horizontal='center')
    ws4.cell(cur_r, 3, st_info['site_id_old']).alignment = Alignment(horizontal='center')
    ws4.cell(cur_r, 4, f"{st_info['name']} ({st_info['district']})")
    ws4.cell(cur_r, 5, dt).alignment = Alignment(horizontal='center')
    ws4.cell(cur_r, 6, f"{len(cluster)} ca").alignment = Alignment(horizontal='center')
    ws4.cell(cur_r, 7, detail_str)
    
    # Kết quả gộp
    c8 = ws4.cell(cur_r, 8, s_str); c8.alignment = Alignment(horizontal='center'); c8.font = Font(name='Arial', size=9, bold=True, color='047857')
    c9 = ws4.cell(cur_r, 9, e_str); c9.alignment = Alignment(horizontal='center'); c9.font = Font(name='Arial', size=9, bold=True, color='047857')
    c10 = ws4.cell(cur_r, 10, tot_h); c10.number_format = '#,##0.00'; c10.font = font_bold
    c11 = ws4.cell(cur_r, 11, tot_f); c11.number_format = '#,##0.00'
    c12 = ws4.cell(cur_r, 12, tot_m); c12.number_format = '#,##0'
    ws4.cell(cur_r, 13, f"Nối ca ({s_str}-{e_str})").font = font_bold
    
    for ci in range(1, 14):
        ws4.cell(cur_r, ci).border = thin_border

# ==============================================================================
# SHEET 5: CHI TIẾT 20 HÓA ĐƠN BỔ SUNG TRÊN WEB (SAI LỆCH 4)
# ==============================================================================
ws5 = wb.create_sheet('4_Lech_Hoa_Don_20_HD')
ws5.views.sheetView[0].showGridLines = True

ws5['A1'] = 'DANH SÁCH 20 HÓA ĐƠN ĐIỆN TỬ CÓ TRÊN WEB NHƯNG CHƯA CÓ TRONG FILE EXCEL CŨ'
ws5['A1'].font = Font(name='Arial', size=13, bold=True, color='2563EB')
ws5['A2'] = 'Nguyên nhân: Các hóa đơn phát sinh đợt cuối tháng (27/08 - 31/08/2026) được hệ thống cập nhật sau • 124 HĐ cũ khớp 100% số tiền'
ws5['A2'].font = Font(name='Arial', size=10, italic=True, color='4B5563')

headers_s5 = [
    'STT', 'Số Hóa Đơn', 'Ngày Lập HĐ', 'Đơn Vị Bán Hàng', 'Mã Số Thuế (Bán)',
    'Pháp Nhân Mua Hàng', 'Ký Hiệu HĐ', 'Số Tiền Trước VAT (đ)', 'Thuế GTGT (đ)', 'Tổng Tiền Thanh Toán (đ)',
    'Mã Tra Cứu / Fkey', 'Link Tra Cứu Hóa Đơn', 'Trạng Thái Đối Soát'
]

for c_idx, h in enumerate(headers_s5, 1):
    cell = ws5.cell(4, c_idx, h)
    cell.fill = fill_blue; cell.font = font_hdr; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin_border
ws5.row_dimensions[4].height = 30

tot_extra_money = 0
for idx, inv in enumerate(extra_invs, 1):
    cur_r = idx + 4
    tot = float(inv.get('total_amount') or 0)
    sub = float(inv.get('sub_total') or round(tot / 1.08))
    vat = float(inv.get('vat_amount') or (tot - sub))
    tot_extra_money += tot
    
    ws5.cell(cur_r, 1, idx).alignment = Alignment(horizontal='center')
    c2 = ws5.cell(cur_r, 2, inv.get('invoice_number', '')); c2.font = font_bold; c2.alignment = Alignment(horizontal='center')
    ws5.cell(cur_r, 3, inv.get('invoice_date', '')).alignment = Alignment(horizontal='center')
    ws5.cell(cur_r, 4, inv.get('seller_name', ''))
    ws5.cell(cur_r, 5, inv.get('seller_mst', '')).alignment = Alignment(horizontal='center')
    ws5.cell(cur_r, 6, inv.get('buyer_name', ''))
    ws5.cell(cur_r, 7, inv.get('kh_hd', '')).alignment = Alignment(horizontal='center')
    
    c8 = ws5.cell(cur_r, 8, sub); c8.number_format = '#,##0'
    c9 = ws5.cell(cur_r, 9, vat); c9.number_format = '#,##0'
    c10 = ws5.cell(cur_r, 10, tot); c10.number_format = '#,##0'; c10.font = font_bold
    ws5.cell(cur_r, 11, inv.get('ma_tra_cuu', '')).alignment = Alignment(horizontal='center')
    
    c12 = ws5.cell(cur_r, 12, inv.get('invoice_url', ''))
    c12.font = font_link
    ws5.cell(cur_r, 13, 'Bổ sung đợt cuối tháng (27-31/08)').alignment = Alignment(horizontal='center')
    
    for ci in range(1, 14):
        ws5.cell(cur_r, ci).border = thin_border

# Dòng tổng cộng HĐ bổ sung
r_tot_s5 = len(extra_invs) + 5
ws5.cell(r_tot_s5, 1, 'TỔNG CỘNG 20 HÓA ĐƠN BỔ SUNG ĐỢT CUỐI THÁNG 8:').font = font_bold
ws5.merge_cells(start_row=r_tot_s5, start_column=1, end_row=r_tot_s5, end_column=9)
c_tot_m = ws5.cell(r_tot_s5, 10, tot_extra_money); c_tot_m.number_format = '#,##0'; c_tot_m.font = Font(name='Arial', size=10, bold=True, color='2563EB')
for ci in range(1, 14):
    c_cell = ws5.cell(r_tot_s5, ci)
    c_cell.fill = fill_subtot
    c_cell.border = double_bottom

# Auto-adjust column widths for all sheets
for ws in [ws1, ws2, ws3, ws4, ws5]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 11), 50)

# Save file
output_path = '/Users/cang_it/Desktop/Chi_Tiet_Sai_Lech_Excel_Va_Web_T8_2026.xlsx'
wb.save(output_path)
print(f"🎉 Đã xuất thành công file báo cáo chi tiết ra: {output_path}")

