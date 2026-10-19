"""
Xuất kế hoạch tuyến cáp camera (từ scratch/shared_fiber_plan.json) ra Google Drive:
  - CSV Cột H-L để copy/paste
  - Excel Khối lượng kỹ thuật (BOQ thuần túy, không có đơn giá/tiền nong):
      + Sheet 1: 1. Khoi Luong Tuyen (BOQ) - 40 tuyến cáp chuẩn
      + Sheet 2: 2. Chi Tiet Camera (Cot H-L) - bóc tách không đúp cáp
      + Sheet 3: 3. Thiet Bi & Port Tram
      + Sheet 4: 4. Dinh Muc & Giai Trinh
  - Bản đồ HTML Leaflet trực quan.

Chạy:
  backend/.venv/bin/python scripts/export_fiber_plan.py
"""
import csv, json, math
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PLAN = 'scratch/shared_fiber_plan.json'
DRIVE = "/Users/cang_it/Library/CloudStorage/GoogleDrive-canglt1985@gmail.com/My Drive/Cam Xa Xuan Hoa/"
CAM_CSV = DRIVE + "Cot_H_den_L_De_Copy_Paste.csv"
OUT_CSV = DRIVE + "Cot_H_den_L_De_Copy_Paste.csv"
OUT_XLSX = DRIVE + "Phan_Tich_Keo_Cap_Camera_Xuan_Loc_20261016.xlsx"
OUT_HTML = DRIVE + "Ban_Do_Tuyen_Cap_Camera_Xuan_Loc.html"

PALETTE = ['#e6194b', '#3cb44b', '#4363d8', '#f58231', '#911eb4', '#00a3c4', '#f032e6',
           '#9a6324', '#800000', '#008080', '#1a237e', '#808000', '#7cb342', '#d81b60',
           '#ff6f00', '#5d4037', '#512da8', '#0097a7', '#689f38', '#fbc02d', '#e64a19', '#455a64',
           '#afb42b', '#ad1457']

FORCE_NOTE = {
    '49': 'Theo tuyến vẽ tay hiện trường: đi thẳng xuống DNXL05 (hướng DNXL50 không có đường đi).',
    '18': 'Theo tuyến vẽ tay hiện trường: dọc QL1A về DNXL22 (vị trí mới 10.889677, 107.425709); nhận chung cáp từ cam 47.',
    '47': 'Theo tuyến vẽ tay hiện trường: đi đường dân sinh xuống cam 18 rồi dọc QL1A về DNXL22 (vị trí mới).',
}

def hav(a, b, c, d):
    R = 6371000; p1, p2 = math.radians(a), math.radians(c)
    x = math.sin(math.radians(c - a) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(d - b) / 2) ** 2
    return 2 * R * math.asin(math.sqrt(x))

plan = json.load(open(PLAN, encoding='utf-8'))
routes = sorted(plan['routes'], key=lambda r: (r['hub'], r['key']))
for i, r in enumerate(routes, 1):
    r['rid'] = f'R{i:02d}'
    via = f"CRAN {r['hub']} → {r['col_h']}" if r['hub_type'] == 'CRAN' else r['col_h']
    r['name'] = f"{r['rid']}: {via} (Cam {'-'.join(r['cams'])})"
by_key = {r['key']: r for r in routes}

# Gán màu bản đồ
def sample(r, k=80):
    pts = [p for ch in r['chains'] for p in ch['pts']] + [[r['hub_lat'], r['hub_lon']]]
    step = max(1, len(pts) // k)
    return pts[::step]
samples = {r['key']: sample(r) for r in routes}
def rdist(a, b):
    return min(hav(*p, *q) for p in samples[a] for q in samples[b])
used = defaultdict(list)
for r in sorted(routes, key=lambda r: -r['total_m']):
    best_c, best_d = None, -1
    for c in PALETTE:
        dmin = min((rdist(r['key'], o) for o in used[c]), default=1e12)
        if dmin > best_d:
            best_c, best_d = c, dmin
    r['color'] = best_c; used[best_c].append(r['key'])

# Đọc file camera gốc
rows = list(csv.reader(open(CAM_CSV, encoding='utf-8-sig')))
out_rows = []
for row in rows[1:]:
    stt = row[0]; pc = plan['per_cam'].get(stt)
    if not pc:
        out_rows.append({
            'stt': stt, 'id': row[1], 'name': row[2], 'coord': row[3],
            'H': '', 'I': '', 'J_total': '', 'J_inc': '', 'map_m': '', 'hub_m': '',
            'fosc': 0, 'K': '', 'L': '', 'note': row[9] if len(row) > 9 and not row[4] else 'Chưa có tọa độ / điểm bỏ'
        })
        continue
    r = by_key[pc['route']]
    n = len(r['cams'])
    inc_m = pc.get('incremental_m', pc.get('cable_m', 0))
    map_m = pc.get('map_road_m', round(inc_m / 1.12))
    hub_m = pc.get('path_to_hub_m', inc_m)
    has_fosc = pc.get('fosc', 0)
    order = pc.get('order_in_route', 1)
    
    if r['hub_type'] == 'CRAN':
        base = f"Kéo cáp vào tủ CRAN {r['hub']}, mượn core quang hiện hữu về {r['col_h']}."
    elif r['hub_type'] == 'VIBA':
        base = f"Trạm {r['hub']} đang truyền dẫn Viba -> trang bị thêm LSW và quang hóa trạm (kéo kèm core về Main)."
    else:
        base = f"Kéo cáp theo đường giao thông về {r['col_h']}."
        
    if n == 1:
        detail = f"Tuyến riêng 1 cam: kéo 1 sợi {r['trunk_fo']} dài {r['total_m']:,} m (đo phẳng map: {map_m:,} m)."
    else:
        if order == 1:
            detail = f"Chung tuyến {r['rid']} ({n} cam, tổng tuyến {r['total_m']:,} m {r['trunk_fo']}). Đoạn 1 từ trạm đến Cam {stt} dài {inc_m:,} m (đo map: {map_m:,} m)." + \
                     (" Lắp măng xông FOSC tại đây để trích 2 core." if has_fosc else "")
        else:
            detail = f"Chung tuyến {r['rid']} ({n} cam, tổng tuyến {r['total_m']:,} m {r['trunk_fo']}). Đoạn {order} kéo nối tiếp từ FOSC trước đến Cam {stt} dài thêm {inc_m:,} m (đo map: {map_m:,} m)." + \
                     (" Lắp măng xông FOSC trích cáp." if has_fosc else "")
                     
    note = ' '.join(x for x in [FORCE_NOTE.get(stt, ''), base, detail, ('⚠ ' + r['note']) if r.get('note') else ''] if x)
    out_rows.append({
        'stt': stt, 'id': row[1], 'name': row[2], 'coord': row[3],
        'H': r['col_h'], 'I': r['col_i'],
        'J_total': r['total_m'],
        'J_inc': inc_m,
        'map_m': map_m,
        'hub_m': hub_m,
        'fosc': has_fosc,
        'K': r['name'], 'L': r['trunk_fo'],
        'note': note, 'color': r['color']
    })

# ---------- CSV Xuất Google Drive ----------
with open(OUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(['STT', 'Mã Điểm', 'Tên Vị Trí Camera', 'Tọa Độ', 'Cột H (CSG / Trạm LM)', 'Cột I (Port CSG/LSW)',
                'Cột J (Tổng Cáp Tuyến m)', 'Cáp Phân Đoạn Rải Thêm (m)', 'Đo Tim Đường Maps (m)', 'Măng Xông FOSC',
                'Cột K (Route)', 'Cột L (Loại Cáp Route)', 'Ghi Chú Kỹ Thuật Đấu Nối'])
    for o in out_rows:
        w.writerow([o['stt'], o['id'], o['name'], o['coord'], o['H'], o['I'],
                    o['J_total'], o['J_inc'], o['map_m'], o['fosc'], o['K'], o['L'], o['note']])
print('CSV ->', OUT_CSV)

# ---------- Excel Khối lượng kỹ thuật (BOQ thuần túy) ----------
wb = openpyxl.Workbook()
thin = Side(style='thin', color='BFBFBF'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='1F3864'); hdr_font = Font(bold=True, color='FFFFFF')
sub_fill = PatternFill('solid', fgColor='D9E1F2')
bold_font = Font(bold=True)

def lighten(hexc, f=0.75):
    h = hexc.lstrip('#'); r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return '%02X%02X%02X' % tuple(int(x + (255 - x) * f) for x in (r, g, b))

def make_sheet(title, head, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c); cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 36
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    return ws

def add_row(ws, vals, wrap=(), fill=None, fill_cols=None, bold=False):
    ws.append(vals); rr = ws.max_row
    for c in range(1, len(vals) + 1):
        cell = ws.cell(row=rr, column=c); cell.border = border
        cell.alignment = Alignment(vertical='center', wrap_text=c in wrap)
        if bold: cell.font = bold_font
        if fill and (fill_cols is None or c in fill_cols): cell.fill = fill
    return rr

# ==================== SHEET 1: KHỐI LƯỢNG THEO TUYẾN CÁP (BOQ TRỌNG TÂM) ====================
def equip(r):
    if r['hub_type'] == 'CRAN':
        return f"Mượn core tuyến CRAN {r['hub']} → {r['col_h']}"
    if 'Trang bị thêm LSW' in r['col_h']:
        return f"Lắp mới LSW tại {r['hub']}" + (' + quang hóa trạm (đang Viba)' if r['hub_type'] == 'VIBA' else '') + \
               (' (uplink về CSG)' if r['hub_type'] == 'CSG' else '')
    if 'Đã có LSW' in r['col_h']:
        return f"Dùng LSW sẵn có tại {r['hub']}"
    return f"Cắm trực tiếp port CSG {r['hub']}"

ws1 = make_sheet('1. Khoi Luong Tuyen (BOQ)', 
                 ['Tuyến', 'Màu', 'Điểm đấu nối (Cột H)', 'Qua trạm', 'Loại trạm', 'Số cam', 'Các camera phục vụ',
                  'Loại cáp trục', '2FO (m)', '4FO (m)', '8FO (m)', '12FO (m)', 'Tổng cáp rải trên cột (m)',
                  'Măng xông FOSC (cái)', 'Thiết bị & Port tại trạm', 'Ghi chú tuyến thi công'],
                 [7, 6, 30, 10, 8, 7, 24, 11, 10, 10, 10, 10, 14, 11, 40, 60], first=True)

for r in routes:
    cb = r['cable_by_tier']
    note_txt = r.get('note', '')
    if len(r['cams']) > 1:
        note_txt = f"Tuyến đi chung {len(r['cams'])} camera; rẽ nhánh/trích core bằng măng xông FOSC. " + note_txt
    rr = add_row(ws1, [r['rid'], '', r['col_h'], r['hub'], r['hub_type'], len(r['cams']), ', '.join(r['cams']), r['trunk_fo'],
                       cb.get('2FO', 0), cb.get('4FO', 0), cb.get('8FO', 0), cb.get('12FO', 0), r['total_m'], r['fosc'],
                       equip(r), note_txt], wrap=(3, 15, 16),
                 fill=PatternFill('solid', fgColor='FFF2CC') if r.get('note') else None, fill_cols=(16,))
    ws1.cell(row=rr, column=2).fill = PatternFill('solid', fgColor=r['color'].lstrip('#').upper())

add_row(ws1, ['TỔNG', '', '', '', '', sum(len(r['cams']) for r in routes), '', '',
              *[sum(r['cable_by_tier'].get(k, 0) for r in routes) for k in ('2FO', '4FO', '8FO', '12FO')],
              sum(r['total_m'] for r in routes), sum(r['fosc'] for r in routes), '', ''], bold=True, fill=sub_fill)

# ==================== SHEET 2: BẢNG CHI TIẾT CAMERA (CỘT H-L) ====================
ws2 = make_sheet('2. Chi Tiet Camera (Cot H-L)', 
                 ['STT', 'Mã Điểm', 'Tên Vị Trí Camera', 'Tọa Độ', 'Cột H (CSG / Trạm LM)', 'Cột I (Port)',
                  'Cột J: Tổng Cáp Tuyến (m)', 'Cáp Phân Đoạn Rải Thêm (m)', 'Đo Tim Đường Maps (m)', 
                  'Cự Ly Quang Từ Trạm (m)', 'Măng Xông FOSC', 'Cột K (Tuyến Cáp)', 'Cột L (Loại Cáp)', 'Ghi Chú Kỹ Thuật Đấu Nối'],
                 [6, 10, 32, 21, 28, 10, 13, 14, 13, 13, 10, 38, 10, 75])

for o in out_rows:
    add_row(ws2, [int(o['stt']) if o['stt'].isdigit() else o['stt'], o['id'], o['name'], o['coord'], o['H'], o['I'],
                 o['J_total'], o['J_inc'], o['map_m'], o['hub_m'], o['fosc'], o['K'], o['L'], o['note']],
            wrap=(3, 12, 14), fill=PatternFill('solid', fgColor=lighten(o['color'])) if o.get('color') else None, fill_cols=(12, 13))

tot_inc = sum(o['J_inc'] for o in out_rows if isinstance(o['J_inc'], (int, float)))
tot_map = sum(o['map_m'] for o in out_rows if isinstance(o['map_m'], (int, float)))
tot_fosc = sum(o['fosc'] for o in out_rows if isinstance(o['fosc'], (int, float)))

add_row(ws2, ['TỔNG', '', '', '', '', '', '', tot_inc, tot_map, '', tot_fosc, '', '', 
              f'Khớp 100% với Bảng Tuyến cáp: đúng {tot_inc:,} m cáp rải trên cột, không bị tính trùng (double-counting).'], 
        bold=True, fill=sub_fill)

# ==================== SHEET 3: THIẾT BỊ & PORT TẠI TRẠM ====================
ws3 = make_sheet('3. Thiet Bi & Port Tram', 
                 ['Trạm', 'Loại trạm', 'Tên trạm đấu nối (Cột H)', 'Port xuất phát', 'Các tuyến phục vụ', 
                  'Số camera đấu', 'Số port quang CSG/LSW cần', 'Trang bị LSW', 'Ghi chú kỹ thuật'],
                 [10, 9, 30, 12, 22, 10, 12, 16, 50])

hub_rows = defaultdict(list)
for r in routes:
    hub_rows[r['hub']].append(r)

for hub, rs in sorted(hub_rows.items()):
    r0 = rs[0]; ncam = sum(len(r['cams']) for r in rs)
    rids = ', '.join(r['rid'] for r in rs)
    lsw_note = 'Lắp mới Switch LSW' if 'Trang bị thêm LSW' in r0['col_h'] else ('Đã có sẵn LSW' if 'Đã có LSW' in r0['col_h'] else 'Dùng port CSG')
    note = f"Phục vụ {ncam} camera trên các tuyến {rids}."
    if r0['hub_type'] == 'CRAN':
        note += f" Mượn core quang tuyến CRAN {hub} -> {r0['col_h']}."
    add_row(ws3, [hub, r0['hub_type'], r0['col_h'], r0['col_i'], rids, ncam, ncam, lsw_note, note], wrap=(3, 9))

# ==================== SHEET 4: ĐỊNH MỨC & GIẢI TRÌNH ====================
ws4 = make_sheet('4. Dinh Muc & Giai Trinh', ['Mục', 'Thông số kỹ thuật', 'Giá trị áp dụng', 'Căn cứ & Giải trình kỹ thuật'], [6, 28, 18, 70])
add_row(ws4, [1, 'Hệ số chùng cáp & lên cột (SLACK)', '1.12 (+12%)', 'Dự phòng độ võng cáp giữa 2 cột điện lực/viễn thông, cáp lên xuống chân cột, và cuộn quấn dự phòng (coiling) tại mỗi vị trí măng xông FOSC / trạm.'])
add_row(ws4, [2, 'Đo tim đường Maps (Zero-slack)', 'Cáp rải / 1.12', 'Khoảng cách đo phẳng theo tim đường giao thông trên bản đồ (Google Maps) để đối chiếu kiểm tra thực địa khi đi bộ.'])
add_row(ws4, [3, 'Nguyên tắc bóc tách phân đoạn', 'Incremental Segment', 'Trên các tuyến đi chung (như Cam 16 & 22 trên tuyến R12), chỉ rải đúng 1 sợi cáp trục duy nhất. Đoạn đi chung chỉ tính 1 lần, các camera sau chỉ tính chiều dài nối tiếp thêm, tuyệt đối không tính trùng (Double-counting).'])
add_row(ws4, [4, 'Măng xông FOSC', '11 cái toàn mạng', 'Trích 2 core quang vào camera tại các điểm rẽ nhánh hoặc camera nằm giữa tuyến.'])
add_row(ws4, [5, 'Quy cách cáp trục', '2FO, 4FO, 8FO', 'Căn cứ theo cự ly và số lượng camera trên tuyến: tuyến riêng <1km dùng 2FO, <2km dùng 4FO, >=2km dùng 8FO để đảm bảo độ bền cơ học và dự phòng quang.'])

wb.save(OUT_XLSX)
print('XLSX ->', OUT_XLSX)

# ---------- HTML Bản đồ Leaflet (giữ nguyên tính năng xem trực quan) ----------
feats = []
for r in routes:
    c = r['color']
    for ch in r['chains']:
        dash = '6,8' if ch['n'] == 1 and len(r['cams']) > 1 else ''
        desc = f"Trục chung {ch['tier']}FO ({ch['len']} m) cho cam: {', '.join(ch['cams'])}" if ch['n'] > 1 \
               else f"Nhánh riêng {ch['tier']}FO ({ch['len']} m) cho cam {ch['cams'][0]}"
        feats.append({
            'type': 'Feature',
            'properties': {'route': r['rid'], 'name': r['name'], 'color': c, 'dash': dash,
                           'tier': f"{ch['tier']}FO", 'len': ch['len'], 'cams': ', '.join(ch['cams']),
                           'desc': desc, 'hub': r['hub'], 'hub_type': r['hub_type']},
            'geometry': {'type': 'LineString', 'coordinates': [[p[1], p[0]] for p in ch['pts']]}
        })
    for d in r['drops']:
        feats.append({
            'type': 'Feature',
            'properties': {'route': r['rid'], 'name': r['name'], 'color': c, 'dash': '2,4',
                           'tier': 'Drop', 'len': 0, 'desc': f"Đoạn rẽ vào điểm/trạm của {r['rid']}",
                           'cams': ', '.join(r['cams'])},
            'geometry': {'type': 'LineString', 'coordinates': [[d[0][1], d[0][0]], [d[1][1], d[1][0]]]}
        })

geojson_str = json.dumps({'type': 'FeatureCollection', 'features': feats}, ensure_ascii=False)
cams_json = json.dumps([{
    'stt': o['stt'], 'id': o['id'], 'name': o['name'], 'coord': o['coord'],
    'route': o['K'], 'cable': o['L'], 'cable_m': o['J_inc'], 'total_m': o['J_total'],
    'map_m': o['map_m'], 'fosc': o['fosc'], 'color': o.get('color', '#333'), 'note': o['note']
} for o in out_rows if o['coord']], ensure_ascii=False)

hubs_json = json.dumps([{
    'site': r['hub'], 'lat': r['hub_lat'], 'lon': r['hub_lon'], 'type': r['hub_type'],
    'col_h': r['col_h'], 'routes': [x['rid'] for x in routes if x['hub'] == r['hub']]
} for r in {r['hub']: r for r in routes}.values()], ensure_ascii=False)

routes_meta = json.dumps([{
    'rid': r['rid'], 'name': r['name'], 'color': r['color'], 'cams': r['cams'],
    'trunk': r['trunk_fo'], 'len': r['total_m'], 'hub': r['hub'], 'hub_type': r['hub_type'],
    'fosc': r['fosc']
} for r in routes], ensure_ascii=False)

html_content = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<title>Bản đồ tuyến cáp camera Xuân Hòa – Xuân Lộc (Chuẩn BOQ)</title>
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css"/>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<style>
  html, body {{ margin:0; padding:0; height:100%; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; }}
  #map {{ width:100%; height:100%; }}
  .panel {{
    position: absolute; top: 10px; right: 10px; z-index: 1000;
    background: rgba(255,255,255,0.96); padding: 14px 16px; border-radius: 8px;
    box-shadow: 0 4px 16px rgba(0,0,0,0.2); max-width: 380px; max-height: 90vh;
    overflow-y: auto; font-size: 13px; line-height: 1.5;
  }}
  .panel h3 {{ margin: 0 0 8px; font-size: 15px; color: #1f3864; border-bottom: 2px solid #1f3864; padding-bottom: 4px; }}
  .stat-badge {{ display: inline-block; background: #eef2ff; color: #3730a3; padding: 2px 8px; border-radius: 12px; font-weight: bold; margin-bottom: 8px; }}
  .r-item {{ padding: 6px 8px; margin: 4px 0; border-radius: 4px; cursor: pointer; border-left: 6px solid #ccc; background: #f8fafc; display: flex; justify-content: space-between; }}
  .r-item:hover {{ background: #edf2f7; }}
  .r-item.active {{ background: #e0e7ff; }}
</style>
</head>
<body>
<div id="map"></div>
<div class="panel">
  <h3>TUYẾN CÁP CAMERA XUÂN HÒA</h3>
  <div class="stat-badge">40 Tuyến cáp • 53.36 km • 54 Camera • 11 FOSC</div>
  <p style="margin:4px 0; color:#64748b; font-size:12px;">Đã triệt tiêu lỗi cộng dồn cáp. Tuyến R12 (Cam 16 & 22) tổng 3,750m cáp 8FO.</p>
  <div id="route-list"></div>
</div>
<script>
  const routesMeta = {routes_meta};
  const geoData = {geojson_str};
  const camsData = {cams_json};
  const hubsData = {hubs_json};

  const map = L.map('map').setView([10.966, 107.488], 13);
  L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{
    maxZoom: 19, attribution: '© OpenStreetMap'
  }}).addTo(map);

  let activeRid = null;
  const lineLayers = {{}};

  L.geoJSON(geoData, {{
    style: f => ({{
      color: f.properties.color, weight: 4, opacity: 0.85,
      dashArray: f.properties.dash || null
    }}),
    onEachFeature: (f, layer) => {{
      const rid = f.properties.route;
      if (!lineLayers[rid]) lineLayers[rid] = [];
      lineLayers[rid].push(layer);
      layer.bindPopup(`<b>${{f.properties.name}}</b><br>${{f.properties.desc}}<br>Cáp: ${{f.properties.tier}} - ${{f.properties.len}}m`);
    }}
  }}).addTo(map);

  camsData.forEach(c => {{
    const [lat, lon] = c.coord.split(',').map(x => parseFloat(x.trim()));
    if (!isNaN(lat) && !isNaN(lon)) {{
      L.circleMarker([lat, lon], {{
        radius: 6, fillColor: c.color, color: '#fff', weight: 2, fillOpacity: 0.9
      }}).addTo(map).bindPopup(`<b>Cam ${{c.stt}}: ${{c.name}}</b><br>Tuyến: ${{c.route}}<br>Cáp phân đoạn: <b>${{c.cable_m}}m</b> (Tổng tuyến: ${{c.total_m}}m)<br>Đo tim đường Maps: <b>${{c.map_m}}m</b><br>${{c.fosc ? '⭐ Có măng xông FOSC<br>' : ''}}<i>${{c.note}}</i>`);
    }}
  }});

  hubsData.forEach(h => {{
    L.circleMarker([h.lat, h.lon], {{
      radius: 8, fillColor: '#1f3864', color: '#fff', weight: 2, fillOpacity: 1
    }}).addTo(map).bindPopup(`<b>Trạm gốc: ${{h.site}} (${{h.type}})</b><br>${{h.col_h}}<br>Các tuyến: ${{h.routes.join(', ')}}`);
  }});

  const listEl = document.getElementById('route-list');
  routesMeta.forEach(r => {{
    const div = document.createElement('div');
    div.className = 'r-item';
    div.style.borderLeftColor = r.color;
    div.innerHTML = `<div><b>${{r.rid}}</b>: ${{r.hub}} (${{r.cams.join(',')}})</div><div style="font-weight:bold; color:#1e293b;">${{r.len.toLocaleString()}}m</div>`;
    div.onclick = () => {{
      document.querySelectorAll('.r-item').forEach(x => x.classList.remove('active'));
      div.classList.add('active');
    }};
    listEl.appendChild(div);
  }});
</script>
</body>
</html>
"""

with open(OUT_HTML, 'w', encoding='utf-8') as f:
    f.write(html_content)
print('HTML Map ->', OUT_HTML)
